# -*- coding: utf-8 -*-
"""IM(Information Memorandum) 실습용 Clean 양식 생성기.

- 실제 IM의 '섹션 구성'만 참고하고, 문안·레이아웃은 새로 만들었다.
- 자산명·소재지·당사자·투자자·대주·금액 등 구체 정보는 넣지 않는다([●] 공란).
- 표는 모두 실제 표 객체, 글자는 세로 가운데 정렬, 서체 KoPubWorld돋움체.
실행: python3 build_im_practice.py
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "20261011_실습_IM작성실습양식_강의용Clean_v01.pptx")

F_L, F_M, F_B = "KoPubWorld돋움체 Light", "KoPubWorld돋움체 Medium", "KoPubWorld돋움체 Bold"
NAVY, NAVY2, BAND = "012353", "002060", "345B86"
INK, GREY, LGREY = "1F1F1F", "595959", "8C8C8C"
HEAD, SHADE, LINE = "DAE3F3", "F2F2F2", "BFBFBF"
GUIDE_BG, GUIDE_LN = "FFF8E5", "F2C14E"
ORANGE = "F78E3F"

W, H = 13.333, 7.5
L, R = 0.45, 13.333 - 0.45
BODY_W = R - L

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
SECTIONS = [
    "Executive Summary", "자산개요", "투자개요", "시장현황 및 입지분석", "투자 및 재원조달",
    "운용전략", "투자수익률 분석", "Key Risks & Mitigants", "시설현황", "Appendix",
]
page_no = [0]


# ---------------------------------------------------------------- 공통 함수
def rgb(h):
    return RGBColor.from_string(h)


def _font(run, size, font=F_L, color=INK, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)


def text(slide, x, y, w, h, lines, size=11, font=F_L, color=INK, align=PP_ALIGN.LEFT,
         bold=False, italic=False, fill=None, line=None, anchor=MSO_ANCHOR.MIDDLE, shape=None,
         margin=0.05):
    """lines: 문자열 또는 [문자열 | (문자열, {옵션})] 목록."""
    if shape is None:
        shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        _no_style(shp)
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    elif shape is not None:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(0.75)
    elif shape is not None:
        shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(margin + 0.03))
    tf.margin_top = tf.margin_bottom = Inches(margin)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        opt = {}
        if isinstance(ln, tuple):
            ln, opt = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        p.space_after = Pt(opt.get("after", 2))
        r = p.add_run()
        r.text = ln
        _font(r, opt.get("size", size), opt.get("font", font), opt.get("color", color),
              opt.get("bold", bold), opt.get("italic", italic))
    return shp


def table(slide, x, y, w, rows, col_w=None, row_h=0.30, size=9.5, head_rows=1, key_cols=0,
          align_num=True, total_last=False, head_fill=HEAD):
    nr, nc = len(rows), max(len(r) for r in rows)
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = shp.table
    # 기본 표 스타일 제거(무늬 없음)
    tblpr = tbl._tbl.tblPr
    tblpr.set("firstRow", "0"); tblpr.set("bandRow", "0")
    if col_w:
        tot = sum(col_w)
        for j, cw in enumerate(col_w):
            tbl.columns[j].width = Inches(w * cw / tot)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
        for j in range(nc):
            v = rows[i][j] if j < len(rows[i]) else ""
            c = tbl.cell(i, j)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.05)
            c.margin_top = c.margin_bottom = Inches(0.02)
            is_head = i < head_rows
            is_key = j < key_cols and not is_head
            is_total = total_last and i == nr - 1
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(head_fill if is_head else (SHADE if (is_key or is_total) else "FFFFFF"))
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            s = str(v)
            if is_head or is_key:
                p.alignment = PP_ALIGN.CENTER
            elif align_num and (s.startswith("[●]") or s[:1].isdigit() or s in ("-", "")):
                p.alignment = PP_ALIGN.CENTER
            else:
                p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = s
            guide = s.startswith("예)") or s.startswith("※")
            _font(r, size, F_B if (is_head or is_total) else (F_M if is_key else F_L),
                  NAVY2 if is_head else (LGREY if guide else INK), italic=guide)
            _cell_border(c)
    return shp


def _cell_border(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcpr.find(qn(tag))
        if ln is None:
            ln = etree.SubElement(tcpr, qn(tag))
        ln.set("w", str(int(Pt(0.5))))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(sf, qn("a:srgbClr")); clr.set("val", LINE)


def merge(tbl_shape, r1, c1, r2, c2):
    t = tbl_shape.table
    for i in range(r1, r2 + 1):
        for j in range(c1, c2 + 1):
            if (i, j) != (r1, c1):
                for para in t.cell(i, j).text_frame.paragraphs:
                    for run in list(para.runs):
                        run._r.getparent().remove(run._r)
    t.cell(r1, c1).merge(t.cell(r2, c2))


def _no_style(shp):
    """도형 기본 스타일(그림자 등 테마 효과) 제거."""
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)


def guide(slide, x, y, w, h, items, title="작성 가이드"):
    lines = [(title, {"font": F_B, "size": 10, "color": "7F6000"})]
    lines += [("• " + t, {"size": 9, "color": GREY}) for t in items]
    text(slide, x, y, w, h, lines, shape=MSO_SHAPE.RECTANGLE, fill=GUIDE_BG, line=GUIDE_LN,
         anchor=MSO_ANCHOR.TOP, margin=0.08)


def placeholder(slide, x, y, w, h, label):
    shp = text(slide, x, y, w, h, [(label, {"color": LGREY, "size": 10, "align": PP_ALIGN.CENTER})],
               shape=MSO_SHAPE.RECTANGLE, fill="FAFAFA", line=LINE)
    shp.line.dash_style = 4  # dash
    return shp


def band(slide, x, y, w, label):
    text(slide, x, y, w, 0.28, [(label, {"font": F_B, "size": 10, "color": "FFFFFF"})],
         shape=MSO_SHAPE.RECTANGLE, fill=BAND, margin=0.03)


def arrow(slide, x1, y1, x2, y2, label=None, color=GREY):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _no_style(c)
    c.line.color.rgb = rgb(color); c.line.width = Pt(1.25)
    ln = c.line._get_or_add_ln()
    te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle")
    if label:
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        text(slide, mx - 0.75, my - 0.17, 1.5, 0.3,
             [(label, {"size": 8.5, "color": GREY, "align": PP_ALIGN.CENTER})], fill="FFFFFF")


def new_slide(sec_idx, title, lead, sub=None):
    s = prs.slides.add_slide(BLANK)
    page_no[0] += 1
    # 머리말: 섹션 번호·이름 + 제목
    text(s, L, 0.28, 0.55, 0.42, [(f"{sec_idx:02d}", {"font": F_B, "size": 16, "color": "FFFFFF",
                                                      "align": PP_ALIGN.CENTER})],
         shape=MSO_SHAPE.RECTANGLE, fill=NAVY)
    text(s, L + 0.65, 0.22, 7.5, 0.28, [(SECTIONS[sec_idx - 1], {"size": 9.5, "color": LGREY})])
    text(s, L + 0.65, 0.46, 9.5, 0.36, [(title, {"font": F_B, "size": 18, "color": NAVY})])
    text(s, R - 3.2, 0.28, 3.2, 0.3, [("[자산명] Information Memorandum", {"size": 9, "color": LGREY,
                                                                         "align": PP_ALIGN.RIGHT})])
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(L), Inches(0.92), Inches(R), Inches(0.92))
    _no_style(ln)
    ln.line.color.rgb = rgb(NAVY); ln.line.width = Pt(1.5)
    # 리드(핵심 메시지)
    lead_lines = [(lead, {"size": 14, "color": INK, "font": F_M})]
    if sub:
        lead_lines.append(("• " + sub, {"size": 11, "color": GREY}))
    text(s, L, 1.0, BODY_W, 0.62 if sub else 0.45, lead_lines, anchor=MSO_ANCHOR.MIDDLE)
    # 꼬리말
    ln2 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(L), Inches(7.05), Inches(R), Inches(7.05))
    _no_style(ln2)
    ln2.line.color.rgb = rgb(LINE); ln2.line.width = Pt(0.5)
    text(s, L, 7.08, 9.5, 0.3, [("IM 작성 실습용 양식(Clean) · 특정 자산·거래·당사자와 무관 · [●]는 실습자가 채울 공란",
                                 {"size": 8, "color": LGREY})])
    text(s, R - 1.0, 7.08, 1.0, 0.3, [(str(page_no[0]), {"size": 8, "color": LGREY, "align": PP_ALIGN.RIGHT})])
    return s


# ---------------------------------------------------------------- 표지
s = prs.slides.add_slide(BLANK)
page_no[0] += 1
text(s, 0, 0, W, H, [""], shape=MSO_SHAPE.RECTANGLE, fill=NAVY)
text(s, 0.9, 1.0, 6, 0.4, [("[운용사명] │ [부서명]", {"size": 12, "color": "C9D3E3"})])
text(s, 0.9, 2.4, 11, 0.9, [("[자산명]", {"font": F_B, "size": 40, "color": "FFFFFF"})])
text(s, 0.9, 3.3, 11, 0.6, [("Information Memorandum", {"font": F_M, "size": 24, "color": "FFFFFF"})])
text(s, 0.9, 4.0, 11, 0.45, [("IM 작성 실습용 양식 (Clean Version)", {"size": 16, "color": ORANGE})])
text(s, 0.9, 5.6, 11.5, 0.9, [
    ("본 양식은 교육 목적의 빈 서식입니다. 특정 자산·거래·당사자·투자자와 무관하며, 모든 수치는 공란([●]) 또는 작성 예시입니다.",
     {"size": 10.5, "color": "C9D3E3"}),
    ("투자 권유 또는 수익을 약속하는 자료가 아닙니다.", {"size": 10.5, "color": "C9D3E3"}),
])
text(s, 0.9, 6.6, 6, 0.4, [("[YYYY.MM]   < Strictly Confidential >", {"size": 11, "color": "FFFFFF"})])

# ---------------------------------------------------------------- 목차
s = prs.slides.add_slide(BLANK)
page_no[0] += 1
text(s, L, 0.35, 6, 0.6, [("Table of Contents", {"font": F_B, "size": 26, "color": NAVY})])
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(L), Inches(1.05), Inches(R), Inches(1.05))
_no_style(ln)
ln.line.color.rgb = rgb(NAVY); ln.line.width = Pt(1.5)
toc_desc = [
    "투자 하이라이트 · 핵심 지표", "물건 개요 · 사진 · 위치", "사업구조도 · 투자구조 · 추진일정",
    "수급 · 임대료 · 입지 · 비교사례", "Sources & Uses · Rent Roll", "임대차 · Re-Tenanting 전략",
    "분석가정 · 현금흐름 · 수익률 · DSCR · 민감도", "리스크 요인과 완화방안", "도면 · 시설 스펙", "용어 및 산식 · 체크리스트",
]
for i, (sec, d) in enumerate(zip(SECTIONS, toc_desc)):
    col, row = divmod(i, 5)
    x = L + 0.2 + col * 6.3
    y = 1.5 + row * 1.0
    text(s, x, y, 0.7, 0.6, [(f"{i+1:02d}", {"font": F_B, "size": 20, "color": "FFFFFF", "align": PP_ALIGN.CENTER})],
         shape=MSO_SHAPE.RECTANGLE, fill=NAVY if i % 2 == 0 else BAND)
    text(s, x + 0.85, y, 5.2, 0.6, [(sec, {"font": F_B, "size": 15, "color": NAVY}),
                                     (d, {"size": 10, "color": GREY})])
text(s, L, 6.6, BODY_W, 0.35, [("※ 실제 IM의 일반적인 목차 체계를 따랐습니다. 자산 유형(오피스·리테일·물류 등)에 따라 섹션을 더하거나 뺄 수 있습니다.",
                                 {"size": 9.5, "color": LGREY})])

# ---------------------------------------------------------------- 01 Executive Summary (지표)
s = new_slide(1, "Executive Summary", "[투자 기회를 한 문장으로: 예) 핵심 입지의 신축 자산을 감정가 대비 할인된 가격에 매입할 기회]",
              "[보조 메시지: 안정적인 임대 현황 및 향후 가치 상승 요인]")
kpis = [("매입가", "[●]억원", "감정가 [●]억원 대비 [▼●]%"), ("Cap. Rate", "[●]%", "NOI ÷ [매입가 / (매입가−보증금)]"),
        ("LTV", "[●]%", "[감정가 / 매입가] 기준, 대출 범위 명시"), ("WALE · 임대율", "[●]년 · [●]%", "기준일 [거래종결일]")]
for i, (k, v, d) in enumerate(kpis):
    x = L + i * (BODY_W + 0.2) / 4
    w = (BODY_W - 0.6) / 4
    text(s, x, 1.8, w, 0.32, [(k, {"font": F_B, "size": 10.5, "color": "FFFFFF", "align": PP_ALIGN.CENTER})],
         shape=MSO_SHAPE.RECTANGLE, fill=BAND)
    text(s, x, 2.12, w, 1.0, [(v, {"font": F_B, "size": 22, "color": NAVY, "align": PP_ALIGN.CENTER}),
                              (d, {"size": 9, "color": GREY, "align": PP_ALIGN.CENTER})],
         shape=MSO_SHAPE.RECTANGLE, fill="FFFFFF", line=LINE)
band(s, L, 3.35, BODY_W, "Tranche별 예상 투자수익률 [추정]")
table(s, L, 3.68, BODY_W, [
    ["구분", "투자금액(억원)", "약정 배당률(CoC)", "예상 IRR", "Equity Multiple", "매각차익 배분", "회수 방식"],
    ["[제1종 종류주식]", "[●]", "[●]%", "[●]%", "[●]x", "[●]%", "[●]"],
    ["[제2종 종류주식]", "[●]", "[●]%", "[●]%", "[●]x", "[●]%", "[●]"],
    ["[제3종 종류주식]", "[●]", "[●]", "[●]", "[●]", "[●]", "[●]"],
    ["보통주", "[●]", "잔여 배분", "[●]%", "[●]x", "[●]%", "[●]"],
], col_w=[1.6, 1.2, 1.3, 1.1, 1.2, 1.2, 2.0], row_h=0.33, key_cols=1)
band(s, L, 5.48, BODY_W, "Key Investment Highlights")
text(s, L, 5.8, BODY_W, 1.15, [
    ("① [가격] 예) 감정가·개발원가 대비 매입가 수준", {"size": 10.5}),
    ("② [입지] 예) 주요 소비지 접근성, 신규 공급 제약", {"size": 10.5}),
    ("③ [임대차] 예) 임차인 신용도, WALE, 임대율, 재계약 시 임대료 상향 여지", {"size": 10.5}),
    ("④ [구조] 예) 장기 보유가 가능한 투자구조, Tranche별 회수 경로", {"size": 10.5}),
], anchor=MSO_ANCHOR.TOP)

# ---------------------------------------------------------------- 01 Investment Highlights 4박스
s = new_slide(1, "Investment Highlights", "[투자 매력 요인을 4가지 관점에서 정리]")
boxes = [("01", "자산 스펙 · 매입 조건", ["준공연도·규모·스펙: [●]", "감정가 대비 할인율: [●]%", "평당 매입가: [●]만원 (비교사례 [●]~[●]만원)", "진입 Cap Rate [●]% / 운용기간 평균 [●]%"]),
         ("02", "입지", ["도심·IC 접근성: [●]km / [●]분", "배후 수요: [●]", "권역 임대료 [●]~[●]원/평, 공실률 [●]%", "향후 공급: [●]"]),
         ("03", "임차인 · 임대차", ["주요 임차인 비중 [●]%", "WALE [●]년 / 임대율 [●]%", "만기 도래 시점별 재계약 계획", "시장임대료 대비 현재 임대료 [●]%"]),
         ("04", "투자 구조", ["운용기간 [●]년, [사모/공모] 구조", "Tranche별 상환·회수 순위", "배당 재원 Buffer: [●]", "Exit 전략: [●]"])]
for i, (n, t, items) in enumerate(boxes):
    x = L + (i % 2) * (BODY_W / 2 + 0.1)
    y = 1.65 + (i // 2) * 2.62
    w = BODY_W / 2 - 0.1
    text(s, x, y, 0.6, 0.5, [(n, {"font": F_B, "size": 16, "color": "FFFFFF", "align": PP_ALIGN.CENTER})],
         shape=MSO_SHAPE.RECTANGLE, fill=NAVY)
    text(s, x + 0.65, y, w - 0.65, 0.5, [(t, {"font": F_B, "size": 13, "color": NAVY})],
         shape=MSO_SHAPE.RECTANGLE, fill=HEAD)
    text(s, x, y + 0.55, w, 1.95, [("• " + it, {"size": 10.5, "after": 5}) for it in items],
         shape=MSO_SHAPE.RECTANGLE, fill="FFFFFF", line=LINE, anchor=MSO_ANCHOR.MIDDLE, margin=0.12)

# ---------------------------------------------------------------- 02 자산개요
s = new_slide(2, "대상자산 개요", "[자산 위치·규모·용도를 한 문장으로 요약]")
table(s, L, 1.65, 6.6, [
    ["구분", "내용"],
    ["물건명", "[●]"], ["위치", "[시·구·동 / 지번]"], ["지역 / 지구", "[용도지역 / 지구]"], ["용도", "[업무시설 / 창고시설 / 판매시설 등]"],
    ["대지면적", "[●]㎡ ([●]평)"], ["연면적", "[●]㎡ ([●]평)"], ["임대면적 / 전용면적", "[●]㎡ / [●]㎡"],
    ["규모", "지하 [●]층 / 지상 [●]층"], ["건폐율 / 용적률", "[●]% (법정 [●]%) / [●]% (법정 [●]%)"],
    ["준공일", "[YYYY.MM.DD]"], ["주요 설비", "[주차 / 접안 / 전력용량 / 층고 / 바닥하중 등]"],
    ["임대 현황", "임대율 [●]% (주요 임차인 [●])"], ["감정평가액", "[●]억원 (평가기관 [●], 기준일 [●])"],
], col_w=[1.6, 4.0], row_h=0.355, key_cols=1)
placeholder(s, 7.25, 1.65, 5.63, 3.0, "[전경 사진]")
placeholder(s, 7.25, 4.75, 2.75, 2.2, "[위치도]")
placeholder(s, 10.13, 4.75, 2.75, 2.2, "[항공 사진]")

# ---------------------------------------------------------------- 03 사업구조도
s = new_slide(3, "사업구조도", "[투자기구 설립 및 매입·운용 구조 요약: 예) 신설 리츠를 설립하여 대상자산을 매입·운용]")
def node(x, y, w, h, title_, sub_, fill=HEAD, tcolor=NAVY):
    text(s, x, y, w, h, [(title_, {"font": F_B, "size": 11.5, "color": tcolor, "align": PP_ALIGN.CENTER}),
                         (sub_, {"size": 9, "color": GREY if fill != NAVY else "C9D3E3", "align": PP_ALIGN.CENTER})],
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill=fill, line=LINE)
node(5.42, 3.35, 2.5, 0.95, "투자기구", "[●] 위탁관리부동산투자회사", fill=NAVY, tcolor="FFFFFF")
node(5.42, 5.65, 2.5, 0.8, "대상자산", "[자산명]")
node(1.0, 3.4, 2.3, 0.85, "매도인", "[●]")
node(10.0, 3.4, 2.3, 0.85, "임차인", "[주요 임차인 외 [●]개사]")
node(1.6, 1.75, 2.6, 0.85, "Equity 투자자", "종류주식 · 보통주")
node(9.1, 1.75, 2.6, 0.85, "대주", "선순위 · 중순위")
node(5.42, 1.75, 2.5, 0.85, "자산관리회사(AMC)", "[●]")
node(1.0, 5.65, 2.3, 0.8, "자산보관·사무수탁", "[●] / [●]")
node(10.0, 5.65, 2.3, 0.8, "담보신탁 수탁자", "[●]")
arrow(s, 3.3, 3.75, 5.42, 3.75, "소유권 이전")
arrow(s, 5.42, 4.0, 3.3, 4.0, "매매대금")
arrow(s, 10.0, 3.7, 7.92, 3.7, "임대료")
arrow(s, 3.4, 2.6, 5.6, 3.35, "투자")
arrow(s, 5.42, 3.6, 3.0, 2.6, None)
text(s, 3.2, 2.85, 1.2, 0.3, [("배당", {"size": 8.5, "color": GREY, "align": PP_ALIGN.CENTER})], fill="FFFFFF")
arrow(s, 9.9, 2.6, 7.75, 3.35, "대출")
arrow(s, 7.92, 3.55, 10.3, 2.6, None)
text(s, 9.3, 2.85, 1.2, 0.3, [("이자", {"size": 8.5, "color": GREY, "align": PP_ALIGN.CENTER})], fill="FFFFFF")
arrow(s, 6.67, 2.6, 6.67, 3.35, "자산관리위탁")
arrow(s, 6.67, 4.3, 6.67, 5.65, "매입 · 운용")
guide(s, L, 6.55, BODY_W, 0.45, ["화살표마다 거래 내용(대금·권리·현금흐름)을 적고, 각 기관명과 역할을 박스 안에 채웁니다."], title="")

# ---------------------------------------------------------------- 03 투자구조 (Sources & Uses 요약)
s = new_slide(3, "투자구조도", "[총투자비 [●]억원을 Equity [●]억원 · Debt [●]억원으로 조달]")
band(s, L, 1.65, 6.1, "자금 사용 (Uses)")
band(s, 6.78, 1.65, 6.1, "자금 조달 (Sources)")
table(s, L, 2.0, 6.1, [
    ["항목", "금액(억원)", "비중"],
    ["부동산 매입금액", "[●]", "[●]%"], ["취득세 등 제세", "[●]", "[●]%"], ["매입·금융 수수료", "[●]", "[●]%"],
    ["실사·예비비", "[●]", "[●]%"], ["이자유보·운영자금", "[●]", "[●]%"], ["합계", "[●]", "100.0%"],
], col_w=[2.6, 1.6, 1.2], row_h=0.38, total_last=True)
table(s, 6.78, 2.0, 6.1, [
    ["항목", "금액(억원)", "비중", "조건"],
    ["임대보증금 승계", "[●]", "[●]%", "-"], ["선순위 대출", "[●]", "[●]%", "[●]개월 / [●]%"],
    ["중순위 대출", "[●]", "[●]%", "[●]개월 / [●]%"], ["[제1종 종류주식]", "[●]", "[●]%", "CoC [●]%"],
    ["[제2종 종류주식]", "[●]", "[●]%", "CoC [●]%"], ["[제3종 종류주식]", "[●]", "[●]%", "[●]"],
    ["보통주", "[●]", "[●]%", "잔여 배분"], ["합계", "[●]", "100.0%", ""],
], col_w=[2.0, 1.3, 1.0, 1.8], row_h=0.38, total_last=True)
guide(s, L, 5.0, 6.1, 1.95, [
    "Uses 합계 = Sources 합계가 되도록 맞춥니다.",
    "LTV는 분모(감정가 또는 매입가)와 대출 범위(보증금 포함 여부)를 함께 적습니다.",
    "Tranche별 상환 순위(대출 → 종류주식 → 보통주)를 구조도에 표시합니다.",
])
placeholder(s, 6.78, 5.5, 6.1, 1.45, "[자본구조 누적 막대: 보증금 / 선순위 / 중순위 / 종류주식 / 보통주]")

# ---------------------------------------------------------------- 03 추진일정
s = new_slide(3, "추진일정", "[목표 거래종결 시점과 총 소요기간: 예) [YYYY.MM] 종결 목표, 약 [●]주 소요]")
weeks = [f"{i}주" for i in range(1, 13)]
rows = [["구분", "업무"] + weeks]
plan = [("매도인 협의", ["경쟁입찰 / 우선협상", "MOU · 독점협상"]), ("실사", ["법률 · 물리 · 재무 · 시장 실사", "감정평가"]),
        ("투자기구", ["설립 · 영업인가(등록)", "주주간계약"]), ("Equity", ["투자자 모집 · 승인", "인수확약서(LOC) 발급"]),
        ("Loan", ["대주 심의", "대출약정 체결"]), ("종결", ["매매계약 · 잔금 · 소유권이전", ""])]
for grp, tasks in plan:
    for t in tasks:
        if t:
            rows.append([grp, t] + [""] * 12)
shp = table(s, L, 1.65, BODY_W, rows, col_w=[1.3, 2.6] + [0.62] * 12, row_h=0.42, key_cols=1)
r = 1
for grp, tasks in plan:
    n = len([t for t in tasks if t])
    if n > 1:
        merge(shp, r, 0, r + n - 1, 0)
    r += n
guide(s, L, 6.25, BODY_W, 0.7, ["해당 주차 셀에 음영을 넣어 막대를 표시하고, 주요 마일스톤(투심·계약·종결)은 ◆로 표시합니다.",
                                "※ 일정은 관계기관 승인과 실사 진행에 따라 변동될 수 있다는 문구를 함께 적습니다."], title="")

# ---------------------------------------------------------------- 04 시장현황
s = new_slide(4, "시장현황", "[시장 판단을 3가지 메시지로: 예) 공급 감소 · 임대료 상승 · 거래가격 회복]")
for i, (t, d) in enumerate([("① 공급", "[신규 공급 전망과 그 원인]"), ("② 수요 · 임대료", "[임대료 추이 · 공실률]"), ("③ 거래시장", "[거래 규모 · Cap Rate 추이]")]):
    x = L + i * (BODY_W + 0.2) / 3
    w = (BODY_W - 0.4) / 3
    text(s, x, 1.65, w, 0.38, [(t, {"font": F_B, "size": 12, "color": "FFFFFF", "align": PP_ALIGN.CENTER})],
         shape=MSO_SHAPE.RECTANGLE, fill=BAND)
    text(s, x, 2.03, w, 0.6, [("[핵심 수치: 예) 전년 대비 [▲/▼●]%]", {"font": F_B, "size": 12, "color": NAVY, "align": PP_ALIGN.CENTER}),
                              (d, {"size": 9.5, "color": GREY, "align": PP_ALIGN.CENTER})],
         shape=MSO_SHAPE.RECTANGLE, fill="FFFFFF", line=LINE)
    placeholder(s, x, 2.75, w, 1.9, "[차트]")
yrs = ["[Y-3]", "[Y-2]", "[Y-1]", "[Y]", "[Y+1]F", "[Y+2]F"]
table(s, L, 4.85, BODY_W, [
    ["구분", "단위"] + yrs + ["출처"],
    ["신규 공급", "만㎡"] + ["[●]"] * 6 + ["[●]"],
    ["재고(누적)", "만㎡"] + ["[●]"] * 6 + ["[●]"],
    ["평균 임대료", "원/평/월"] + ["[●]"] * 6 + ["[●]"],
    ["공실률", "%"] + ["[●]"] * 6 + ["[●]"],
    ["거래 Cap Rate", "%"] + ["[●]"] * 6 + ["[●]"],
], col_w=[1.6, 1.0] + [1.05] * 6 + [1.6], row_h=0.34, key_cols=1)

# ---------------------------------------------------------------- 04 입지분석
s = new_slide(4, "입지분석", "[입지 강점: 예) 도심 [●]km 이내, 주요 IC [●]분 접근]")
placeholder(s, L, 1.65, 6.6, 5.3, "[광역 위치도 · 반경 표시 · 주요 도로망]")
table(s, 7.25, 1.65, 5.63, [
    ["구분", "목적지", "거리(km)", "소요시간(분)"],
    ["IC / JC", "[●]", "[●]", "[●]"], ["IC / JC", "[●]", "[●]", "[●]"], ["도심(CBD 등)", "[●]", "[●]", "[●]"],
    ["주요 소비지", "[●]", "[●]", "[●]"], ["대중교통", "[●]", "[●]", "[●]"],
], col_w=[1.4, 1.8, 1.1, 1.3], row_h=0.36, key_cols=1)
table(s, 7.25, 4.0, 5.63, [
    ["입지 평가 항목", "평가", "근거"],
    ["접근성", "[상/중/하]", "[●]"], ["배후 수요", "[상/중/하]", "[●]"], ["인력 수급", "[상/중/하]", "[●]"],
    ["공급 제약", "[상/중/하]", "[●]"], ["규제·민원", "[상/중/하]", "[●]"],
], col_w=[1.6, 1.0, 3.0], row_h=0.36, key_cols=1)
text(s, 7.25, 6.25, 5.63, 0.6, [("※ 거리·시간은 측정 도구와 기준 시점을 각주로 적습니다.", {"size": 9, "color": LGREY})])

# ---------------------------------------------------------------- 04 비교사례
s = new_slide(4, "임대 · 매매 비교사례", "[본건 임대료·매입가가 비교사례 대비 어느 수준인지 요약]")
cases = ["사례 1", "사례 2", "사례 3", "사례 4", "사례 5", "본건"]
table(s, L, 1.65, BODY_W, [
    ["구분"] + cases,
    ["소재지"] + ["[●]"] * 6, ["사진"] + [""] * 6, ["준공연도"] + ["[●]"] * 6, ["연면적(평)"] + ["[●]"] * 6,
    ["임대료(원/평/월)"] + ["[●]"] * 6, ["관리비(원/평/월)"] + ["[●]"] * 6, ["Rent-Free(개월/년)"] + ["[●]"] * 6,
    ["Effective Rent"] + ["[●]"] * 6, ["공실률(%)"] + ["[●]"] * 6,
    ["거래가(평당, 만원)"] + ["[●]"] * 6, ["거래 Cap Rate(%)"] + ["[●]"] * 6, ["거래시점"] + ["[●]"] * 6,
], col_w=[1.8] + [1.55] * 6, row_h=0.33, key_cols=1)
shp = s.shapes[-1]
shp.table.rows[2].height = Inches(0.62)
text(s, L, 6.62, BODY_W, 0.4, [("※ Effective Rent = (월 임대료 × (12 − Rent-Free 개월) ÷ 12) + 월 관리비. 출처(시장보고서·실사보고서)와 기준 시점을 적습니다.",
                                 {"size": 9, "color": LGREY})])

# ---------------------------------------------------------------- 05 Sources & Uses 상세
s = new_slide(5, "Sources & Uses", "[총투자비 [●]억원, 매입가 대비 부대비용 [●]%]")
table(s, L, 1.65, BODY_W, [
    ["구분", "항목", "금액(백만원)", "비중", "산정 근거"],
    ["Uses", "부동산 매입금액", "[●]", "[●]%", "매매계약서상 매매대금"],
    ["", "매입수수료", "[●]", "[●]%", "AMC 매입보수 매입가의 [●]%"],
    ["", "취득세 등", "[●]", "[●]%", "과세표준의 [●]%"],
    ["", "실사 · 자문 비용", "[●]", "[●]%", "감정평가 · 법률 · 물리 · 재무 · 시장 실사"],
    ["", "금융 · 주선 수수료", "[●]", "[●]%", "Tranche별 수수료율 [●]%"],
    ["", "이자유보 · 예비비", "[●]", "[●]%", "이자 [●]개월분 등"],
    ["", "Uses 합계", "[●]", "100.0%", ""],
    ["Sources", "임대보증금", "[●]", "[●]%", "승계 보증금"],
    ["", "선순위 대출", "[●]", "[●]%", "금리 [●]%, 만기 [●]개월"],
    ["", "중순위 대출", "[●]", "[●]%", "금리 [●]%, 만기 [●]개월"],
    ["", "종류주식", "[●]", "[●]%", "종류별 CoC [●]%"],
    ["", "보통주", "[●]", "[●]%", ""],
    ["", "Sources 합계", "[●]", "100.0%", ""],
], col_w=[1.2, 2.4, 1.6, 1.0, 5.2], row_h=0.33)
shp = s.shapes[-1]
merge(shp, 1, 0, 7, 0); merge(shp, 8, 0, 13, 0)
for rr in (1, 8):
    c = shp.table.cell(rr, 0); c.fill.solid(); c.fill.fore_color.rgb = rgb(SHADE)
    for p in c.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            _font(run, 10, F_B, NAVY2)
for rr in (7, 13):
    for cc in range(1, 5):
        c = shp.table.cell(rr, cc); c.fill.solid(); c.fill.fore_color.rgb = rgb(SHADE)
        for p in c.text_frame.paragraphs:
            for run in p.runs:
                _font(run, 9.5, F_B, INK)

# ---------------------------------------------------------------- 05 Rent Roll
s = new_slide(5, "임대차 현황 (Rent Roll)", "[임대율 [●]%, WALE [●]년, 주요 임차인 비중 [●]%]")
rr_rows = [["층", "임차인", "용도", "임대면적(평)", "보증금(원/평)", "월 임대료(원/평)", "월 관리비(원/평)", "계약 만기", "인상률", "비고"]]
for f in ["[B●]", "[1F]", "[2F]", "[3F]", "[4F]", "[5F]", "[●F]", "[●F]"]:
    rr_rows.append([f, "[●]", "[●]", "[●]", "[●]", "[●]", "[●]", "[YYYY.MM]", "[●]%", ""])
rr_rows.append(["합계 / 평균", "", "", "[●]", "[●]", "[●]", "[●]", "WALE [●]년", "", ""])
table(s, L, 1.65, BODY_W, rr_rows, col_w=[0.8, 1.8, 1.0, 1.2, 1.2, 1.3, 1.3, 1.2, 0.8, 1.6], row_h=0.36, total_last=True)
placeholder(s, L, 5.4, 6.1, 1.55, "[임차인별 면적 비중 원형 차트]")
placeholder(s, 6.78, 5.4, 6.1, 1.55, "[연도별 만기 도래 면적 막대 차트]")

# ---------------------------------------------------------------- 06 운용전략
s = new_slide(6, "운용전략 — Re-Tenanting", "[만기 도래 시 임대료 조정 등 운용 전략의 핵심 목표]")
years = [f"{i}년차" for i in range(1, 11)]
table(s, L, 1.65, BODY_W, [
    ["구분"] + years,
    ["연도"] + ["[YYYY]"] * 10,
    ["임대료(원/평/월)"] + ["[●]"] * 10,
    ["관리비(원/평/월)"] + ["[●]"] * 10,
    ["Effective Rent"] + ["[●]"] * 10,
    ["임대율(%)"] + ["[●]"] * 10,
    ["주요 이벤트"] + [""] * 10,
], col_w=[1.8] + [1.0] * 10, row_h=0.38, key_cols=1)
table(s, L, 4.5, BODY_W, [
    ["전략", "주요 내용", "기대 효과", "관련 리스크"],
    ["임대료 조정", "[만기 시 시장임대료 수준으로 재계약]", "[NOI [●]% 증가]", "[임차인 이탈]"],
    ["공실 해소 · 재임대", "[신규 임차인 유치 계획]", "[임대율 [●]%]", "[다운타임 · Rent-Free]"],
    ["비용 효율화", "[관리 용역 · 에너지 비용 절감]", "[운영비 [●]% 절감]", "[서비스 품질]"],
    ["자본적 지출", "[설비 개선 · 용도 전환]", "[자산가치 제고]", "[공사비 · 공사기간]"],
], col_w=[1.8, 4.2, 2.8, 2.8], row_h=0.48, key_cols=1)

# ---------------------------------------------------------------- 07 분석가정
s = new_slide(7, "분석가정", "[현금흐름 추정에 사용한 주요 가정 요약]")
table(s, L, 1.65, 6.1, [
    ["구분", "가정", "근거"],
    ["운용기간", "[●]개월", "[●]"], ["개시일 / 종료일", "[YYYY.MM] / [YYYY.MM]", ""],
    ["임대료 인상률", "연 [●]%", "임대차계약"], ["갱신 시 임대료", "[●]원/평", "시장 실사"],
    ["공실률 · 다운타임", "[●]% · [●]개월", "[●]"], ["Rent-Free", "[●]개월/년", "[●]"],
    ["관리비 수입 / 비용", "[●] / [●]", "[●]"], ["운영비용 상승률", "연 [●]%", "[●]"],
], col_w=[1.8, 2.2, 2.1], row_h=0.4, key_cols=1)
table(s, 6.78, 1.65, 6.1, [
    ["구분", "가정", "근거"],
    ["선순위 금리", "[●]% ([고정/변동])", "금융확약서"], ["중순위 금리", "[●]%", "금융확약서"],
    ["대출 만기 · 리파이낸싱", "[●]개월 · [●]", "[●]"], ["AMC 보수", "매입 [●]% / 운용 [●]% / 매각 [●]%", ""],
    ["Exit Cap Rate", "[●]%", "[진입 Cap + [●]bp]"], ["매각 비용", "매각가의 [●]%", "[●]"],
    ["법인세 · 배당", "[배당가능이익의 90% 이상 배당 등]", "관련 법령"], ["할인율", "[●]%", "[●]"],
], col_w=[1.8, 2.4, 1.9], row_h=0.4, key_cols=1)
guide(s, L, 5.45, BODY_W, 1.5, ["모든 가정에 근거 자료(계약서·실사보고서·확약서 등)와 기준일을 적습니다.",
                                "Cap Rate의 분모(매입가 / 매입가−보증금 / 부대비용 포함 여부)와 IRR의 레버리지 전·후 구분을 명시합니다.",
                                "추정치는 [추정], 확인되지 않은 값은 [확인 필요]로 표시하고, '보장'·'확정' 같은 표현은 쓰지 않습니다."])

# ---------------------------------------------------------------- 07 현금흐름
s = new_slide(7, "추정 현금흐름 (Cash Flow)", "[운용기간 NOI 연평균 [●]% 증가, 평균 Cap Rate [●]%] [추정]")
cf_years = [f"Y{i}" for i in range(1, 11)] + ["합계"]
cf_items = ["임대료 수입", "관리비 수입", "기타 수입", "(−) 공실·Rent-Free", "총수입", "(−) 운영비용", "NOI",
            "(−) AMC 운용보수", "(−) 대출 이자", "(−) 기타 비용", "배당가능 현금흐름", "Cap Rate(매입가 기준, %)"]
rows = [["구분(백만원)"] + cf_years] + [[it] + ["[●]"] * 11 for it in cf_items]
shp = table(s, L, 1.65, BODY_W, rows, col_w=[2.2] + [0.95] * 11, row_h=0.36, key_cols=1, size=9)
for rr in (5, 7, 11):
    for cc in range(0, 12):
        c = shp.table.cell(rr, cc); c.fill.solid(); c.fill.fore_color.rgb = rgb(HEAD if rr == 7 else SHADE)
        for p in c.text_frame.paragraphs:
            for run in p.runs:
                _font(run, 9, F_B, NAVY2 if rr == 7 else INK)
text(s, L, 6.4, BODY_W, 0.5, [("※ NOI = 임대수입 + 기타수입 − 운영비용(자본적 지출 제외). 단위·반올림 기준을 표 위에 적습니다.",
                                 {"size": 9, "color": LGREY})])

# ---------------------------------------------------------------- 07 예상투자수익률
s = new_slide(7, "예상 투자수익률 (Tranche별)", "[Tranche별 예상 배당수익률과 IRR 요약] [추정]")
fy = [f"FY{i}" for i in range(1, 11)]
rows = [["Tranche", "항목"] + fy + ["합계/평균"]]
for t in ["[제1종 종류주식]", "[제2종 종류주식]", "보통주"]:
    for it in ["투자잔액", "운영배당", "매각차익(C.G)", "Yield(%)"]:
        rows.append([t, it] + ["[●]"] * 11)
shp = table(s, L, 1.65, BODY_W, rows, col_w=[1.6, 1.2] + [0.85] * 11, row_h=0.3, key_cols=2, size=8.5)
for k in range(3):
    merge(shp, 1 + k * 4, 0, 4 + k * 4, 0)
table(s, L, 5.65, BODY_W, [
    ["구분", "[제1종 종류주식]", "[제2종 종류주식]", "보통주", "전체 Equity"],
    ["IRR / Equity Multiple", "[●]% / [●]x", "[●]% / [●]x", "[●]% / [●]x", "[●]% / [●]x"],
], col_w=[2.2, 2.3, 2.3, 2.3, 2.3], row_h=0.36, key_cols=1)
text(s, L, 6.45, BODY_W, 0.45, [("※ 예상수익률은 가정에 따른 추정치이며 실제 결과와 다를 수 있습니다. 운용기간·Exit Cap·매각비용 가정을 함께 적습니다.",
                                  {"size": 9, "color": LGREY})])

# ---------------------------------------------------------------- 07 DSCR
s = new_slide(7, "DSCR 분석", "[대주 및 종류주식 기준 원리금·배당 상환능력 요약] [추정]")
yrs = [f"Y{i}" for i in range(1, 11)]
rows = [["구분(백만원)"] + yrs + ["평균", "Min"]]
for it in ["NOI", "선순위 이자", "중순위 이자", "선순위 DSCR(x)", "중순위 포함 DSCR(x)", "[제1종] 배당", "[제1종] 포함 커버리지(x)",
           "[제2종] 배당", "[제2종] 포함 커버리지(x)"]:
    rows.append([it] + ["[●]"] * 12)
shp = table(s, L, 1.65, BODY_W, rows, col_w=[2.4] + [0.85] * 12, row_h=0.36, key_cols=1, size=9)
table(s, L, 5.5, BODY_W, [
    ["재무약정", "기준", "미달 시 조치", "근거"],
    ["DSCR(이자지급력지수)", "[●]x 이상", "[이자 [●]개월분 추가 유보]", "대출약정서 제[●]조"],
    ["배당 제한", "[중순위 기준 [●]x 미만]", "[보통주 배당 금지]", "대출약정서 제[●]조"],
], col_w=[2.4, 2.4, 3.6, 2.8], row_h=0.36, key_cols=1)
text(s, L, 6.65, BODY_W, 0.3, [("※ DSCR 산식(NOI ÷ 연 원리금 / 이자)과 평가 주기를 대출약정서 정의와 일치시킵니다.", {"size": 9, "color": LGREY})])

# ---------------------------------------------------------------- 07 민감도
s = new_slide(7, "민감도 분석", "[임대료·공실률·금리·Exit Cap 변동에 따른 수익률 변화 요약] [추정]")
band(s, L, 1.65, 6.1, "① 갱신 임대료 변동 vs IRR")
table(s, L, 1.98, 6.1, [
    ["구분", "−15%", "−10%", "−5%", "Base", "+5%", "+10%", "+15%"],
    ["E.Rent(원/평)"] + ["[●]"] * 7, ["[제1종] IRR"] + ["[●]"] * 7, ["[제2종] IRR"] + ["[●]"] * 7, ["보통주 IRR"] + ["[●]"] * 7,
], col_w=[1.6] + [0.8] * 7, row_h=0.34, key_cols=1, size=9)
band(s, 6.78, 1.65, 6.1, "② 공실률 변동 vs IRR")
table(s, 6.78, 1.98, 6.1, [
    ["구분", "0%", "5%", "10%", "15%", "20%", "25%", "30%"],
    ["NOI(백만원)"] + ["[●]"] * 7, ["[제1종] IRR"] + ["[●]"] * 7, ["[제2종] IRR"] + ["[●]"] * 7, ["보통주 IRR"] + ["[●]"] * 7,
], col_w=[1.6] + [0.8] * 7, row_h=0.34, key_cols=1, size=9)
band(s, L, 3.95, 6.1, "③ 금리 × Exit Cap Rate (보통주 IRR)")
caps = ["[●−50bp]", "[●−25bp]", "Base", "[●+25bp]", "[●+50bp]"]
table(s, L, 4.28, 6.1, [["금리(행)·Exit Cap(열)"] + caps] + [[r] + ["[●]"] * 5 for r in ["−100bp", "−50bp", "Base", "+50bp", "+100bp"]],
      col_w=[1.6] + [0.9] * 5, row_h=0.34, key_cols=1, size=9)
guide(s, 6.78, 3.95, 6.1, 3.0, [
    "Base 값이 수익률 분석 페이지의 수치와 같은지 확인합니다.",
    "변동 폭(±5%, ±25bp 등)은 시장 변동성을 고려해 정합니다.",
    "임계값(배당 재원 부족, 재무약정 미달이 시작되는 지점)을 별도로 표시합니다.",
    "결과 해석은 '보장' 같은 표현 없이 '추정 시' 등으로 적습니다.",
])

# ---------------------------------------------------------------- 08 Key Risks
s = new_slide(8, "Key Risks & Mitigants", "[주요 리스크와 완화 방안 요약]")
rows = [["구분", "리스크 요인", "내용", "완화 방안", "잔여 위험"]]
for cat, items in [("A. 임차인", ["임대차 재계약 / 만기 집중", "임차인 집중도"]), ("B. 재무", ["금리 · 차환", "배당 재원 부족"]),
                   ("C. 시장", ["공급 증가", "Exit Cap Rate 상승"]), ("D. 운영 · 법률", ["자산관리 · 시설", "인허가 · 이해상충"])]:
    for it in items:
        rows.append([cat, it, "[●]", "[●]", "[상/중/하]"])
shp = table(s, L, 1.65, BODY_W, rows, col_w=[1.4, 2.4, 3.6, 3.6, 1.2], row_h=0.58, key_cols=1)
for k in range(4):
    merge(shp, 1 + k * 2, 0, 2 + k * 2, 0)

# ---------------------------------------------------------------- 09 시설현황
s = new_slide(9, "시설현황", "[건물 스펙 요약: 예) 전층 접안, 층고 [●]m, 바닥하중 [●]t/㎡]")
placeholder(s, L, 1.65, 6.1, 2.55, "[건축 단면도]")
placeholder(s, L, 4.35, 2.97, 2.6, "[층별 평면도 1]")
placeholder(s, L + 3.13, 4.35, 2.97, 2.6, "[층별 평면도 2]")
table(s, 6.78, 1.65, 6.1, [
    ["구분", "사양", "비고"],
    ["구조", "[●]", ""], ["층고(유효)", "[●]m", ""], ["바닥하중", "[●]t/㎡", ""], ["기둥 간격", "[●]m × [●]m", ""],
    ["접안시설(Dock)", "[●]개 / [램프·전층 접안 여부]", ""], ["화물 엘리베이터", "[●]대 / [●]t", ""],
    ["전력 용량", "[●]kW", ""], ["냉난방 · 공조", "[●]", ""], ["소방 · 방재", "[스프링클러 등]", ""],
    ["주차", "[●]대 (화물차 [●]대)", ""], ["인증", "[친환경 · 에너지 인증]", ""],
    ["물리실사 의견", "[중대 하자 유무 / 단기 Capex [●]억원]", "물리실사 보고서"],
], col_w=[1.6, 3.0, 1.5], row_h=0.4, key_cols=1)

# ---------------------------------------------------------------- 10 Appendix 용어·산식
s = new_slide(10, "Appendix — 용어 및 산식", "[IM에서 쓰는 지표의 정의를 통일]")
table(s, L, 1.65, BODY_W, [
    ["지표", "산식", "작성 시 유의사항"],
    ["NOI", "임대수입 + 기타수입 − 운영비용", "자본적 지출(Capex) 제외"],
    ["Cap Rate", "NOI ÷ 매입가", "분모(매입가 / 매입가−보증금)와 부대비용 포함 여부 명시"],
    ["LTV", "대출금 ÷ 감정평가액(또는 매입가)", "분모와 보증금 포함 여부 명시"],
    ["DSCR", "NOI ÷ 연 원리금(또는 이자)", "대출약정서 정의와 일치"],
    ["WALE", "임대료 가중 평균 잔여 임대기간", "기준일과 중도해지 옵션 반영 여부 명시"],
    ["공실률", "공실 면적 ÷ 임대가능 면적", "Rent-Free 구간 포함 여부 명시"],
    ["Effective Rent", "(월 임대료 × (12 − Rent-Free) ÷ 12) + 월 관리비", "시점 차이 조정 여부 명시"],
    ["CoC(Cash on Cash)", "연 배당금 ÷ 투자원금", "종류주식은 약정 배당률과 구분"],
    ["IRR", "투자·회수 현금흐름의 내부수익률", "레버리지 전·후, 운용기간·Exit Cap·매각비용 가정 병기"],
    ["Equity Multiple", "총 회수액 ÷ 투입 Equity", "Tranche별로 따로 산출"],
], col_w=[1.8, 4.4, 6.0], row_h=0.42, key_cols=1)

# ---------------------------------------------------------------- 10 Appendix 체크리스트
s = new_slide(10, "Appendix — 작성 점검표", "[배포 전 점검 항목]")
table(s, L, 1.65, BODY_W, [
    ["#", "점검 항목", "확인"],
    ["1", "모든 수치에 출처(자료명·버전·기준일)를 붙였는가", "☐"],
    ["2", "같은 지표가 페이지마다 같은 값인가 (Executive Summary ↔ 본문 ↔ 수익률 분석)", "☐"],
    ["3", "Sources = Uses 합계가 일치하는가", "☐"],
    ["4", "Cap Rate · LTV · DSCR의 분모와 정의를 명시했는가", "☐"],
    ["5", "추정치는 [추정], 미확인 값은 [확인 필요]로 표시했는가", "☐"],
    ["6", "'보장' · '확정' 등 수익을 약속하는 표현을 쓰지 않았는가", "☐"],
    ["7", "일정·조건이 최신 계약서·확약서와 일치하는가", "☐"],
    ["8", "비밀유지 문구와 배포 대상·범위를 표지에 적었는가", "☐"],
    ["9", "단위(억원/백만원, ㎡/평)와 반올림 기준이 일관적인가", "☐"],
    ["10", "리스크 페이지에 완화 방안과 잔여 위험을 함께 적었는가", "☐"],
], col_w=[0.6, 10.4, 1.2], row_h=0.42)

# ---------------------------------------------------------------- 저장 (문서 속성 정리)
cp = prs.core_properties
cp.author = ""
cp.last_modified_by = ""
cp.title = "IM 작성 실습용 양식(Clean)"
cp.subject = "강의용"
cp.keywords = "IM, 실습, 양식, Clean"
cp.comments = ""
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
