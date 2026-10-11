# -*- coding: utf-8 -*-
"""강의용 표준서식(Clean) 공통 빌더.

블록 DSL:
  ("title", 텍스트)        문서 제목(가운데, 굵게)
  ("sub", 텍스트)          부제(가운데)
  ("center", 텍스트)       가운데 정렬 본문
  ("right", 텍스트)        오른쪽 정렬 본문
  ("h", 텍스트)            조 제목(굵게)
  ("h2", 텍스트)           항 제목(굵게, 들여쓰기 없음)
  ("p", 텍스트)            본문
  ("i1"/"i2"/"i3", 텍스트) 들여쓴 항목(단계별)
  ("table", [행...], 폭[cm] 목록 또는 None)  첫 행은 머리글
  ("kv", [(항목, 내용)...]) 2열 표(항목 열 음영)
  ("sign", [(당사자, 상호, 대표)...]) 서명란
  ("pb",)                  쪽 나눔
  ("gap",)                 빈 줄
"""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "맑은 고딕"
BODY_PT = 10
HEADER_TEXT = "강의용 예시 서식(Clean) · 특정 거래·당사자와 무관 · [●]는 공란"
FOOTER_NOTE = "본 서식은 교육 목적의 일반 예시이며 법률 자문이 아닙니다. 실제 거래에는 외부 법무법인 검토가 필요합니다."


def _set_font(run, size=BODY_PT, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), FONT)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _para(doc_or_cell, text, size=BODY_PT, bold=False, align=None, left_cm=0.0,
          hanging_cm=0.0, before=0, after=4, color=None, line=1.3):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if left_cm or hanging_cm:
        pf.left_indent = Cm(left_cm + hanging_cm)
        pf.first_line_indent = Cm(-hanging_cm)
    if align:
        p.alignment = align
    elif "\n" in text:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    _set_font(r, size, bold, color)
    return p


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=9):
    cell.text = ""
    lines = str(text).split("\n")
    p = cell.paragraphs[0]
    for i, ln in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(ln)
        _set_font(r, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(doc, rows, widths=None, header=True, key_col=False):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            val = row[j] if j < len(row) else ""
            c = t.cell(i, j)
            is_head = (header and i == 0) or (key_col and j == 0)
            _cell_text(c, val, bold=is_head,
                       align=WD_ALIGN_PARAGRAPH.CENTER if is_head else WD_ALIGN_PARAGRAPH.LEFT)
            if is_head:
                _shade(c, "D9E2F3" if (header and i == 0) else "F2F2F2")
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
            if j < len(widths):
                gc.set(qn("w:w"), str(int(widths[j] * 567)))
        for j, w in enumerate(widths):
            t.columns[j].width = Cm(w)
            for i in range(len(rows)):
                t.cell(i, j).width = Cm(w)
    # 표 뒤 간격
    _para(doc, "", after=2)
    return t


def _add_page_field(par):
    r = par.add_run()
    _set_font(r, 8)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._element.append(f1); r._element.append(it); r._element.append(f2)


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    st.element.get_or_add_rPr()
    rf = st.element.rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); st.element.rPr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), FONT)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(hp.add_run(HEADER_TEXT), 8, color="7F7F7F")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(fp.add_run(FOOTER_NOTE + "   - "), 7.5, color="7F7F7F")
    _add_page_field(fp)
    _set_font(fp.add_run(" -"), 7.5, color="7F7F7F")
    # 문서 속성 정리(작성자 등 식별정보 제거)
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.company = "" if hasattr(cp, "company") else None
    cp.comments = ""
    cp.keywords = "강의용, 예시서식, Clean"
    return doc


def render(blocks, out_path, title_prop):
    doc = new_doc()
    doc.core_properties.title = title_prop
    for b in blocks:
        k = b[0]
        if k == "title":
            _para(doc, b[1], size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6)
        elif k == "sub":
            _para(doc, b[1], size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
        elif k == "center":
            _para(doc, b[1], align=WD_ALIGN_PARAGRAPH.CENTER)
        elif k == "right":
            _para(doc, b[1], align=WD_ALIGN_PARAGRAPH.RIGHT)
        elif k == "h":
            _para(doc, b[1], size=11, bold=True, before=10, after=4).paragraph_format.keep_with_next = True
        elif k == "h2":
            _para(doc, b[1], bold=True, before=6, after=3).paragraph_format.keep_with_next = True
        elif k == "p":
            _para(doc, b[1])
        elif k == "i1":
            _para(doc, b[1], left_cm=0.3, hanging_cm=0.6)
        elif k == "i2":
            _para(doc, b[1], left_cm=0.9, hanging_cm=0.6, after=2)
        elif k == "i3":
            _para(doc, b[1], left_cm=1.5, hanging_cm=0.6, after=2)
        elif k == "table":
            _table(doc, b[1], b[2] if len(b) > 2 else None)
        elif k == "kv":
            _table(doc, [list(x) for x in b[1]], b[2] if len(b) > 2 else [4.0, 12.4],
                   header=False, key_col=True)
        elif k == "sign":
            rows = [["구분", "상호", "대표자 / 인"]] + [[a, c, d + "   (인)"] for a, c, d in b[1]]
            _table(doc, rows, [3.0, 7.4, 6.0])
        elif k == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif k == "gap":
            _para(doc, "", after=2)
        else:
            raise ValueError(k)
    doc.save(out_path)
    return out_path
