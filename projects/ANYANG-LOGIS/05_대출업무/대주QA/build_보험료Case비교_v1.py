import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
OUT='/home/user/Investment/projects/ANYANG-LOGIS/05_대출업무/대주QA/20260930_QA_패키지보험료_Case비교_v01.xlsx'
wb=openpyxl.Workbook()
F=lambda **k: Font(name='맑은 고딕', **k)
HF=PatternFill('solid',fgColor='345B86'); IN=PatternFill('solid',fgColor='FFF2CC'); OUT_=PatternFill('solid',fgColor='DAE3F3'); TOT=PatternFill('solid',fgColor='D9E1F2'); G=PatternFill('solid',fgColor='F2F2F2')
th=Side(style='thin',color='BFBFBF'); BD=Border(left=th,right=th,top=th,bottom=th)
def c(ws,ref,v,fmt=None,fill=None,b=False,al=None,col=None):
    x=ws[ref]; x.value=v; x.font=F(bold=b,color=col) if col else F(bold=b); x.border=BD
    if fmt: x.number_format=fmt
    if fill: x.fill=fill
    x.alignment=Alignment(horizontal=al or ('right' if isinstance(v,(int,float)) or (isinstance(v,str) and v.startswith('=')) else 'left'),vertical='center',wrap_text=True)
    return x
def hdr(ws,row,vals):
    for i,v in enumerate(vals,1):
        x=ws.cell(row=row,column=i,value=v); x.font=F(bold=True,color='FFFFFF'); x.fill=HF; x.border=BD; x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)

# ── 산정근거 ────────────────────────────────────────────
g=wb.active; g.title='산정근거'
g['A1']='산정 근거 (노란 셀 = 입력값)'; g['A1'].font=F(bold=True,size=13)
hdr(g,3,['구분','항목','값','단위','출처·비고'])
rows=[
 ('기본','연면적',95474.59,'㎡','견적서, 건축물대장'),
 ('기본','평 환산계수',3.3058,'㎡/평',''),
 ('기본','연면적(평)','=C4/C5','평',''),
 ('Case1 요율','Sec I PAR 가입금액',119000000000,'원','첨부 견적서(안양 물류 Package Insurance 예상 보험료)'),
 ('Case1 요율','Sec I PAR 보험료',583100000,'원','동상'),
 ('Case1 요율','Sec I 요율 (역산)','=C8/C7','%','Case2 동일 적용'),
 ('Case1 요율','Sec III BI 가입금액(매출액·임대료)',31020000000,'원','동상, 보상기간 24개월분'),
 ('Case1 요율','Sec III BI 보험료',151998000,'원','동상'),
 ('Case1 요율','Sec III 요율 (역산)','=C11/C10','%','Case2 동일 적용'),
 ('Case1 요율','Sec III 보상기간',24,'개월','동상'),
 ('Case1 요율','월 매출액(임대료) 환산','=C10/C13','원/월','가입금액 ÷ 보상기간'),
 ('재조달원가','창고시설·사무실 등 (지2~8층)',60639.97,'㎡','미래새한 감정서 p.119~120, 재조달원가 1,800,000원/㎡'),
 ('재조달원가','램프·차로·BERTH 등 (지2~8층)',34785.76,'㎡','동상, 1,100,000원/㎡'),
 ('재조달원가','옥탑 계단실 등',103.4,'㎡','동상, 1,100,000원/㎡'),
 ('재조달원가','경비실 등 (1층)',8.82,'㎡','동상, 1,500,000원/㎡'),
 ('재조달원가','창고 재조달원가 단가',1800000,'원/㎡',''),
 ('재조달원가','램프·옥탑 재조달원가 단가',1100000,'원/㎡',''),
 ('재조달원가','경비실 재조달원가 단가',1500000,'원/㎡',''),
 ('재조달원가','건물 재조달원가 합계','=C15*C19+(C16+C17)*C20+C18*C21','원','감가 전 금액. 감정서 건물 적산가액 141,641,521,920원 ÷ 잔가율 96%와 일치'),
 ('재조달원가','(참고) 감정서 건물 적산가액',141641521920,'원','감가수정 후(내용연수 50년, 경과 2년)'),
 ('공사기간','착공신고일','2021-03-26','일자','CBRE 물리실사 §2.1'),
 ('공사기간','사용승인일(준공)','2023-11-24','일자','건축물대장·CBRE'),
 ('공사기간','공사기간(일)','=DATEVALUE(C25)-DATEVALUE(C24)','일',''),
 ('공사기간','Case2 보상기간','=ROUNDUP(C26/(365/12),0)','개월','공사기간 월 환산 후 올림 (973일 ≈ 32개월)'),
 ('모델','기준 재무모델 보험료(1차년도)',742284700,'원','A&R L72: 연면적 기준, 월 평당 2,141.79원, 연 2% 감소 가정'),
]
for i,(a,b_,v,u,s) in enumerate(rows, start=4):
    c(g,f'A{i}',a); c(g,f'B{i}',b_)
    fmt='0.0000%' if u=='%' else ('#,##0.00' if u in('㎡','㎡/평','평') else '#,##0')
    fill=IN if not (isinstance(v,str) and v.startswith('=')) else OUT_
    c(g,f'C{i}',v,fmt,fill); c(g,f'D{i}',u,al='center'); c(g,f'E{i}',s)
g['C24'].number_format='@'; g['C25'].number_format='@'
for i,w in enumerate([12,34,20,8,62],1): g.column_dimensions[L(i)].width=w

# ── 비교표 ──────────────────────────────────────────────
s=wb.create_sheet('비교표',0)
s['A1']='안양물류센터 패키지보험 예상 보험료 — Case 1(기존안) vs Case 2(재조달원가·공사기간 기준)'; s['A1'].font=F(bold=True,size=13)
s['A2']='동일 요율 적용 (Sec I 0.49%, Sec III 0.49% — 기존 견적서 역산). 보험기간 1년, 창고(상온 100%), 단위: 원'
hdr(s,4,['담보구분','항목','Case 1 가입금액·조건','Case 1 보험료','Case 2 가입금액·조건','Case 2 보험료','증감(보험료)','증감률','Case 2 산정 기준'])
R=[
 ('Sec I. 재물손해 (PAR)','보험가입금액 (건물·부속설비)','=산정근거!C7','=산정근거!C8','=산정근거!C22','=ROUND(E5*산정근거!C9,0)','미래새한 감정서 건물 재조달원가, 요율 Case1 동일'),
 ('','자기부담금 (1사고당)','10% of claim, Min. 1억','', '10% of claim, Min. 1억','','Case1 동일'),
 ('Sec II. 기계손해 (MB)','보험가입금액','Included in PAR',0,'Included in PAR',0,'Case1 동일'),
 ('Sec III. 기업휴지 (BI&MLOP)','보험가입금액 (매출액·임대료)','=산정근거!C10','=산정근거!C11','=산정근거!C14*E9','=ROUND(E8*산정근거!C12,0)','월 임대료 × 보상기간, 요율 Case1 동일'),
 ('','보상기간 (개월)','=산정근거!C13','','=산정근거!C27','','준공 당시 공사기간(2021.3.26~2023.11.24)'),
 ('','공제기간','7일','','7일','','Case1 동일'),
 ('Sec IV. 배상책임 (GL)','시설소유자 50억 / 임차인재고자산 30억','5,000,000,000 / 3,000,000,000',7000000,'5,000,000,000 / 3,000,000,000',7000000,'Case1 동일'),
 ('승강기배상책임','법정 보상한도','법정',146700,'법정',146700,'Case1 동일'),
 ('전기차충전시설배상책임','법정 보상한도','법정',20000,'법정',20000,'Case1 동일'),
 ('가스배상책임','법정 보상한도','법정',20000,'법정',20000,'Case1 동일'),
]
for r,(a,b_,c1,p1,c2,p2,note) in enumerate(R, start=5):
    c(s,f'A{r}',a,b=True); c(s,f'B{r}',b_)
    c(s,f'C{r}',c1,'#,##0'); c(s,f'D{r}',p1 if p1!='' else None,'#,##0')
    c(s,f'E{r}',c2,'#,##0',fill=OUT_ if isinstance(c2,str) and c2.startswith('=') else None); c(s,f'F{r}',p2 if p2!='' else None,'#,##0',fill=OUT_ if isinstance(p2,str) and p2.startswith('=') else None)
    if p1!='' : c(s,f'G{r}',f'=F{r}-D{r}','#,##0;[Red]-#,##0'); c(s,f'H{r}',f'=IF(D{r}=0,"-",F{r}/D{r}-1)','0.0%')
    else: c(s,f'G{r}',None); c(s,f'H{r}',None)
    c(s,f'I{r}',note)
t=15
c(s,f'A{t}','합계',b=True,fill=TOT); c(s,f'B{t}','',fill=TOT); c(s,f'C{t}','',fill=TOT); c(s,f'E{t}','',fill=TOT)
c(s,f'D{t}','=SUM(D5:D14)','#,##0',TOT,True); c(s,f'F{t}','=SUM(F5:F14)','#,##0',TOT,True)
c(s,f'G{t}',f'=F{t}-D{t}','#,##0',TOT,True); c(s,f'H{t}',f'=F{t}/D{t}-1','0.0%',TOT,True); c(s,f'I{t}','',fill=TOT)
# 평당 보험료
hdr(s,17,['구분','항목','Case 1','','Case 2','','증감','증감률','비고'])
P=[('평당 보험료','연간 (원/평·년)','=D15/산정근거!$C$6','=F15/산정근거!$C$6','연면적 28,881평 기준'),
   ('','월간 (원/평·월)','=C18/12','=E18/12','기준 재무모델 A&R L72 단위(월 평당)와 동일'),
   ('','연간 (원/㎡·년)','=D15/산정근거!$C$4','=F15/산정근거!$C$4','연면적 95,474.59㎡'),
   ('부보율','PAR 가입금액 ÷ 건물 재조달원가','=C5/산정근거!$C$22','=E5/산정근거!$C$22','Case1은 재조달원가의 약 81% → 비례보상 적용 가능성'),
   ('BI','보상기간 (개월)','=C9','=E9',''),
   ('재무모델 영향','1차년도 보험료 − 모델 반영액','=D15-산정근거!$C$28','=F15-산정근거!$C$28','모델(742.28백만) = Case1')]
for r,(a,b_,v1,v2,note) in enumerate(P, start=18):
    c(s,f'A{r}',a,b=True); c(s,f'B{r}',b_)
    fmt='0.0%' if a=='부보율' else '#,##0'
    c(s,f'C{r}',v1,fmt); c(s,f'D{r}',None); c(s,f'E{r}',v2,fmt,OUT_); c(s,f'F{r}',None)
    if a=='부보율': c(s,f'G{r}',f'=E{r}-C{r}','+0.0%p;-0.0%p'); c(s,f'H{r}',None)
    else: c(s,f'G{r}',f'=E{r}-C{r}','#,##0;[Red]-#,##0'); c(s,f'H{r}',f'=IF(C{r}=0,"-",E{r}/C{r}-1)','0.0%')
    c(s,f'I{r}',note)
s['A25']='※ Case 2는 보험사 견적이 아닌 동일 요율 가정의 추정치이며, 실제 요율은 가입금액 증가·보상기간 연장에 따라 보험사 인수심사로 달라질 수 있음. 자기부담금·배상책임 한도는 Case 1과 동일 가정.'
s['A26']='※ 출처: 기존 견적서(첨부 이미지), 미래새한 감정평가서(2026.09) 건물 재조달원가, CBRE 물리실사(착공 2021.3.26·사용승인 2023.11.24), 기준 재무모델 v02 A&R L72.'
for a in ('A25','A26'): s[a].font=F(size=9,color='404040')
for i,w in enumerate([26,32,24,15,24,15,15,9,40],1): s.column_dimensions[L(i)].width=w
for r in range(5,24): s.row_dimensions[r].height=30
s.freeze_panes='A5'
wb.calculation.fullCalcOnLoad=True
wb.save(OUT); print('saved',OUT)
