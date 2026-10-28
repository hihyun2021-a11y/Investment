# -*- coding: utf-8 -*-
"""기표일(2026-10-29) 기관별 수수료·보수·이자유보금 이체내역 — 운용역 제공 이체 양식(순번~입금계좌) 기준.

근거
- Equity 인수수수료: 04_Equity투자자/투자자별/[채결본] 인수계약서·주선 및 인수계약서(2026-10-26) 제5조, 각 계좌사본
- 대출관련수수료: [체결본] 수수료약정서 10건(2026-10-27) 제2조(우리은행 주선은 제3조), 대주법인서류 통장사본
- 담보신탁보수: 부동산담보신탁계약서(2026-10-20 v02.3) 별지3 신탁재산 관리보수, 우리은행 신탁부 계좌표제부
- 이자유보금: 대출약정서(v.final) 이자적립금 조항(인출일 당일 3개월분 이자), CD금리 10/26 3.26%·10/27 3.27%·10/28 3.27%
출력: 20261028_금융_기표일이체내역_수수료보수이자유보금_v02.xlsx
v02(2026-10-28): 삼성증권 인수수수료를 면세(금융용역) 기준 비과세로 반영, 1~4탭을 취합한 '0.통합리스트' 탭을 맨 앞에 추가
"""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "20261028_금융_기표일이체내역_수수료보수이자유보금_v02.xlsx")

FONT = "맑은 고딕"
HDR = PatternFill("solid", fgColor="A9D08E")
BODY = PatternFill("solid", fgColor="E2EFDA")
INPUT = PatternFill("solid", fgColor="FFF2CC")
TOT = PatternFill("solid", fgColor="D9D9D9")
thin = Side(style="thin", color="000000")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)
L = Alignment(horizontal="left", vertical="center", wrap_text=True)
R = Alignment(horizontal="right", vertical="center")
NUM = "#,##0"

OUT_ACC = ("운영계좌", "우리은행", "1005-804-963776")  # 양식 기본값(출금처)

HEAD = ["순번", "항목", "계정", "금액", "과세유형", "출금처", "출금은행", "출금계좌", "입금처", "입금은행", "입금계좌",
        "부가세(10%)", "실제 이체금액", "지급기한", "근거·비고"]
WIDTH = [6, 34, 10, 17, 10, 11, 10, 17, 30, 10, 20, 14, 17, 22, 60]


def f(bold=False, size=10, color="000000"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def list_sheet(wb, title, caption, rows, n_blank=3):
    """rows: dict(item, amt, tax, payee, bank, acc, due, note[, amt_formula])"""
    ws = wb.create_sheet(title)
    ws["A1"] = "필수입력"
    ws["A1"].font = f(color="808080")
    ws["B1"] = caption
    ws["B1"].font = f(bold=True, size=11)
    for j, (h, w) in enumerate(zip(HEAD, WIDTH), 1):
        c = ws.cell(row=2, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = f(bold=True), HDR, C, BOX
        ws.column_dimensions[get_column_letter(j)].width = w
    r = 3
    for i, d in enumerate(rows, 1):
        vals = [i, d["item"], "", d.get("amt_formula", d["amt"]), d["tax"], *OUT_ACC, d["payee"], d["bank"], d["acc"],
                f'=IF(E{r}="과세",ROUNDDOWN(D{r}*0.1,0),0)', f"=D{r}+L{r}", d["due"], d["note"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.fill, c.border = f(), BODY, BOX
            c.alignment = R if j in (4, 12, 13) else (L if j in (2, 9, 14, 15) else C)
            if j in (4, 12, 13):
                c.number_format = NUM
        if d.get("flag"):
            for j in (5, 15):
                ws.cell(row=r, column=j).fill = INPUT
        r += 1
    for _ in range(n_blank):  # 양식처럼 추가 입력용 빈 행(출금처 기본값만)
        for j, v in enumerate([r - 2, "", "", None, "", *OUT_ACC, "", "", "", None, None, "", ""], 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.fill, c.border, c.alignment = f(), BODY, BOX, C
        r += 1
    last = r - 1
    ws.cell(row=r, column=2, value="합계")
    ws.cell(row=r, column=4, value=f"=SUM(D3:D{last})")
    ws.cell(row=r, column=12, value=f"=SUM(L3:L{last})")
    ws.cell(row=r, column=13, value=f"=SUM(M3:M{last})")
    for j in range(1, len(HEAD) + 1):
        c = ws.cell(row=r, column=j)
        c.font, c.fill, c.border = f(bold=True), TOT, BOX
        c.alignment = R if j in (4, 12, 13) else C
        if j in (4, 12, 13):
            c.number_format = NUM
    ws.freeze_panes = "C3"
    ws.sheet_view.zoomScale = 90
    return ws, r


wb = Workbook()
wb.remove(wb.active)

# ---------------------------------------------------------------- 1. Equity 인수수수료
EQ = "04_Equity투자자/투자자별"
eq_rows = [
    dict(item="인수수수료(제1종 종류주식, 주선 및 잔액인수)", amt=480_000_000, tax="과세",
         payee="주식회사 우리투자증권", bank="우리은행", acc="1003-700-005935",
         due="거래종결일로부터 1영업일 이내",
         note=f"인수약정금 240억원 × 2% (부가가치세 별도). [채결본] 인수계약서(제1종 종류주식)_우리투자증권_20261026 제5조 제1항 / 계좌: 주식회사 우리투자증권_우리은행 계좌 사본"),
    dict(item="인수수수료(제1종 종류주식)", amt=36_000_000, tax="과세",
         payee="엠지캐피탈(주)", bank="우리은행", acc="894-000383-01-015",
         due="거래종결일로부터 5영업일 이내",
         note="인수약정금 30억원 × 1.2% (부가가치세 별도). [채결본] 인수계약서(제1종 종류주식)_엠지캐피탈_20261026 제5조 / 계좌: [엠지캐피탈] 우리은행 모계좌 (계약서 기재 계좌와 동일)"),
    dict(item="인수수수료(삼성증권, 제2종 종류주식 주선 및 잔액인수)", amt=1_200_000_000, tax="비과세",
         payee="삼성증권(주)", bank="우리은행", acc="061-219568-03-103",
         due="거래종결일로부터 5영업일 이내",
         note="인수약정금 500억원 × 2.4%. [채결본] 주선 및 인수계약서_삼성증권_20261026 제5조(부가가치세 문구 없음). 면세 기준: 증권사의 주선(투자중개업)·잔액인수(투자매매업) 용역 → 금융용역 면세(부가가치세법 제26조 제1항 제11호, 시행령 제40조) [내부 검토, 삼성증권 계산서 발행 여부 확인] / 계좌: [삼성증권] 계좌사본_우리은행"),
]
list_sheet(wb, "1.Equity인수수수료", "Equity 인수수수료 이체내역 (종류주식 인수계약 체결본 2026-10-26 기준)", eq_rows)

# ---------------------------------------------------------------- 2. 대출관련수수료
D = "대주법인서류"
SH = "[체결본] 대출참여수수료약정서(Tr_A-1_신협대주단)_20261027 제2조 표(대출약정금의 1.0%)"
loan_rows = [
    ("대출참여수수료(Tr.A-1 롯데캐피탈)", 540_000_000, "롯데캐피탈(주)", "우리은행", "364-032300-03-001",
     "300억 × 1.8%. [체결본] 대출참여수수료약정서(Tr_A-1_롯데캐피탈) 제2조 / 계좌: ■롯데캐피탈 입금계좌 사본"),
    ("대출참여수수료(Tr.A-1 우리카드)", 120_000_000, "(주)우리카드", "우리은행", "1005-504-287838",
     "100억 × 1.2%. [체결본] 대출참여수수료약정서(Tr_A-1_우리카드) 제2조 / 계좌: [우리카드] 수수료 및 원리금 상환 통장표지"),
    ("대출참여수수료(Tr.A-1 신용협동조합중앙회)", 1_000_000, "신협중앙회", "농협은행", "081-01-168860",
     SH + " / 계좌: 대주 통지처(261023 vf) 기재 수취계좌 — 통장사본 미접수", True),
    ("대출참여수수료(Tr.A-1 서로신협)", 40_000_000, "서로신용협동조합", "신협", "131-000-941666", SH + " / 계좌: 서로 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 수지신협)", 28_000_000, "수지신협", "신협", "131-014-493393", SH + " / 계좌: 수지 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 정선신협)", 24_000_000, "정선신용협동조합", "신협", "131-019-058167", SH + " / 계좌: 정선 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 연희신협)", 14_000_000, "연희신협", "신협", "131-000-189579", SH + " / 계좌: 연희 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 금모래신협)", 18_000_000, "금모래신협", "신협", "131-000-001170", SH + " / 계좌: 금모래 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 남인천신협)", 10_000_000, "남인천신협", "신협", "131-000-942952", SH + " / 계좌: 남인천 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 미호신협)", 13_000_000, "미호신협", "신협", "131-002-821894", SH + " / 계좌: 미호신협 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 대구대서신협)", 13_000_000, "대구대서신협", "신협", "131-005-005290", SH + " / 계좌: 대구대서 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 한빛신협)", 17_000_000, "한빛신용협동조합", "신협", "131-003-657209", SH + " / 계좌: 한빛 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 양동행복신협)", 19_000_000, "양동행복신협", "신협", "131-021-819738", SH + " / 계좌: 양동행복 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 서귀포동부신협)", 12_000_000, "서귀포동부신협", "신협", "131-005-410020", SH + " / 계좌: 서귀포동부 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 광주어룡신협)", 27_000_000, "광주어룡신협", "신협", "131-004-571106", SH + " / 계좌: 광주어룡 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 광주미래신협)", 14_000_000, "광주미래신용협동조합", "신협", "131-004-612864", SH + " / 계좌: 광주미래 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 새안양신협)", 40_000_000, "새안양신용협동조합", "신협", "131-001-041701", SH + " / 계좌: 새안양 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 남수원신협)", 3_000_000, "남수원신용협동조합", "신협", "131-001-156258", SH + " / 계좌: 남수원 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 북인천신협)", 7_000_000, "북인천신용협동조합", "신협", "131-002-051950", SH + " / 계좌: 북인천신협 사업자등록증,계좌"),
    ("대출참여수수료(Tr.A-1 수협은행)", 100_000_000, "수협은행사당역지점", "수협은행", "1010-1159-8651",
     "100억 × 1.0%. [체결본] 대출참여수수료약정서(Tr_A-1_수협은행) 제2조 / 계좌: 수협은행 업무용통장 사본(사당역지점)"),
    ("대출주선수수료(Tr.A-1 에스비아이저축은행)", 100_000_000, "(주)에스비아이저축은행", "우리은행", "1005-704-857219",
     "정액 1억원. [체결본] 대출주선수수료약정서(Tr_A_1_SBI저축은행) 제2조 / 계좌: [SBI기업금융1팀] 예치금계좌_우리은행_통장사본"),
    ("대출주선수수료(Tr.A-2 우리은행)", 830_000_000, "(주)우리은행 투자금융1부", "우리은행", "1005-001-222719",
     "정액 8.3억원. [체결본] 대출주선수수료약정서(Tr_A_2_우리은행) 제3조 / 계좌: 【우리은행】통장사본_투자금융부_222719"),
    ("대출참여수수료(Tr.A-2 우리은행)", 560_000_000, "(주)우리은행 투자금융1부", "우리은행", "1005-001-222719",
     "700억 × 0.8%. [체결본] 대출참여수수료약정서(Tr_A_2_우리은행) 제2조 / 계좌: 【우리은행】통장사본_투자금융부_222719 (주선수수료와 같은 계좌로 기재 — 구분 입금 여부 우리은행 확인)"),
    ("대출취급수수료(Tr.B 키움캐피탈)", 200_000_000, "키움캐피탈 주식회사", "우리은행", "1005-003-573134",
     "200억 × 1.0%. [체결본] 대출취급수수료약정서(Tr_B_키움캐피탈) 제2조 / 계좌: 키움캐피탈 기업금융2팀 통장사본"),
    ("대출취급수수료(Tr.B 엠지캐피탈)", 64_000_000, "엠지캐피탈(주)", "우리은행", "894-000383-01-015",
     "100억 × 0.64%. [체결본] 대출취급수수료약정서(Tr_B_MG캐피탈) 제2조 / 계좌: [엠지캐피탈] 우리은행 모계좌"),
    ("대출참여수수료(Tr.B 디비저축은행)", 80_000_000, "(주)디비저축은행", "국민은행", "006-01-0381-794",
     "80억 × 1.0%. [체결본] 대출참여수수료약정서(Tr_B_DB저축은행) 제2조 / 계좌: [DB저축은행]_국민은행 예치금계좌"),
]
rows = []
for t in loan_rows:
    item, amt, payee, bank, acc, note = t[:6]
    rows.append(dict(item=item, amt=amt, tax="비과세", payee=payee, bank=bank, acc=acc, due="인출일(2026-10-29)",
                     note="부가가치세 없음. " + note, flag=len(t) > 6))
assert sum(r["amt"] for r in rows) == 2_894_000_000
list_sheet(wb, "2.대출관련수수료", "대출관련수수료 이체내역 (수수료약정서 체결본 2026-10-27 기준, 25개 대주·26건)", rows)

# ---------------------------------------------------------------- 3. 담보신탁보수
trust = [dict(item="담보신탁보수(신탁재산 관리보수)", amt=50_000_000, tax="비과세", payee="(주)우리은행 신탁부",
              bank="우리은행", acc="003-441153-01-001", due="수익권증서 발급 시(일시 선납)",
              note="금 5천만원, 부가가치세 없음. 부동산담보신탁계약서(2026-10-20 v02.3) 제18조·별지3 / 계좌: □신탁보수입금계좌표제부_우리은행신탁부. ※ 신탁원부 변경등기 시 관리보수 150만원 별도(특약 제16조) — 이번 이체 대상 아님")]
list_sheet(wb, "3.담보신탁보수", "담보신탁보수 이체내역 (수탁자: 우리은행)", trust)

# ---------------------------------------------------------------- 이자유보금 계산
LENDERS = [  # (Tranche, 대주, 약정액)
    ("A-1", "롯데캐피탈(주)", 30_000_000_000), ("A-1", "(주)우리카드", 10_000_000_000), ("A-1", "신용협동조합중앙회", 100_000_000),
    ("A-1", "서로신용협동조합", 4_000_000_000), ("A-1", "수지신용협동조합", 2_800_000_000), ("A-1", "정선신용협동조합", 2_400_000_000),
    ("A-1", "연희신용협동조합", 1_400_000_000), ("A-1", "금모래신용협동조합", 1_800_000_000), ("A-1", "남인천신용협동조합", 1_000_000_000),
    ("A-1", "미호신용협동조합", 1_300_000_000), ("A-1", "대구대서신용협동조합", 1_300_000_000), ("A-1", "한빛신용협동조합", 1_700_000_000),
    ("A-1", "양동행복신용협동조합", 1_900_000_000), ("A-1", "서귀포동부신용협동조합", 1_200_000_000), ("A-1", "광주어룡신용협동조합", 2_700_000_000),
    ("A-1", "광주미래신용협동조합", 1_400_000_000), ("A-1", "새안양신용협동조합", 4_000_000_000), ("A-1", "남수원신용협동조합", 300_000_000),
    ("A-1", "북인천신용협동조합", 700_000_000), ("A-1", "수협은행", 10_000_000_000), ("A-1", "(주)에스비아이저축은행", 10_000_000_000),
    ("A-2", "(주)우리은행", 70_000_000_000),
    ("B", "키움캐피탈(주)", 20_000_000_000), ("B", "엠지캐피탈(주)", 10_000_000_000), ("B", "(주)디비저축은행", 8_000_000_000),
]
assert sum(x[2] for x in LENDERS) == 198_000_000_000

ws = wb.create_sheet("이자유보금_계산")
ws.column_dimensions["A"].width = 30
for col, w in zip("BCDEFGH", (24, 18, 12, 12, 12, 18, 40)):
    ws.column_dimensions[col].width = w
ws["A1"] = "이자유보계좌 이자적립금(3개월분) 계산 — 인출일 2026-10-29"
ws["A1"].font = f(bold=True, size=12)
ws["A2"] = ("대출약정서(v.final): 차주는 인출일 당일 Tranche A·B 대출금의 각 3개월분 이자 상당액을 각 이자유보계좌에 적립. "
            "이자 = 원금 × 연이율 × 일수 ÷ 365, 초일 산입·말일 불산입, 대주별 원 미만 절사.")
ws["A2"].font = f(size=9)
ws.merge_cells("A2:H2")
ws.row_dimensions[2].height = 30
ws["A2"].alignment = L

inp = [
    ("인출일(최초 이자기간 시작일)", dt.date(2026, 10, 29), "기준 일정"),
    ("최초 이자지급일(3개월 응당일)", dt.date(2027, 1, 29), "2027-01-29(금) 영업일 — 조정 없음"),
    ("이자기간 일수", "=B5-B4", "초일 산입·말일 불산입"),
    ("연 일수", 365, "2026~2027년 평년"),
    ("CD금리 ① 2026-10-26(월)", 0.0326, "운용역 제공"),
    ("CD금리 ② 2026-10-27(화)", 0.0327, "운용역 제공"),
    ("CD금리 ③ 2026-10-28(수)", 0.0327, "운용역 제공"),
    ("CD금리 3영업일 단순평균", "=AVERAGE(B8:B10)", "반올림 전"),
    ("A-2 기준금리", "=ROUND(B11,4)", "소수점 아래 셋째 자리에서 반올림(%) → 3.27%"),
    ("A-2 가산금리", 0.0143, "대출약정서 제5조"),
    ("Tranche A-1 금리(고정)", 0.05, "대출약정서 제5조"),
    ("Tranche A-2 적용금리", "=B12+B13", "기준금리 + 1.43%"),
    ("Tranche B 금리(고정)", 0.06, "대출약정서 제5조"),
]
ws["A3"], ws["B3"], ws["C3"] = "항목", "값", "비고"
for j in range(1, 4):
    c = ws.cell(row=3, column=j)
    c.font, c.fill, c.border, c.alignment = f(bold=True), HDR, BOX, C
for i, (k, v, n) in enumerate(inp, 4):
    ws.cell(row=i, column=1, value=k)
    ws.cell(row=i, column=2, value=v)
    ws.cell(row=i, column=3, value=n)
    for j in range(1, 4):
        c = ws.cell(row=i, column=j)
        c.font, c.border = f(), BOX
        c.alignment = L if j != 2 else R
    b = ws.cell(row=i, column=2)
    if isinstance(v, dt.date) or i in (4, 5):
        b.number_format = "yyyy-mm-dd"
    elif i in (6, 7):
        b.number_format = "0"
    else:
        b.number_format = "0.0000%"
    if i in (8, 9, 10):
        b.fill = INPUT
ws.merge_cells("C3:E3")
for i in range(4, 17):
    ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=5)

hr = 19
heads = ["대주", "Tranche", "대출약정금", "적용금리", "일수", "연일수", "3개월분 이자(원 미만 절사)", "입금 이자유보계좌"]
for j, h in enumerate(heads, 1):
    c = ws.cell(row=hr, column=j, value=h)
    c.font, c.fill, c.border, c.alignment = f(bold=True), HDR, BOX, C
r = hr + 1
first = r
for tr, name, amt in LENDERS:
    rate = {"A-1": "$B$14", "A-2": "$B$15", "B": "$B$16"}[tr]
    acc = "Tranche B 이자유보계좌 1005-504-970089" if tr == "B" else "Tranche A 이자유보계좌 1005-504-965088"
    vals = [name, tr, amt, f"={rate}", "=$B$6", "=$B$7", f"=ROUNDDOWN(C{r}*D{r}*E{r}/F{r},0)", acc]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border = f(), BOX
        c.alignment = R if j in (3, 4, 5, 6, 7) else (C if j == 2 else L)
    ws.cell(row=r, column=3).number_format = NUM
    ws.cell(row=r, column=4).number_format = "0.00%"
    ws.cell(row=r, column=7).number_format = NUM
    r += 1
last = r - 1
sums = {}
for label, cond in (("Tranche A-1 소계", "A-1"), ("Tranche A-2 소계", "A-2"), ("Tranche B 소계", "B")):
    ws.cell(row=r, column=1, value=label)
    ws.cell(row=r, column=3, value=f'=SUMIF($B${first}:$B${last},"{cond}",$C${first}:$C${last})')
    ws.cell(row=r, column=7, value=f'=SUMIF($B${first}:$B${last},"{cond}",$G${first}:$G${last})')
    sums[cond] = r
    r += 1
ws.cell(row=r, column=1, value="Tranche A 이자적립금 (A-1 + A-2) → 1005-504-965088")
ws.cell(row=r, column=3, value=f"=C{sums['A-1']}+C{sums['A-2']}")
ws.cell(row=r, column=7, value=f"=G{sums['A-1']}+G{sums['A-2']}")
rowA = r
r += 1
ws.cell(row=r, column=1, value="Tranche B 이자적립금 → 1005-504-970089")
ws.cell(row=r, column=3, value=f"=C{sums['B']}")
ws.cell(row=r, column=7, value=f"=G{sums['B']}")
r += 1
ws.cell(row=r, column=1, value="이자적립금 합계")
ws.cell(row=r, column=3, value=f"=C{rowA}+C{rowA + 1}")
ws.cell(row=r, column=7, value=f"=G{rowA}+G{rowA + 1}")
for rr in range(last + 1, r + 1):
    for j in range(1, 9):
        c = ws.cell(row=rr, column=j)
        c.font, c.fill, c.border = f(bold=True), TOT, BOX
        c.alignment = R if j in (3, 7) else L
    ws.cell(row=rr, column=3).number_format = NUM
    ws.cell(row=rr, column=7).number_format = NUM
r += 2
ws.cell(row=r, column=1, value="참고: 연이자 × 3/12 방식").font = f(bold=True)
for k, (label, cond) in enumerate((("Tranche A-1", "A-1"), ("Tranche A-2", "A-2"), ("Tranche B", "B")), 1):
    ws.cell(row=r + k, column=1, value=label).font = f()
    rate = {"A-1": "$B$14", "A-2": "$B$15", "B": "$B$16"}[cond]
    c = ws.cell(row=r + k, column=7, value=f"=ROUNDDOWN(C{sums[cond]}*{rate}*3/12,0)")
    c.font, c.number_format = f(), NUM
r += 5
for t in ["※ 적립액은 최초 이자기간(2026-10-29~2027-01-29, 92일) 실제 일수로 대주별 이자를 계산·절사해 합산한 금액이며, 첫 이자지급일에 실제로 지급할 이자와 같습니다(기존 이자적립금 산정 엑셀 v02와 같은 방식).",
          "※ '3개월분 이자 상당액'을 연이자×3/12로 볼 경우 위 참고 금액이 됩니다. 산정 방식은 대주(우리은행)와 최종 확인하세요 [확인 필요].",
          "※ Tranche A 이자유보계좌는 A-1과 A-2 이자적립금을 함께 받습니다."]:
    ws.cell(row=r, column=1, value=t).font = f(size=9)
    r += 1
ws.freeze_panes = "A4"

# ---------------------------------------------------------------- 4. 이자유보금 예치
CALC = "이자유보금_계산"
ir = [
    dict(item="이자적립금 예치(Tranche A-1, 3개월분)", amt=None, amt_formula=f"='{CALC}'!G{sums['A-1']}", tax="비과세",
         payee="Tranche A 이자유보계좌(차주 명의)", bank="우리은행", acc="1005-504-965088", due="인출일 당일(2026-10-29)",
         note="연 5.00%, 2026-10-29~2027-01-29(92일). 계산: '이자유보금_계산' 탭"),
    dict(item="이자적립금 예치(Tranche A-2, 3개월분)", amt=None, amt_formula=f"='{CALC}'!G{sums['A-2']}", tax="비과세",
         payee="Tranche A 이자유보계좌(차주 명의)", bank="우리은행", acc="1005-504-965088", due="인출일 당일(2026-10-29)",
         note="CD 3.27%(10/26 3.26%·10/27 3.27%·10/28 3.27% 평균, 셋째 자리 반올림) + 1.43% = 연 4.70%, 92일. 계산: '이자유보금_계산' 탭"),
    dict(item="이자적립금 예치(Tranche B, 3개월분)", amt=None, amt_formula=f"='{CALC}'!G{sums['B']}", tax="비과세",
         payee="Tranche B 이자유보계좌(차주 명의)", bank="우리은행", acc="1005-504-970089", due="인출일 당일(2026-10-29)",
         note="연 6.00%, 92일. 계산: '이자유보금_계산' 탭"),
]
for d in ir:
    d["note"] += " / 대출약정서상 대출금입금계좌(1005-504-949583)에서 직접 이체할 수 있음 — 출금계좌 확인"
list_sheet(wb, "4.이자유보금예치", "이자유보계좌 이자적립금 예치 내역 (대출약정서 v.final, 인출일 당일 3개월분 이자)", ir)
wb.move_sheet(CALC, offset=1)

# ---------------------------------------------------------------- 요약
ws = wb.create_sheet("요약", 0)
ws["A1"] = "기표일(2026-10-29) 수수료·보수·이자유보금 이체내역 요약"
ws["A1"].font = f(bold=True, size=12)
heads = ["구분", "건수", "금액(공급가액)", "부가세", "실제 이체금액", "비고"]
for j, (h, w) in enumerate(zip(heads, (24, 8, 18, 15, 18, 60)), 1):
    c = ws.cell(row=3, column=j, value=h)
    c.font, c.fill, c.border, c.alignment = f(bold=True), HDR, BOX, C
    ws.column_dimensions[get_column_letter(j)].width = w
items = [("1.Equity인수수수료", 3, "우리투자증권·엠지캐피탈 과세(계약서 '부가가치세 별도'), 삼성증권 면세(비과세) 기준"),
         ("2.대출관련수수료", 26, "25개 대주(우리은행 주선·참여 2건), 전부 부가가치세 없음"),
         ("3.담보신탁보수", 1, "신탁재산 관리보수, 부가가치세 없음"),
         ("4.이자유보금예치", 3, "Tranche A(A-1+A-2)·B 이자유보계좌")]
r = 4
for sn, n, note in items:
    tot_row = {"1.Equity인수수수료": 3 + 3 + 3, "2.대출관련수수료": 3 + 26 + 3, "3.담보신탁보수": 3 + 1 + 3, "4.이자유보금예치": 3 + 3 + 3}[sn]
    vals = [sn, n, f"='{sn}'!D{tot_row}", f"='{sn}'!L{tot_row}", f"='{sn}'!M{tot_row}", note]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border = f(), BOX
        c.alignment = R if j in (3, 4, 5) else (C if j == 2 else L)
        if j in (3, 4, 5):
            c.number_format = NUM
    r += 1
ws.cell(row=r, column=1, value="합계")
for j, col in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
    ws.cell(row=r, column=j, value=f"=SUM({col}4:{col}{r - 1})")
for j in range(1, 7):
    c = ws.cell(row=r, column=j)
    c.font, c.fill, c.border = f(bold=True), TOT, BOX
    c.alignment = R if j in (3, 4, 5) else C
    if j in (3, 4, 5):
        c.number_format = NUM
r += 2
for t in ["· 각 탭 A~K열은 운용역 제공 이체 양식(순번·항목·계정·금액·과세유형·출금처·출금은행·출금계좌·입금처·입금은행·입금계좌)과 같고, L~O열(부가세·실제 이체금액·지급기한·근거)은 확인용으로 덧붙였습니다.",
          "· '금액'은 양식 예시(인수수수료 36,000,000 과세)와 같이 공급가액 기준이며, 과세 건은 부가세 10%를 더한 '실제 이체금액'을 함께 표시합니다.",
          "· 출금처는 양식 기본값(운영계좌 우리은행 1005-804-963776)으로 채웠습니다. 계정 열은 비워 두었습니다.",
          "· 노란 칸은 확인이 필요한 항목입니다(신협중앙회 통장사본 미접수). 삼성증권 인수수수료는 면세(비과세) 기준으로 반영했습니다.",
          "· '0.통합리스트' 탭은 1~4탭 데이터 행을 수식으로 그대로 가져오므로, 각 탭을 고치면 자동 반영됩니다(탭에 행을 추가하면 통합리스트에도 행 추가 필요).",
          "· 수익 확약이 아닌 이체 준비용 집계표이며, 금액·계좌는 체결본 약정서와 통장사본을 기준으로 합니다."]:
    ws.cell(row=r, column=1, value=t).font = f(size=9)
    r += 1

# ---------------------------------------------------------------- 0. 통합리스트 (1~4탭 취합, 맨 앞)
ws = wb.create_sheet("0.통합리스트", 0)
ws["A1"] = "필수입력"
ws["A1"].font = f(color="808080")
ws["B1"] = "기표일(2026-10-29) 이체 통합 리스트 — 1.Equity인수수수료 · 2.대출관련수수료 · 3.담보신탁보수 · 4.이자유보금예치 취합"
ws["B1"].font = f(bold=True, size=11)
LH = HEAD[:14] + ["구분"]
LW = WIDTH[:14] + [20]
for j, (h, w) in enumerate(zip(LH, LW), 1):
    c = ws.cell(row=2, column=j, value=h)
    c.font, c.fill, c.border, c.alignment = f(bold=True), HDR, BOX, C
    ws.column_dimensions[get_column_letter(j)].width = w
SRC = [("1.Equity인수수수료", 3, "Equity 인수수수료"), ("2.대출관련수수료", 26, "대출관련수수료"),
       ("3.담보신탁보수", 1, "담보신탁보수"), ("4.이자유보금예치", 3, "이자유보금 예치")]
r = 3
for sn, n, label in SRC:
    for k in range(n):
        src_row = 3 + k
        for j in range(1, 15):
            col = get_column_letter(j)
            v = r - 2 if j == 1 else (f"=IF('{sn}'!{col}{src_row}=\"\",\"\",'{sn}'!{col}{src_row})" if j == 3 else f"='{sn}'!{col}{src_row}")
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.fill, c.border = f(), BODY, BOX
            c.alignment = R if j in (4, 12, 13) else (L if j in (2, 9, 14) else C)
            if j in (4, 12, 13):
                c.number_format = NUM
        c = ws.cell(row=r, column=15, value=label)
        c.font, c.fill, c.border, c.alignment = f(), BODY, BOX, C
        if sn == "2.대출관련수수료" and k == 2:  # 신협중앙회 통장사본 미접수
            ws.cell(row=r, column=11).fill = INPUT
        r += 1
last = r - 1
ws.cell(row=r, column=2, value="합계")
for j, col in ((4, "D"), (12, "L"), (13, "M")):
    ws.cell(row=r, column=j, value=f"=SUM({col}3:{col}{last})")
for j in range(1, 16):
    c = ws.cell(row=r, column=j)
    c.font, c.fill, c.border = f(bold=True), TOT, BOX
    c.alignment = R if j in (4, 12, 13) else C
    if j in (4, 12, 13):
        c.number_format = NUM
r += 1
for label, tag in (("  └ Equity 인수수수료", "Equity 인수수수료"), ("  └ 대출관련수수료", "대출관련수수료"),
                   ("  └ 담보신탁보수", "담보신탁보수"), ("  └ 이자유보금 예치", "이자유보금 예치")):
    ws.cell(row=r, column=2, value=label)
    for j, col in ((4, "D"), (12, "L"), (13, "M")):
        ws.cell(row=r, column=j, value=f'=SUMIF($O$3:$O${last},"{tag}",{col}$3:{col}${last})')
    for j in range(1, 16):
        c = ws.cell(row=r, column=j)
        c.font, c.border = f(), BOX
        c.alignment = R if j in (4, 12, 13) else L
        if j in (4, 12, 13):
            c.number_format = NUM
    r += 1
r += 1
for t in ["· 1~4탭 데이터 행을 수식으로 가져옵니다. 금액·계좌는 각 탭에서 고치면 여기도 바뀝니다.",
          "· '금액'은 공급가액, '실제 이체금액'은 과세 건 부가세 10% 포함 금액입니다. 삼성증권 인수수수료는 면세(비과세) 기준입니다.",
          "· 근거(약정서 조항·통장사본)는 각 탭 O열을 참고하세요."]:
    ws.cell(row=r, column=1, value=t).font = f(size=9)
    r += 1
ws.freeze_panes = "C3"
ws.sheet_view.zoomScale = 90
wb.active = 0

for w in wb.worksheets:
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.orientation = "landscape"
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
wb.save(OUT)
print("saved", OUT)
