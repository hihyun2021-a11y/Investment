# -*- coding: utf-8 -*-
"""대출금 입금(인출) 요청 공문 날인·발송 기안(결재)문 — 사내 기안 양식(1.대상공문~5.발송방법, 첨부) 기준.
출력: 20261028_기안_대출금입금요청공문_날인발송_v01.docx
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "20261028_기안_대출금입금요청공문_날인발송_v01.docx")
FONT = "맑은 고딕"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.2)
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.2


def para(text="", indent=0.0, hanging=0.0, bold=False, size=None, align=None, before=0):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.left_indent = Cm(indent + hanging)
    if hanging:
        f.first_line_indent = Cm(-hanging)
    f.space_before = Pt(before)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
    return p


para("제목: 코람코라이프로지스 담보대출 대출금 입금(인출) 요청 공문 날인 및 발송의 건", bold=True, size=12)
para()
para("코람코라이프인프라리츠의 자회사인 코람코라이프로지스의 담보대출 인출을 위하여, 대주기관이 요청한 아래와 같은 "
     "내용의 대출금 입금(인출) 요청 공문을 날인 및 발송하고자 하오니 승인하여 주시기를 바랍니다.")
para("- 아 래 -", align=WD_ALIGN_PARAGRAPH.CENTER, before=8)

para("1. 대상공문", bold=True, before=8)
para("가. 대출금 입금(인출) 요청 공문", indent=0.4)
for k, v in [
    ("명의", "주식회사 코람코라이프로지스위탁관리부동산투자회사 (대표이사 원재이)"),
    ("제목", "코람코라이프로지스리츠 담보대출 인출(대출금 입금) 요청의 건"),
    ("수신", "우리은행 등 25개 대주기관"),
    ("일자", "2026. 10. 28."),
    ("내용", "2026. 10. 27. 체결한 대출약정서에 따라 인출요청서를 제출하였으므로, 각 대주기관이 약정한 대출금을 "
             "아래와 같이 입금하여 줄 것을 요청함"),
]:
    para(f"- {k} : {v}", indent=0.8, hanging=1.6)
for t in ["① 입금일 : 2026. 10. 29.(목)",
          "② 입금계좌 : 우리은행 1005-504-949583 (예금주: 코람코라이프로지스위탁관리부동산투자회사)",
          "③ 입금금액 : 대주별 약정액 전액, 합계 1,980억원 (Tranche A-1 900억원, Tranche A-2 700억원, Tranche B 380억원)",
          "④ 별첨 : 차주 인출요청서 25부(대주별 각 1부)"]:
    para(t, indent=2.4, hanging=0.5)
para("- 사유 : 대출약정서 제4조 제1항에 따라 인출예정일 1영업일 전(10/28)에 인출요청서를 제출하며, 대주기관이 "
     "내부 기표 절차용으로 인출요청서와 함께 차주 명의의 입금 요청 공문을 요청함", indent=0.8, hanging=1.6)

para("2. 체결사유", bold=True, before=8)
para("코람코라이프로지스의 담보대출 인출(2026. 10. 29. 예정)을 위하여 대주기관이 인출요청서와 함께 차주 명의의 "
     "대출금 입금 요청 공문 제출을 요청하였음. 공문은 이미 체결된 대출약정서와 제출된 인출요청서의 내용(인출일, "
     "입금계좌, 대주별 금액)을 그대로 확인하여 요청하는 것으로, 새로운 의무나 추가 부담이 생기지 않음. 인출 "
     "전일까지 대주기관에 전달해야 하므로 날인 및 발송이 필요함.", indent=0.4)

para("3. 대출개요", bold=True, before=8)
for t in ["가. 대주 : 우리은행 등 25개 대주기관 (Tranche A 22개사, Tranche B 3개사)",
          "나. 차주 : 코람코라이프로지스위탁관리부동산투자회사",
          "다. 용도 : 안양물류센터 신규자산 편입 담보대출",
          "라. 금액 : 선순위 담보대출(1,600억원), 중순위 담보대출(380억원), 합계 1,980억원"]:
    para(t, indent=0.4, hanging=0.6)
for t in ["- 선순위 Tranche A-1 900억원(고정 연 5.00%), Tranche A-2 700억원(변동 CD+1.43%)",
          "- 중순위 Tranche B 380억원(고정 연 6.00%)"]:
    para(t, indent=2.2, hanging=0.3)
para("마. 일정 : 대출약정 체결 2026. 10. 27. / 인출요청서 제출 2026. 10. 28. / 인출 2026. 10. 29.", indent=0.4, hanging=0.6)

para("4. 승인을 구하는 사항", bold=True, before=8)
para("첨부1. 공문의 날인(법인인감) 및 발송", indent=0.4)

para("5. 발송방법 / 발송처", bold=True, before=8)
para("날인 후 스캔본 발송(원본은 요청 시 송부) / 우리은행 등 25개 대주기관 업무 담당자", indent=0.4)

para("첨부1. 대출금 입금(인출) 요청 공문", before=10)
para("첨부2. 차주 인출요청서(날인본) 25부.  끝.")

doc.core_properties.author = "코람코 상장리츠팀"
doc.save(OUT)
print("saved", OUT)
