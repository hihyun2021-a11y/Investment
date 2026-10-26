# -*- coding: utf-8 -*-
"""차주(코람코라이프로지스) 제출서류·날인·인출일 업무 진행 체크리스트(한 시트) 생성.

출처(운용역 업로드 2026-10-26, 원본 불변):
- 법무법인 세한 「소유권이전 및 신탁등기 필요서류」(2026-10-13)
- 법무법인 세한 「소유권이전등기 필요서류」(2026-10-22)
- 태평양 「코람코라이프로지스리츠 담보대출 관련 준비서류 및 필요서류 목록」(Checklist 25918417.v2)
- 이자적립금·수수료: 05_대출업무/20261026_금융_이자적립금_이자스케줄_기표일수수료_v02.xlsx
출력: 05_대출업무/20261026_금융_차주제출서류_운용지시체크리스트_v01.xlsx
"""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "20261026_금융_차주제출서류_운용지시체크리스트_v01.xlsx")
D = dt.date
SEHAN_T, SEHAN_O, BKL = "세한(신탁등기)", "세한(소유권이전)", "태평양 체크리스트"

# (단계, 구분, 서류·업무, 부수, 날인·형식, 제출처, 기한, 담당, 출처, 비고)
ROWS = [
 # 1. 발급 서류
 ("1. 발급 서류", "법인서류", "법인인감증명서", "4부", "원본(3개월 내)", "세한 3부(이전 2·신탁 1), 태평양 1부", D(2026,10,27), "사무수탁사", f"{SEHAN_O}·{SEHAN_T}·{BKL}", "10/26~27 발급 권장"),
 ("1. 발급 서류", "법인서류", "법인등기사항전부증명서", "3부", "원본(3개월 내)", "세한 2부(이전 1·신탁 1), 태평양 1부", D(2026,10,27), "사무수탁사", f"{SEHAN_O}·{SEHAN_T}·{BKL}", ""),
 ("1. 발급 서류", "납세", "지방세 납세증명서(신탁등기용)", "1부", "원본", "세한", D(2026,10,28), "사무수탁사", SEHAN_T, "담보부동산 목록(관양동 934·934-1·912-2 토지, 건물) 정확히 기재"),
 ("1. 발급 서류", "납세", "국세 완납증명서", "1부", "원본", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, "인출선행조건"),
 ("1. 발급 서류", "납세", "지방세 완납증명서", "1부", "원본", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, "인출선행조건"),
 ("1. 발급 서류", "법인서류", "사업자등록증 사본", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "법인서류", "정관 사본", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "법인서류", "이사회의사록 사본(대출약정 체결 승인, 10/26)", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "계좌", "프로젝트계좌 통장 사본 – 대출금입금계좌(1005-504-949583)", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "계좌", "프로젝트계좌 통장 사본 – 수입금계좌(1005-104-954081)", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "계좌", "프로젝트계좌 통장 사본 – 운영계좌(1005-804-963776)", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "계좌", "프로젝트계좌 통장 사본 – Tranche A 이자유보계좌(1005-504-965088)", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "계좌", "프로젝트계좌 통장 사본 – Tranche B 이자유보계좌(1005-504-970089)", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "계좌", "프로젝트계좌 통장 사본 – 보험금수령계좌(1005-204-955717)", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "계좌", "프로젝트계좌 통장 사본 – 임대차보증금계좌(1005-704-972584)", "1부", "원본대조필", "태평양(대주)", D(2026,10,27), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "주금", "주금납입 증빙서류", "1부", "사본(원본대조필)", "태평양(대주)", D(2026,10,28), "사무수탁사", BKL, "10/27 주금 납입 후"),
 ("1. 발급 서류", "주금", "3종 종류주식 상계합의서 사본", "1부", "원본대조필", "태평양(대주)", D(2026,10,28), "사무수탁사", BKL, "계약금 채권과 주금납입채무 상계"),
 ("1. 발급 서류", "주금", "3종 종류주식 상계 관련 차주 동의서 사본", "1부", "원본대조필", "태평양(대주)", D(2026,10,28), "사무수탁사", BKL, ""),
 ("1. 발급 서류", "법인서류", "위임장(법인인감) + 법인인감증명서 + 수임인 신분증 사본", "필요 시", "법인인감", "태평양(대주)·세한", D(2026,10,27), "사무수탁사", BKL, "대표이사가 아닌 자가 서명할 때만"),
 # 2. 날인 – 소유권이전등기
 ("2. 날인 – 소유권이전등기", "등기", "등기신청위임장(소유권이전)", "1부", "법인도장(세한 안내) [법인인감 통일 여부 확인]", "세한", D(2026,10,28), "사무수탁사", SEHAN_O, ""),
 ("2. 날인 – 소유권이전등기", "신고", "부동산거래계약신고서", "1부", "법인인감 + 간인", "세한", D(2026,10,28), "사무수탁사", SEHAN_O, "날인본 원본"),
 ("2. 날인 – 소유권이전등기", "신고", "부동산거래계약신고 위임장", "1부", "법인인감", "세한", D(2026,10,28), "사무수탁사", SEHAN_O, ""),
 ("2. 날인 – 소유권이전등기", "취득세", "취득세 신고서", "1부", "법인인감", "세한", D(2026,10,28), "사무수탁사", SEHAN_O, ""),
 ("2. 날인 – 소유권이전등기", "취득세", "취득가액 내역표", "1부", "법인인감", "세한", D(2026,10,28), "사무수탁사", SEHAN_O, ""),
 ("2. 날인 – 소유권이전등기", "취득세", "계정별원장", "1부", "법인인감", "세한", D(2026,10,28), "사무수탁사", SEHAN_O, "사무수탁사 작성"),
 ("2. 날인 – 소유권이전등기", "취득세", "지방세 감면신청 위임장", "1부", "법인인감", "세한", D(2026,10,28), "사무수탁사", SEHAN_O, ""),
 ("2. 날인 – 소유권이전등기", "취득세", "취득부동산 사용계획서", "1부", "법인인감", "세한", D(2026,10,28), "사무수탁사", SEHAN_O, ""),
 ("2. 날인 – 소유권이전등기", "계약", "부동산매매계약서(체결본 원본)", "1부", "법인인감", "세한", D(2026,10,28), "운용역·사무수탁사", SEHAN_O, "10/26 체결본"),
 # 3. 날인 – 신탁등기
 ("3. 날인 – 신탁등기", "등기", "등기신청위임장(신탁)", "1부", "법인인감", "세한", D(2026,10,28), "사무수탁사", SEHAN_T, ""),
 ("3. 날인 – 신탁등기", "계약", "부동산담보신탁계약서 및 특약", "2부", "법인인감 + 간인(천공)", "우리은행 신탁부(은행 보관용은 등기소 제출)", D(2026,10,27), "사무수탁사", SEHAN_T, "위탁자·수탁자 각 1부"),
 # 4. 날인 – 대출
 ("4. 날인 – 대출", "약정", "담보대출약정서", "1부", "법인인감 + 수입인지 35만원", "태평양", D(2026,10,27), "사무수탁사", BKL, "인지세 부담 주체 확인"),
 ("4. 날인 – 대출", "약정", "예금채권근질권설정계약서", "1부", "법인인감", "태평양", D(2026,10,27), "사무수탁사", BKL, ""),
 ("4. 날인 – 대출", "약정", "예금채권근질권설정 통지서(우리은행 앞)", "1부", "법인인감 + 확정일자", "우리은행(계좌개설은행)", D(2026,10,29), "사무수탁사", BKL, "확정일자 사본은 인출일 당일 제출"),
 ("4. 날인 – 대출", "약정", "보험근질권설정계약서", "1부", "법인인감", "태평양", D(2026,10,27), "사무수탁사", BKL, ""),
 ("4. 날인 – 대출", "약정", "보험근질권설정 통지서(보험회사 앞)", "1부", "법인인감 + 확정일자", "보험회사", D(2026,10,29), "사무수탁사", BKL, "승계 보험: 서식 가-2"),
 ("4. 날인 – 대출", "수수료", "수수료약정서 10건(우리은행 주선·참여, 에스비아이 주선, 롯데·우리카드·수협·신협대주단 참여, 키움·엠지 취급, 디비 참여)", "각 1부", "법인인감", "각 대주", D(2026,10,27), "사무수탁사", "수수료약정서", "롯데캐피탈만 2부 작성"),
 ("4. 날인 – 대출", "인출", "인출요청서", "1부", "법인인감", "우리은행(대리)", D(2026,10,28), "사무수탁사", BKL, "인출일 10/29 기재"),
 # 5. 인출일 처리
 ("5. 인출일(10/29) 처리", "자금", "Tranche A 이자적립금 입금(1005-504-965088) 및 입금증빙", "1,958,213,686원", "입금증", "태평양(대주)", D(2026,10,29), "사무수탁사", "이자 엑셀 v02", "A-2 기준금리 [추정] → 10/26~28 CD 평균으로 확정"),
 ("5. 인출일(10/29) 처리", "자금", "Tranche B 이자적립금 입금(1005-504-970089) 및 입금증빙", "574,684,930원", "입금증", "태평양(대주)", D(2026,10,29), "사무수탁사", "이자 엑셀 v02", ""),
 ("5. 인출일(10/29) 처리", "자금", "기표일 수수료 지급 26건", "2,894,000,000원", "이체확인증", "각 대주", D(2026,10,29), "사무수탁사", "이자 엑셀 v02", "기관별 계좌는 통장사본 수령 후 통보"),
 ("5. 인출일(10/29) 처리", "자금", "취득세 납부", "[금액 확인 필요]", "영수증", "세한(등기 접수 전)", D(2026,10,29), "사무수탁사", SEHAN_O, ""),
 ("5. 인출일(10/29) 처리", "제출", "예금·보험근질권 확정일자 통지서 사본 제출", "각 1부", "사본", "태평양(대주)", D(2026,10,29), "사무수탁사", BKL, ""),
 # 6. 인출 후
 ("6. 인출 후", "제출", "보험회사 승낙서(확정일자) 사본", "1부", "사본", "태평양(대주)", D(2026,11,9), "사무수탁사", BKL, "인출 후 7영업일"),
 ("6. 인출 후", "제출", "계좌개설은행(우리은행) 승낙서(확정일자) 사본 – 3개 계좌", "각 1부", "사본", "태평양(대주)", D(2026,11,9), "사무수탁사", BKL, "인출 후 7영업일"),
 ("6. 인출 후", "제출", "보험배서증권 원본(근질권 문언 기재)", "1부", "원본", "태평양(대주)", D(2026,11,12), "사무수탁사", BKL, "인출 후 10영업일"),
]

THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="1F3864")
SEC = PatternFill("solid", fgColor="D9E1F2")
INPUT = PatternFill("solid", fgColor="FFF2CC")
F = lambda **k: Font(**{"name": "맑은 고딕", "size": 10, **k})
WRAP = Alignment(vertical="center", wrap_text=True)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "체크리스트"
ws.sheet_view.showGridLines = False
ws["A1"] = "안양물류센터 — 차주(코람코라이프로지스) 제출서류·날인·인출일 업무 체크리스트"
ws["A1"].font = Font(name="맑은 고딕", size=14, bold=True)
ws["A2"] = ("기준 일정: 10/26 이사회·주총 → 10/27 약정 체결·주금 납입 → 10/28 인출요청 → 10/29 인출·등기 접수. "
            "출처: 세한 신탁등기(10/13)·소유권이전등기(10/22) 필요서류, 태평양 대출 체크리스트. 사본은 원본대조필, 원본 증명서는 3개월 이내 발급분.")
ws["A2"].font = F(color="595959")
ws.merge_cells("A2:M2")
ws["A2"].alignment = WRAP
ws.row_dimensions[2].height = 30

# 진행 현황 요약 (4행)
FIRST, LAST = 7, 6 + len(ROWS)
rng = f"$K${FIRST}:$K${LAST}"
summ = [("전체", f"=COUNTA($D${FIRST}:$D${LAST})"), ("완료", f'=COUNTIF({rng},"완료")'),
        ("진행중", f'=COUNTIF({rng},"진행중")'), ("미착수", f'=COUNTIF({rng},"미착수")'),
        ("해당없음", f'=COUNTIF({rng},"해당없음")'), ("진행률", "=IFERROR(C5/(B5-F5),0)"),
        ("기한 경과·미완료", f'=SUMPRODUCT(($H${FIRST}:$H${LAST}<TODAY())*($K${FIRST}:$K${LAST}<>"완료")*($K${FIRST}:$K${LAST}<>"해당없음"))')]
ws["A4"] = "진행 현황"
ws["A4"].font = F(bold=True)
ws["A4"].alignment = CEN
ws.merge_cells("A4:A5")
for k, (lab, fml) in enumerate(summ):
    c1 = ws.cell(4, 2 + k, lab)
    c2 = ws.cell(5, 2 + k, fml)
    c1.font, c1.fill, c1.alignment, c1.border = F(bold=True, color="FFFFFF"), HEAD, CEN, BOX
    c2.font, c2.alignment, c2.border = F(bold=True, size=11), CEN, BOX
ws.cell(5, 7).number_format = "0%"

hdr = ["No", "단계", "구분", "서류·업무", "부수·금액", "날인·형식", "제출처", "기한", "담당", "출처",
       "진행상태", "완료일", "비고·메모"]
for i, h in enumerate(hdr, 1):
    c = ws.cell(6, i, h)
    c.font, c.fill, c.alignment, c.border = F(bold=True, color="FFFFFF"), HEAD, CEN, BOX
ws.row_dimensions[6].height = 22

prev = None
for i, (stage, kind, item, qty, seal, to, due, who, src, note) in enumerate(ROWS):
    r = FIRST + i
    vals = [i + 1, stage, kind, item, qty, seal, to, due, who, src, "미착수", None, note]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.font, c.border, c.alignment = F(), BOX, WRAP
    for j in (1, 3, 5, 8, 9, 11, 12):
        ws.cell(r, j).alignment = CEN
    ws.cell(r, 8).number_format = "yyyy-mm-dd"
    ws.cell(r, 12).number_format = "yyyy-mm-dd"
    ws.cell(r, 11).fill = INPUT
    ws.cell(r, 12).fill = INPUT
    ws.cell(r, 13).fill = INPUT
    if stage != prev:
        ws.cell(r, 2).font = F(bold=True)
    prev = stage

dv = DataValidation(type="list", formula1='"미착수,진행중,완료,해당없음"', allow_blank=False,
                    showErrorMessage=True, errorTitle="진행상태", error="목록에서 선택하세요.")
ws.add_data_validation(dv)
dv.add(f"K{FIRST}:K{LAST}")
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
ws.add_data_validation(dvd)
dvd.add(f"L{FIRST}:L{LAST}")

rows_ref = f"A{FIRST}:M{LAST}"
ws.conditional_formatting.add(rows_ref, FormulaRule(formula=[f'$K{FIRST}="완료"'], fill=PatternFill("solid", fgColor="E2EFDA"), font=Font(color="7F7F7F")))
ws.conditional_formatting.add(rows_ref, FormulaRule(formula=[f'$K{FIRST}="해당없음"'], fill=PatternFill("solid", fgColor="EDEDED"), font=Font(color="A6A6A6", strike=True)))
ws.conditional_formatting.add(rows_ref, FormulaRule(formula=[f'AND($K{FIRST}<>"완료",$K{FIRST}<>"해당없음",$H{FIRST}<TODAY())'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws.conditional_formatting.add(rows_ref, FormulaRule(formula=[f'$K{FIRST}="진행중"'], fill=PatternFill("solid", fgColor="DDEBF7")))

nr = LAST + 2
for t in ["※ 진행상태(노란 칸)는 목록에서 선택: 미착수 / 진행중 / 완료 / 해당없음. 완료=초록, 진행중=파랑, 기한 경과·미완료=주황으로 표시됩니다.",
          "※ 확인 필요: ① 소유권이전 위임장 매수인 날인(세한 '법인도장' 표기) 법인인감 통일 여부 ② 대출약정서 수입인지 35만원 부담 주체 ③ 취득세 금액 ④ A-2 기준금리 확정(10/26~28 CD 평균) 후 이자적립금 금액.",
          "※ 차주 외 서류(KLI 리파이낸싱 동의 공문·종류주식 매입확약서, 매도인 매도용 인감증명서·등기필증 등)는 이 목록에서 제외했습니다.",
          "※ 세한 등기법무지원팀 02-560-5171~2."]:
    ws.cell(nr, 1, t).font = F(color="595959", size=9)
    ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=13)
    nr += 1

for col, w in zip("ABCDEFGHIJKLM", [5, 20, 8, 52, 15, 24, 28, 12, 13, 22, 11, 12, 34]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = f"E{FIRST}"
ws.auto_filter.ref = f"A6:M{LAST}"
ws.print_title_rows = "6:6"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("saved", OUT, len(ROWS), "rows")
