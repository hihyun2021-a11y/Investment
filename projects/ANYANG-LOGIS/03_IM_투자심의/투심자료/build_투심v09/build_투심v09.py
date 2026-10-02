# -*- coding: utf-8 -*-
"""ANYANG-LOGIS 본 투심자료 v08(운용역 취합본) → v09 보완.

사용자 지시(2026-10-02):
  ① 목차 페이지 세부 테마 정합  ② 전 페이지 표 우측 상단 단위 표기 (예: (단위: 백만원, %))
  ③ 투자심의위원회 개요 → 첨부 계약서(매매계약 BKL v11, 주주간계약 BKL v30) 기준 최신화
  ④ 재무수치 → 재무모델 260930 기준 갱신
"""
import copy, json, pathlib, re, sys
from lxml import etree
import openpyxl
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
SRC = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else HERE.parent / "20261002_투심자료_ANYANG-LOGIS_v08.pptx"
MODEL = pathlib.Path(sys.argv[3]) if len(sys.argv) > 3 else HERE.parents[2] / "04_Equity투자자/models/20260930_재무모델_Financial_Modeling_v03.xlsm"
OVERVIEW = HERE / "투심개요_계약최신화_문안.json"
OUT = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "out.pptx"
TAG = "재무모델 2026-09-30"
REIT = "주식회사 코람코라이프로지스위탁관리부동산투자회사"
LOG = []

# ── 모델 ─────────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(MODEL, data_only=True, read_only=True)
def grab(name, r0=1, r1=200, c0=1, c1=80):
    d = {}
    for row in wb[name].iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1):
        for c in row:
            if c.value is not None:
                d[c.coordinate] = c.value
    return d
AR, IM, SD = grab("A&R", 1, 140, 1, 70), grab("IM", 1, 170, 1, 40), grab("Sell-down", 1, 95, 1, 32)
A = lambda k: AR.get(k, 0)
I = lambda k: IM.get(k, 0)
D = lambda k: SD.get(k, 0)

def n0(v):
    if v is None or v == 0 or (isinstance(v, (int, float)) and abs(v) < 0.5):
        return "-"
    return f"{v:,.0f}" if v > 0 else f"({abs(v):,.0f})"
pc = lambda v, nd=2: f"{v:.{nd}%}"
num = lambda v: "-" if not v else f"{v:,.2f}"

prs = Presentation(str(SRC))
S = lambda n: prs.slides[n - 1]

def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)

def find(slide, sid):
    for sh in walk(slide.shapes):
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)

def set_para(p, text):
    runs = p.runs
    if not runs:
        r = p.add_run()
        end = p._p.find(qn("a:endParaRPr"))
        if end is not None:
            rp = copy.deepcopy(end); rp.tag = qn("a:rPr")
            old = r._r.find(qn("a:rPr"))
            if old is not None: r._r.remove(old)
            r._r.insert(0, rp)
        runs = [r]
    runs[0].text = text
    for x in runs[1:]:
        p._p.remove(x._r)

def set_lines(tf, lines):
    tf = tf.text_frame if hasattr(tf, "text_frame") else tf
    body = tf._txBody
    while len(tf.paragraphs) < len(lines):
        body.append(copy.deepcopy(tf.paragraphs[-1]._p))
    while len(tf.paragraphs) > len(lines):
        body.remove(tf.paragraphs[-1]._p)
    for p, t in zip(tf.paragraphs, lines):
        set_para(p, t)

def put(tbl, i, j, text):
    cell = tbl.cell(i, j)
    set_lines(cell.text_frame, str(text).split("\n"))

def put_row(tbl, i, vals, j0=0):
    for k, v in enumerate(vals):
        put(tbl, i, j0 + k, v)

def replace_text(slide_or_shape, old, new, count_log=True):
    shapes = walk(slide_or_shape.shapes) if hasattr(slide_or_shape, "shapes") else [slide_or_shape]
    hit = 0
    for sh in shapes:
        frames = []
        if sh.has_text_frame:
            frames.append(sh.text_frame)
        if getattr(sh, "has_table", False) and sh.has_table:
            frames += [c.text_frame for r in sh.table.rows for c in r.cells]
        for tf in frames:
            for p in tf.paragraphs:
                if old in p.text:
                    t = p.text.replace(old, new)
                    set_para(p, t); hit += 1
    if count_log:
        LOG.append(f"replace '{old}'→'{new}': {hit}")
    return hit

def table(slide, sid):
    return find(slide, sid).table

# ══════════════════════════════════════════════════════════════════════
# ① 목차
# ══════════════════════════════════════════════════════════════════════
set_lines(find(S(2), 9), [
    "심의안건 및 추진일정",
    "안건 Ⅰ. 매매계약서 체결의 건",
    "안건 Ⅱ. 주주간계약서 체결의 건",
    "안건 Ⅲ. 코람코라이프인프라리츠 매입확약서(LOC) 날인의 건",
    "참고. 관계사 인수확약서, 참여의향서 등",
    "참고. 당사 예상수익",
    "[별첨] 투자 대상 물건 설명자료",
])

def rebuild(tf_shape, items):
    """items: (text, kind) kind=head/sub/blank. 기존 문단을 서식 견본으로 재사용."""
    tf = tf_shape.text_frame
    ps = tf.paragraphs
    tpl = {"head": None, "sub": None, "blank": None}
    for p in ps:
        t = p.text
        if not t.strip():
            if tpl["blank"] is None: tpl["blank"] = p._p
        elif t.lstrip().startswith("-"):
            if tpl["sub"] is None: tpl["sub"] = p._p
        else:
            if tpl["head"] is None: tpl["head"] = p._p
    tpl["blank"] = tpl["blank"] if tpl["blank"] is not None else tpl["sub"]
    tpl = {k: copy.deepcopy(v) for k, v in tpl.items()}
    body = tf._txBody
    for p in list(ps):
        body.remove(p._p)
    from pptx.text.text import _Paragraph
    for text, kind in items:
        el = copy.deepcopy(tpl[kind]); body.append(el)
        para = _Paragraph(el, tf)
        if kind == "blank":
            for r in para.runs: el.remove(r._r)
        else:
            set_para(para, text)

SUB = "     - "
rebuild(find(S(14), 9), [
    ("Investment Highlights", "head"),
    ("Ⅰ. 사업개요", "head"),
    (SUB + "투자자산 개요", "sub"), (SUB + "입지분석", "sub"), (SUB + "시장현황 및 입지분석", "sub"),
    (SUB + "유사사례 검토", "sub"), (SUB + "적정 임대료 및 매매가", "sub"),
    ("", "blank"),
    ("Ⅱ. 투자계획", "head"),
    (SUB + "리츠개요", "sub"), (SUB + "투자구조", "sub"), (SUB + "투자 Milestone", "sub"),
    (SUB + "재원조달", "sub"), (SUB + "운영가정", "sub"), (SUB + "투자수익률", "sub"),
    (SUB + "대출가정 및 DSCR", "sub"), (SUB + "민감도분석", "sub"),
])
rebuild(find(S(14), 11), [
    ("Ⅲ. 운영계획", "head"),
    (SUB + "운용전략 (Re-Tenanting)", "sub"), (SUB + "상장리츠협업 CPM", "sub"),
    ("", "blank"),
    ("Ⅳ. 모리츠(KLI리츠) 투자계획", "head"),
    (SUB + "투자개요", "sub"), (SUB + "투자전략", "sub"), (SUB + "투자효과", "sub"),
    (SUB + "투자의사결정 수권절차", "sub"), (SUB + "회수전략", "sub"),
    ("", "blank"),
    ("Ⅴ. 리스크 요인분석 및 관리방안", "head"),
    (SUB + "주요 리스크 및 대응방안", "sub"),
])

# ══════════════════════════════════════════════════════════════════════
# ③ 투자심의위원회 개요 (법률 검토 문안 JSON)
# ══════════════════════════════════════════════════════════════════════
if OVERVIEW.exists():
    OV = json.loads(OVERVIEW.read_text(encoding="utf-8"))
    def rows_into(tbl, rows, start=1):
        for i, row in enumerate(rows, start=start):
            if i >= len(tbl.rows):
                LOG.append(f"!! 행 초과 {row[0]}"); break
            for j, v in enumerate(row):
                if j < len(tbl.columns):
                    put(tbl, i, j, v)
    o = OV.get("slide3", {})
    if o:
        if o.get("lead"): set_lines(find(S(3), 3), o["lead"])
        t = table(S(3), 11)
        for i, v in enumerate(o.get("table", []), start=1):
            put(t, i, 2, v[0] if isinstance(v, list) else v)
        if o.get("footnote"): set_lines(find(S(3), 109), o["footnote"] if isinstance(o["footnote"], list) else [o["footnote"]])
    for sn, sid, fid in ((4, 8, 16), (6, 19, 16), (8, 10, 16)):
        o = OV.get(f"slide{sn}", {})
        if o.get("rows"): rows_into(table(S(sn), sid), o["rows"])
        if o.get("footnote"): set_lines(find(S(sn), fid), [o["footnote"]])
    o = OV.get("slide5", {})
    if o.get("rows"): rows_into(table(S(5), 15), o["rows"])
    if o.get("footnote"): set_lines(find(S(5), 13), [o["footnote"]])
    o = OV.get("slide7", {})
    if o.get("footnote"): set_lines(find(S(7), 16), [o["footnote"]])
    if o.get("cap_rows"): rows_into(table(S(7), 19), o["cap_rows"])
    o = OV.get("slide9", {})
    if o.get("table1"): rows_into(table(S(9), 17), o["table1"], start=0)
    if o.get("table2"): rows_into(table(S(9), 18), o["table2"], start=0)
    if o.get("footnote"): set_lines(find(S(9), 16), [o["footnote"]])
    t = table(S(7), 19)
    for i, nm in ((1, "이지스리츠3호펀드"), (4, "엠지캐피탈"), (7, "이지스K리츠펀드"), (9, "코크렙안양(매도인)")):
        put(t, i, 1, nm)
    set_lines(find(S(8), 7), ["주요 조건 (정관)"])
    for sn, fid in ((4, 16), (5, 13), (6, 16), (7, 16), (8, 16), (9, 16)):
        find(S(sn), fid).width = Emu(10009067)
    LOG.append("overview applied")

# ══════════════════════════════════════════════════════════════════════
# ④ 재무수치 (재무모델 260930)
# ══════════════════════════════════════════════════════════════════════
# 리츠명(가칭 → 확정 상호)
for sl in prs.slides:
    replace_text(sl, "(가칭) 주식회사 케이라이프로지스1호위탁관리부동산투자회사", REIT, count_log=False)
set_lines(find(S(28), 62), ["코람코라이프", "로지스리츠", "(KLL리츠)"])
set_lines(find(S(28), 189), ["코람코라이프로지스리츠", "(KLL리츠)"])
set_lines(find(S(28), 55), ["이지스·키움 등 FI", "삼성증권 등"])
set_lines(find(S(28), 52), ["매도인 측"])
t = table(S(29), 137)
put_row(t, 1, ["1종우선주", "FI(이지스 등)", "FI(이지스 등)", "KLI리츠"])
put_row(t, 2, ["2종우선주", "KLI·FI", "KLI리츠", "KLI리츠"])
put_row(t, 3, ["3종우선주", "매도인 측", "매도인 측", "원본감자"])
set_lines(find(S(30), 15), ["  제1종우선주 매입확약 조건 주요 Term"])
set_lines(find(S(30), 18), ["  제2종우선주 매입확약 조건 주요 Term"])
set_lines(find(S(42), 15), ["  제1종우선주 매입확약 조건 주요 Term"])
set_lines(find(S(42), 18), ["  제2종우선주 매입확약 조건 주요 Term"])

# p15 Investment Highlights
replace_text(find(S(15), 29), "IRR 16.44%", f"IRR {I('J102'):.2%}")
replace_text(find(S(15), 30), "E. Multiple 4.35 예상", f"E. Multiple {I('J103'):.2f} 예상")
replace_text(find(S(15), 42), "10년 평균 추정 IRR 9.24%", f"10년 평균 추정 IRR {I('J50'):.2%}")

# p26 리츠개요
s = S(26)
t = table(s, 11)
put_row(t, 1, [pc(I("C49")), pc(I("J49")), pc(I("C101")), pc(I("J101")), pc(I("U101"))], 1)
put_row(t, 2, [pc(A("G66")), pc(A("K66")), pc(A("G96")), pc(A("K96")), pc(I("U102"))], 1)
put_row(t, 3, [pc(I("C50")), pc(I("J50")), pc(I("C102")), pc(I("J102")), pc(I("U103"))], 1)
set_lines(find(s, 12), [
    f"Equity 전체 연평균 CoC {I('U101'):.2%}(매각차익 제외), IRR {I('U103'):.2%} 예상",
    f"1종 CoC {I('C49'):.2%}·IRR {I('C50'):.2%}, 2종 CoC {I('J49'):.2%}·IRR {I('J50'):.2%}, "
    f"보통주 CoC {I('J101'):.2%}·IRR {I('J102'):.2%}"])
find(s, 7).width = Emu(10009067)
set_lines(find(s, 7), [f"총 자산가액: 부동산 장부가액 {A('P28')/1e8:,.0f}억원 및 여유현금 100억원 예정 / 수익률·보수는 {TAG} 기준 "
                       f"(Equity 전체 IRR: IM 탭 {I('U103'):.2%})"])

# p31 재원조달
s = S(31)
t = table(s, 3)
due = I("R47"); fin = I("R48")
put_row(t, 4, [n0(due), f"{due*1000/A('T13'):.1%}"], 1)
put_row(t, 5, [n0(fin), f"{fin*1000/A('T13'):.1%}"], 1)
t = table(s, 29)
put(t, 2, 5, "우리은행 금융확약서(LOC) 확보(9/21)")
put(t, 3, 5, "키움캐피탈·MG캐피탈 등 LOC 확보")
put(t, 4, 5, "이지스·키움·애큐온·MG 인수 (주주간계약 제2.1조)")
put(t, 5, 5, "KLI 160억·FI 340억 인수, 12개월 내 KLI 매입")
put(t, 6, 3, "무배당·무의결권 (발행가 23,000원)")
put(t, 6, 5, "매도인 측 재투자 (주주간계약 제5조 유상감자)")
put(t, 7, 5, "KLI리츠 이사회 승인 예정(10/14)")

# p35·36 투자자별 수익률 (IM 탭)
def inv_rows(cols, r0, rt):
    b, d, g, tt, y = cols
    rows = []
    for k in range(20):
        r = r0 + k
        bal = I(f"{b}{r}")
        rows.append([f"FY {k+1}", n0(bal), n0(I(f"{d}{r}")), n0(I(f"{g}{r}")), n0(I(f"{tt}{r}")),
                     (pc(I(f"{y}{r}"), 1) if bal else "-")])
    rows.append(["합계", "", n0(I(f"{d}{rt}")), n0(I(f"{g}{rt}")), n0(I(f"{tt}{rt}")), ""])
    rows.append(["투자금액", n0(I(f"{b}{rt+1}")), "", "", "", ""])
    rows.append(["CoC_C.G제외", pc(I(f"{b}{rt+2}")), "", "", "", ""])
    rows.append(["IRR_C.G포함", pc(I(f"{b}{rt+3}")), "", "", "", ""])
    rows.append(["E.Multiple", f"x {I(f'{b}{rt+4}'):.2f}", "", "", "", ""])
    return rows
for (sn, sid, cols, r0, rt) in ((35, 36, ("C", "D", "E", "F", "G"), 27, 47), (35, 38, ("J", "K", "L", "M", "N"), 27, 47),
                                (36, 2, ("C", "D", "E", "F", "G"), 79, 99), (36, 43, ("J", "K", "L", "M", "N"), 79, 99)):
    t = table(S(sn), sid)
    for i, row in enumerate(inv_rows(cols, r0, rt), start=1):
        put_row(t, i, row)
j2 = [I(f"K{r}") for r in range(27, 47)]
short = [k + 1 for k, v in enumerate(j2) if v < 1875 - 0.5]
over = [k + 1 for k, v in enumerate(j2) if v > 1875 + 0.5]
set_lines(find(S(35), 5), [
    f"사업기간 10년 가정 투자자별 수익률은 1종 IRR {I('C50'):.2%}, 2종 IRR {I('J50'):.2%} 예상",
    f"제1종 연 7.0%·제2종 연 7.5% 누적배당 — 2종은 FY{short[0]}~{short[-1]} 미지급분을 FY{over[0]}~{over[-1]}에 누적 지급(3년차 Rent-Free 영향)"])
set_lines(find(S(35), 47), [f"Source: {TAG} IM 탭(B25:N51), FY는 반기"])
set_lines(find(S(36), 11), [
    f"3종은 무배당·2년 후 원본 유상감자, 보통주 IRR {I('J102'):.2%}·E.M {I('J103'):.2f}x 예상",
    f"보통주는 초기 1년 연 7.5% 배당 후 잔여배당, 매각차익의 80%({n0(I('L98'))}백만원) 배분"])
set_lines(find(S(36), 109), [f"Source: {TAG} IM 탭(B77:N103). 3종 E.Multiple 0.00x는 배당·C.G 기준 산출값으로 "
                             "FY4 원본 유상감자(100억원) 회수분 제외 (FY는 반기)"])

# p37 DSCR (A&R X68:AI81)
s = S(37)
t = table(s, 8)
YR = ["Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI"]
spec = [(2, 68, n0), (3, 69, n0), (4, 70, n0), (5, 71, num), (6, 72, num), (7, 73, n0), (8, 74, num), (9, 75, num),
        (10, 76, n0), (11, 77, num), (12, 78, num), (13, 79, n0), (14, 80, num), (15, 81, num)]
for i, r, f in spec:
    put(t, i, 2, f(A(f"X{r}")) if A(f"X{r}") else "")
    put(t, i, 3, f(A(f"Y{r}")) if A(f"Y{r}") else "")
    put_row(t, i, [f(A(f"{c}{r}")) for c in YR], 4)
set_lines(find(s, 5), [
    f"선순위 단순 DSCR은 Carry 2개년 최소 {A('X71'):,.2f}·평균 {A('Y71'):,.2f}, 중순위 최소 {A('X74'):,.2f} 수준",
    f"2종 우선주 누적 DSCR은 운영기간 최소 {min(A(c + '81') for c in YR[:-1]):,.2f} 수준 (3년차 Rent-Free 영향, 미지급분은 이후 누적 지급)"])

# p38 민감도 (A&R AW~BC)
s = S(38)
S7 = ["AW", "AX", "AY", "AZ", "BA", "BB", "BC"]
S6 = ["AX", "AY", "AZ", "BA", "BB", "BC"]
t = table(s, 6)
put_row(t, 1, [f"{A(c + '61'):,.0f}천원/평" for c in S7], 1)
for i, r in enumerate((62, 63, 64, 65, 66), start=2):
    put_row(t, i, [pc(A(f"{c}{r}")) for c in S7], 1)
t = table(s, 16)
for i, r in enumerate((71, 72, 73, 74, 75), start=1):
    put_row(t, i, [pc(A(f"{c}{r}")) for c in S7], 1)
t = table(s, 28)
for i, r in enumerate((79, 80), start=1):
    put_row(t, i, [pc(A(f"{c}{r}")) for c in S6], 2)
for i, r in enumerate((81, 82, 83, 84, 85), start=3):
    put_row(t, i, [pc(A(f"{c}{r}")) for c in S6], 2)
for i, r in enumerate((86, 87, 88, 89), start=8):
    put_row(t, i, [num(A(f"{c}{r}")) for c in S6], 2)
set_lines(find(s, 29), ["* Carry 기간(24개월) 내 DSCR 최소치 / Base 열은 모델 시나리오 산출값으로 기준 Case(IM 탭)와 소폭 차이"])

# p41 KLI 투자개요 — Sell-down 탭
s = S(41)
t = table(s, 6)
def coc_cg(dv, cg, inv, months):
    return (dv + cg) / inv * 12 / months
RET = {
    "1종 우선주": (D("H19"), coc_cg(D("J17"), D("J18"), D("I15"), D("H20")), D("I20")),
    "2종 우선주": (D("H30"), coc_cg(D("J28"), D("J29"), D("I26"), D("H31")), D("I31")),
    "보통주": (D("H41"), coc_cg(D("J39"), D("J40"), D("I37"), D("H42")), D("I42")),
    "KLI 수익률": (D("AC90"), (D("AA88") + D("AB88")) / D("AC89") * 12 / D("H42"), D("AC92")),
}
for i in range(1, len(t.rows)):
    k = t.cell(i, 0).text.strip()
    if k in RET:
        put_row(t, i, [pc(v) for v in RET[k]], 1)
for sh in walk(s.shapes):
    if sh.has_text_frame and "Sell-down 탭" in sh.text_frame.text:
        set_lines(sh, [f"※ {TAG} Sell-down 탭 (보통주 ’26.10 출자, 2종 ’27.10·1종 ’28.10 매입, ’36.10 매각) / "
                       "CoC(매각차익 포함) = (운영배당+매각차익배분)÷투자원금÷보유기간"])

# 모델 표기 일괄 갱신
for sl in prs.slides:
    replace_text(sl, "재무모델 2026-09-15 인가ver(v02)", TAG, count_log=False)
    replace_text(sl, "기준 재무모델 v02", TAG, count_log=False)

# ══════════════════════════════════════════════════════════════════════
# ② 단위 표기 (표·차트 우측 상단)
# ══════════════════════════════════════════════════════════════════════
UNIT_TPL = copy.deepcopy(find(S(31), 33)._element)
def unit_label(slide, sid, text, dy=0):
    sh = find(slide, sid)
    el = copy.deepcopy(UNIT_TPL)
    right = sh.left + sh.width
    w = 2600000
    xfrm = el.find(qn("p:spPr")).find(qn("a:xfrm"))
    off = xfrm.find(qn("a:off")); ext = xfrm.find(qn("a:ext"))
    off.set("x", str(int(right - w))); off.set("y", str(int(sh.top - 246221 + dy)))
    ext.set("cx", str(w))
    for cnv in el.iter(qn("p:cNvPr")):
        cnv.set("id", str(9000 + len(LOG))); cnv.set("name", f"단위 {sid}")
    slide.shapes._spTree.append(el)
    from pptx.shapes.autoshape import Shape
    ps = list(el.iter(qn("a:p")))
    for extra in ps[1:]:
        extra.getparent().remove(extra)
    rs = list(ps[0].iter(qn("a:r")))
    for extra in rs[1:]:
        ps[0].remove(extra)
    rs[0].find(qn("a:t")).text = text
    LOG.append(f"unit S{prs.slides.index(slide)+1} #{sid} {text}")

UNITS = [
    (7, 19, "(단위: 백만원, 원, 주, %)"), (12, 9, "(단위: 억원)"),
    (19, 28, "(단위: 명, 분, 만원/평)"), (20, 366, "(단위: 평, 원/평)"), (21, 435, "(단위: ㎡)"),
    (23, 26, "(단위: 평, 원/평, 개월)"), (24, 22, "(단위: 평, 억원, 만원/평, %)"),
    (25, 22, "(단위: 원/평, %)"), (25, 23, "(단위: 원/평)"),
    (32, 46, "(단위: 평, 원/평, %)"), (33, 39, "(단위: 원/평, 백만원, %)"), (33, 17, "(단위: 백만원)"),
    (48, 88, "(단위: %)"),
]
for sn, sid, txt in UNITS:
    unit_label(S(sn), sid, txt)
# 기존 단위 표기 서식 통일 "(단위: 백만원, %)"
for sl in prs.slides:
    for sh in walk(sl.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip().startswith(("(단위", "(IRR")):
            t0 = sh.text_frame.text.strip()
            t1 = re.sub(r"\(\s*IRR\s*,\s*단위\s*:\s*%\s*\)", "(단위: %, IRR 기준)", t0)
            t1 = re.sub(r"\(\s*단위\s*:\s*", "(단위: ", t1)
            t1 = re.sub(r"\s*,\s*", ", ", t1)
            if t1 != t0:
                set_lines(sh, [t1])
            if sh.width < 1800000:
                r = sh.left + sh.width
                sh.width = Emu(1800000); sh.left = Emu(r - 1800000)
                for p_ in sh.text_frame.paragraphs:
                    p_.alignment = 3  # 오른쪽

# 공통: 세로 중간 정렬 (사내 표준)
for sl in prs.slides:
    for el in sl.shapes._spTree.iter(qn("a:tcPr")):
        el.set("anchor", "ctr")

prs.save(str(OUT))
print("\n".join(LOG[-40:]))
print("saved", OUT)
