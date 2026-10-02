# -*- coding: utf-8 -*-
"""ANYANG-LOGIS 본 투심자료 v10 → v12 (v11 구성 + 3~11p 표를 실제 표 객체로).

v12 (2026-10-02 운용역 지시): v11의 3~11p 내용은 그대로 두고, 도형·텍스트상자로 표처럼 그린 부분을
실제 PowerPoint 표로 변경 — 심의안건 표(4p), 계약서 주요 내용 '구분|주요내용' 표(6~11p, 내장 하위표 포함 한 개의 표).
12p(매입확약서)는 지시 범위 밖이라 v11 형태 유지.

(이하 v11 설명)

사용자 지시(2026-10-02): 사내 표준 투심자료(현대차 실물자산 유동화 프로젝트 투심자료 v1.0, 260325)의
'내용 전달 구성 방식'만 참고하여 투자심의위원회 개요 섹션을 재구성한다.
 - 첨부자료의 수치는 사용하지 않는다.
 - 기존 투심자료(v10)의 내용·수치는 수정하지 않는다(배치·표현 형식만 변경).
 - 예) 임대료는 층별·임차인별 임대료, 보증금 등으로 표시.
"""
import copy, math, pathlib, re, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_PATTERN, MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
SRC = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else HERE.parent / "20261002_투심자료_ANYANG-LOGIS_v10.pptx"
OUT = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "20261002_투심자료_ANYANG-LOGIS_v12.pptx"
F_B, F_M, F_L = "KoPubWorld돋움체 Bold", "KoPubWorld돋움체 Medium", "KoPubWorld돋움체 Light"
NAVY, BAND, INK, GRAY, RED = "002060", "345B86", "1F1F1F", "595959", "C00000"
LBLUE, LGRAY, MGRAY, ORANGE, LORANGE = "DAE3F3", "F2F2F2", "D9D9D9", "F78E3F", "FDE9D9"
RULE = "BFBFBF"
SPA = "부동산매매계약서 BKL v11 clean(2026-09-30)"
SHA = "주주간계약서 BKL v30.5 markup(2026-10-02, 최종)"
LEGAL = "본 검토는 내부 검토이며 최종 법률자문은 외부 법무법인 확인이 필요함."

prs = Presentation(str(SRC))
SL = list(prs.slides)                       # v10 슬라이드 원본 순서 (1-base: SL[n-1])
IN = lambda v: Emu(int(round(v * 914400)))
L0, R0 = 0.57, 11.24
W0 = R0 - L0

# ── 공통 도구 ────────────────────────────────────────────────────────────
def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)

def find(slide, sid):
    for sh in walk(slide.shapes):
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)

def set_font(run, name, size, color=INK, bold=False):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)

def set_para(p, text):
    runs = p.runs
    if not runs:
        r = p.add_run()
        end = p._p.find(qn("a:endParaRPr"))
        if end is not None:
            rp = copy.deepcopy(end); rp.tag = qn("a:rPr")
            old = r._r.find(qn("a:rPr"))
            if old is not None: r._r.remove(old)
            r._r.insert(0, rp)
        runs = [r]
    runs[0].text = text
    for x in runs[1:]:
        p._p.remove(x._r)

def set_lines(shape, lines):
    tf = shape.text_frame
    body = tf._txBody
    while len(tf.paragraphs) < len(lines):
        body.append(copy.deepcopy(tf.paragraphs[-1]._p))
    while len(tf.paragraphs) > len(lines):
        body.remove(tf.paragraphs[-1]._p)
    for p, t in zip(tf.paragraphs, lines):
        set_para(p, t)

def band_label(slide, text, idx=0):
    groups = [g for g in slide.shapes if g.shape_type == 6 and any(
        c.has_text_frame and c.text_frame.text.strip() for c in g.shapes)]
    g = sorted(groups, key=lambda x: x.top)[idx]
    tb = [c for c in g.shapes if c.has_text_frame and c.text_frame.text.strip()][0]
    set_lines(tb, [text])
    tb.width = Emu(max(tb.width, 1950000))       # 띠 도형(2.21M) 안에서 글자 잘림 방지
    return g

TOKEN = re.compile(r"(\*\*.+?\*\*|\^.+?\^)")

def rich(p, text, size, color=INK, font=F_L, bold_font=F_B):
    """'**굵게**', '^위첨자^' 표기를 run으로 분리."""
    for part in TOKEN.split(text):
        if not part:
            continue
        r = p.add_run()
        if part.startswith("**"):
            r.text = part[2:-2]; set_font(r, bold_font, size, color)
        elif part.startswith("^"):
            r.text = part[1:-1]; set_font(r, font, size, color)
            r._r.get_or_add_rPr().set("baseline", "30000")
        else:
            r.text = part; set_font(r, font, size, color)

def bullet_ppr(p, kind, indent=0.15):
    pPr = p._p.get_or_add_pPr()
    for ch in list(pPr):
        pPr.remove(ch)
    sb = etree.SubElement(pPr, qn("a:spcBef")); etree.SubElement(sb, qn("a:spcPts"), val="200")
    if kind == "bullet":
        pPr.set("marL", str(int(indent * 914400))); pPr.set("indent", str(-int(indent * 914400)))
        etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
        etree.SubElement(pPr, qn("a:buChar"), char="•")
    else:
        pPr.set("marL", str(int(indent * 914400)) if kind == "cont" else "0"); pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))

def write(tf, lines, size=9, color=INK, font=F_L, align=PP_ALIGN.CENTER, bold_font=F_B):
    """lines: str 또는 (text, dict). '• ' 시작=글머리표, '  ' 시작=들여쓴 연속행, '▲' 시작=검토의견(적색)."""
    tf.clear(); tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, ln in enumerate(lines):
        txt, opt = (ln, {}) if isinstance(ln, str) else ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        col = opt.get("color", color)
        if txt.startswith("• "):
            bullet_ppr(p, "bullet"); txt = txt[2:]
        elif txt.startswith("  "):
            bullet_ppr(p, "cont"); txt = txt.strip()
        elif txt.startswith("▲"):
            bullet_ppr(p, "cont"); col = RED
        rich(p, txt, opt.get("size", size), col, opt.get("font", font), opt.get("bold_font", bold_font))

def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=0.75, dash=False):
    sh = slide.shapes.add_shape(kind, IN(x), IN(y), IN(w), IN(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
        if dash: sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        sh.line.fill.background()
    return sh

def box(slide, x, y, w, h, lines, fill=None, line=None, kind=MSO_SHAPE.RECTANGLE, lw=0.75, dash=False, **kw):
    sh = shape(slide, kind, x, y, w, h, fill, line, lw, dash)
    write(sh.text_frame, lines, **kw)
    return sh

def text(slide, x, y, w, h, lines, **kw):
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    write(tb.text_frame, lines, **kw)
    return tb

def line(slide, x1, y1, x2, y2, color=RULE, width=0.75, dash=False, head=False):
    ln = slide.shapes.add_connector(1, IN(x1), IN(y1), IN(x2), IN(y2))
    ln.line.color.rgb = RGBColor.from_string(color); ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        etree.SubElement(ln.line._get_or_add_ln(), qn("a:tailEnd"), type="triangle", w="med", len="med")
    return ln

def ln_el(tag, w=3175, color=RULE):
    el = etree.Element(qn(f"a:{tag}"), w=str(w), cap="flat", cmpd="sng", algn="ctr")
    if color is None:
        etree.SubElement(el, qn("a:noFill"))
    else:
        sf = etree.SubElement(el, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=color)
    return el

def table(slide, rows, x, y, widths, row_h=0.30, size=9, head=True, label_col=True, aligns=None,
          colors=None, total_rows=(), fills=None, merges=(), heights=None):
    nr, nc = len(rows), len(widths)
    gf = slide.shapes.add_table(nr, nc, IN(x), IN(y), IN(sum(widths)), IN(row_h * nr))
    t = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    for i, w in enumerate(widths):
        t.columns[i].width = IN(w)
    for ri, row in enumerate(rows):
        t.rows[ri].height = IN(heights[ri] if heights else row_h)
        for ci in range(nc):
            val = row[ci] if ci < len(row) else ""
            cell = t.cell(ri, ci)
            is_head = head and ri == 0
            is_tot = ri in total_rows
            is_label = label_col and ci == 0 and not is_head
            kind_font = F_B if (is_head or is_tot) else (F_M if is_label else F_L)
            al = PP_ALIGN.CENTER if (is_head or is_label) else (aligns[ci] if aligns else PP_ALIGN.LEFT)
            col = (colors or {}).get((ri, ci), INK)
            lines = val.split("\n") if isinstance(val, str) else val
            write(cell.text_frame, lines, size=size, font=kind_font, align=al, color=col)
            tcPr = cell._tc.get_or_add_tcPr()
            for ch in list(tcPr): tcPr.remove(ch)
            tcPr.set("anchor", "ctr"); tcPr.set("marL", "54000"); tcPr.set("marR", "36000")
            tcPr.set("marT", "9525"); tcPr.set("marB", "9525")
            tcPr.append(ln_el("lnL", 9525, None) if ci == 0 else ln_el("lnL"))
            tcPr.append(ln_el("lnR", 9525, None) if ci == nc - 1 else ln_el("lnR"))
            tcPr.append(ln_el("lnT", 6350, "000000") if ri == 0 else ln_el("lnT"))
            tcPr.append(ln_el("lnB", 6350, "000000") if ri == nr - 1 else ln_el("lnB"))
            fill = (fills or {}).get((ri, ci)) or (MGRAY if is_head else (LGRAY if (is_label or is_tot) else None))
            if fill:
                sf = etree.SubElement(tcPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=fill)
            else:
                etree.SubElement(tcPr, qn("a:noFill"))
    for (r1, c1, r2, c2) in merges:
        t.cell(r1, c1).merge(t.cell(r2, c2))
    return t

def unit(slide, right, y, txt):
    tb = slide.shapes.add_textbox(IN(right - 3.0), IN(y), IN(3.0), IN(0.2))
    write(tb.text_frame, [txt], size=8, color=INK, align=PP_ALIGN.RIGHT)
    tb.text_frame.margin_right = 0

# ── 표준양식형 '구분 | 주요내용' 표 (하위 표 내장 가능) ───────────────────────
def text_w(s, size):
    w = 0.0
    for ch in s.replace("**", "").replace("^", ""):
        o = ord(ch)
        if 0xAC00 <= o <= 0xD7A3 or 0x3130 <= o <= 0x318F or o in (0x2160, 0x2161, 0x2162, 0x203B, 0x00B7, 0x2192):
            w += 0.95
        elif ch == " ":
            w += 0.32
        elif ch.isupper() or ch in "%@&W":
            w += 0.68
        else:
            w += 0.56
    return w * size / 72

def lines_h(lines, size, width):
    n = 0
    for ln in lines:
        txt = ln if isinstance(ln, str) else ln[0]
        n += max(1, math.ceil(text_w(txt.strip(), size) / max(width - 0.3, 0.3)))
    return n * size * 1.22 / 72 + (len(lines) - 1) * 2 / 72 + 0.08

def sub_dims(data, opt, width):
    ws = opt["w"]; k = (opt.get("tw") or width) / sum(ws); ws = [v * k for v in ws]
    size = opt.get("size", 8); hs = []
    for row in data:
        n = 1
        for txt, w in zip(row, ws):
            m = sum(max(1, math.ceil(text_w(t, size) * 1.08 / max(w - 0.2, 0.2))) for t in txt.split("\n"))
            n = max(n, m)
        hs.append(max(opt.get("rh", 0.22), n * size * 1.22 / 72 + 0.07))
    return ws, hs

def block_h(b, width):
    kind, data, opt = b[0], b[1], (b[2] if len(b) > 2 else {})
    if kind == "b":
        return opt.get("h") or lines_h(data, opt.get("size", 9), width)
    return opt.get("h") or sum(sub_dims(data, opt, width)[1])

def place_block(slide, b, x, y, width):
    kind, data, opt = b[0], b[1], (b[2] if len(b) > 2 else {})
    h = block_h(b, width)
    if kind == "b":
        text(slide, x, y, width, h, data, size=opt.get("size", 9), align=PP_ALIGN.LEFT)
    else:
        ws, hs = sub_dims(data, opt, width)
        table(slide, data, x, y, ws, heights=hs, size=opt.get("size", 8),
              label_col=opt.get("label_col", False), aligns=opt.get("al"), total_rows=opt.get("tot", ()),
              colors=opt.get("colors"), fills=opt.get("fills"), merges=opt.get("merges", ()))
    return h

def grid(slide, y, rows, widths=(1.55, W0 - 1.55), header=("구분", "주요내용"), x=L0, head_h=0.28, gap=0.04, min_h=0.34):
    """rows: [(label, [block...], min_h)] — label은 '이름\\n(조항)' 형식, 둘째 줄부터 작은 회색 글씨."""
    cx = x
    for w, hd in zip(widths, header):
        box(slide, cx, y, w, head_h, [(hd, {"font": F_B, "size": 9.5})], fill=MGRAY)
        cx += w
    line(slide, x, y, x + sum(widths), y, color="000000", width=1.0)
    y += head_h
    for ri, row in enumerate(rows):
        label, blocks = row[0], row[1]
        mh = row[2] if len(row) > 2 else min_h
        cw = widths[1] - 0.2
        inner = sum(block_h(b, cw) for b in blocks) + gap * (len(blocks) - 1)
        h = max(mh, inner + 0.12)
        parts = label.split("\n")
        box(slide, x, y, widths[0], h, [(p, {"size": 7.5, "color": GRAY}) if p.startswith("(") else (p, {"font": F_M, "size": 9})
                                        for p in parts], fill=LGRAY)
        by = y + (h - inner) / 2
        for b in blocks:
            by += place_block(slide, b, x + widths[0] + 0.1, by, cw) + gap
        last = ri == len(rows) - 1
        line(slide, x, y + h, x + sum(widths), y + h, color="000000" if last else RULE, width=1.0 if last else 0.5)
        y += h
    return y


# ── 실제 표 객체 버전 '구분 | 주요내용' (v12) ──────────────────────────────
def cell_style(cell, lines, size=9, font=F_L, align=PP_ALIGN.LEFT, color=INK, fill=None,
               t=(6350, RULE), b=(6350, RULE), l=(6350, RULE), r=(6350, RULE), marL=54000):
    write(cell.text_frame, lines, size=size, font=font, align=align, color=color)
    tcPr = cell._tc.get_or_add_tcPr()
    for ch in list(tcPr): tcPr.remove(ch)
    tcPr.set("anchor", "ctr"); tcPr.set("marL", str(marL)); tcPr.set("marR", "36000")
    tcPr.set("marT", "9525"); tcPr.set("marB", "9525")
    for tag, (w, c) in (("lnL", l), ("lnR", r), ("lnT", t), ("lnB", b)):
        tcPr.append(ln_el(tag, w, c))
    if fill:
        sf = etree.SubElement(tcPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=fill)
    else:
        etree.SubElement(tcPr, qn("a:noFill"))

def label_lines(label):
    return [(p, {"size": 7.5, "color": GRAY}) if p.startswith("(") else (p, {"font": F_M, "size": 9})
            for p in label.split("\n")]

def grid_tbl(slide, y, rows, widths=(1.55, W0 - 1.55), header=("구분", "주요내용"), x=L0, head_h=0.28, gap=0.04, min_h=0.34):
    """grid()와 같은 입력으로 '한 개의 실제 표'를 만든다. 내장 하위표는 열 경계를 합쳐 같은 표의 행·병합 셀로 구성."""
    LW, CW = widths
    subs = {}
    bnd = [0.0, CW]
    for row in rows:
        for b in row[1]:
            if b[0] == "t":
                ws, hs = sub_dims(b[1], b[2], CW)
                subs[id(b)] = (ws, hs)
                acc = 0.0
                for w in ws[:-1]:
                    acc += w; bnd.append(acc)
    cuts = []
    for v in sorted(bnd):
        if not cuts or v - cuts[-1] > 0.04:
            cuts.append(v)
    cuts[-1] = CW
    near = lambda v: min(range(len(cuts)), key=lambda j: abs(cuts[j] - v))
    cols = [LW] + [cuts[i + 1] - cuts[i] for i in range(len(cuts) - 1)]
    NC = len(cols)
    spec = [("h", None, head_h, -1)]                # (종류, 데이터, 높이, 그룹번호)
    for gi, row in enumerate(rows):
        blocks = row[1]; mh = row[2] if len(row) > 2 else min_h
        items = []
        for b in blocks:
            if b[0] == "b":
                items.append(["b", b, lines_h(b[1], (b[2] if len(b) > 2 else {}).get("size", 9), CW) + 0.10])
            else:
                ws, hs = subs[id(b)]
                for ri, h in enumerate(hs):
                    items.append(["t", (b, ri), h])
        tot = sum(it[2] for it in items)
        if tot < mh:
            tgt = [it for it in items if it[0] == "b"] or items
            for it in tgt:
                it[2] += (mh - tot) / len(tgt)
        spec += [(k, d, h, gi) for k, d, h in items]
    nr = len(spec)
    gf = slide.shapes.add_table(nr, NC, IN(x), IN(y), IN(sum(cols)), IN(sum(sp[2] for sp in spec)))
    tb = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    for i, w in enumerate(cols):
        tb.columns[i].width = IN(w)
    NONE, THIN, DARK = (9525, None), (6350, RULE), (12700, "000000")
    merges = []
    for ri, (kind, data, h, gi) in enumerate(spec):
        tb.rows[ri].height = IN(h)
        first, last = ri == 0, ri == nr - 1
        prev_same = ri > 0 and spec[ri - 1][3] == gi
        next_same = ri < nr - 1 and spec[ri + 1][3] == gi
        top = DARK if first else (NONE if prev_same and kind == "b" and spec[ri - 1][0] == "b" else THIN)
        bot = DARK if last else (NONE if next_same and kind == "b" and spec[ri + 1][0] == "b" else THIN)
        if kind == "h":
            cell_style(tb.cell(ri, 0), [(header[0], {"font": F_B, "size": 9.5})], align=PP_ALIGN.CENTER, fill=MGRAY,
                       t=top, b=bot, l=NONE)
            for ci in range(1, NC):
                cell_style(tb.cell(ri, ci), [(header[1], {"font": F_B, "size": 9.5})] if ci == 1 else [""],
                           align=PP_ALIGN.CENTER, fill=MGRAY, t=top, b=bot, r=NONE if ci == NC - 1 else THIN)
            merges.append((ri, 1, ri, NC - 1))
            continue
        # 구분(라벨) 열 — 그룹 첫 행에 쓰고 그룹 전체 병합
        g_start = ri
        while g_start - 1 > 0 and spec[g_start - 1][3] == gi:
            g_start -= 1
        g_end = ri
        while g_end + 1 < nr and spec[g_end + 1][3] == gi:
            g_end += 1
        if not prev_same:
            if g_end > ri:
                merges.append((ri, 0, g_end, 0))
            lab = rows[gi][0]
        else:
            lab = None
        g_top = THIN
        g_bot = DARK if g_end == nr - 1 else THIN
        cell_style(tb.cell(ri, 0), label_lines(lab) if lab else [""], align=PP_ALIGN.CENTER, fill=LGRAY,
                   t=g_top, b=g_bot, l=NONE)
        if kind == "b":
            b = data; opt = b[2] if len(b) > 2 else {}
            for ci in range(1, NC):
                cell_style(tb.cell(ri, ci), b[1] if ci == 1 else [""], size=opt.get("size", 9), t=top, b=bot, r=NONE,
                           marL=91440)
            if NC > 2:
                merges.append((ri, 1, ri, NC - 1))
            continue
        b, sri = data; opt = b[2]; ws, _ = subs[id(b)]
        srow = b[1][sri]; nsub = len(b[1])
        bounds = [0.0]
        for w in ws:
            bounds.append(bounds[-1] + w)
        size = opt.get("size", 8); aligns = opt.get("al"); tots = opt.get("tot", ())
        fills = opt.get("fills") or {}; colors = opt.get("colors") or {}
        spans = []
        for sci in range(len(ws)):
            c0 = 1 + near(bounds[sci]); c1 = near(bounds[sci + 1])
            spans.append((c0, c1))
            is_head = sri == 0; is_tot = sri in tots
            is_lab = opt.get("label_col") and sci == 0 and not is_head
            fnt = F_B if (is_head or is_tot) else (F_M if is_lab else F_L)
            al = PP_ALIGN.CENTER if (is_head or is_lab) else (aligns[sci] if aligns else PP_ALIGN.LEFT)
            fill = fills.get((sri, sci)) or (MGRAY if is_head else (LGRAY if (is_lab or is_tot) else None))
            val = srow[sci] if sci < len(srow) else ""
            lines = val.split("\n")
            for ci in range(c0, c1 + 1):
                cell_style(tb.cell(ri, ci), lines if ci == c0 else [""], size=size, font=fnt, align=al,
                           color=colors.get((sri, sci), INK), fill=fill, t=top, b=bot,
                           r=NONE if c1 == NC - 1 else THIN)
        row0 = ri - sri
        covered = set()
        for (r1, c1_, r2, c2_) in opt.get("merges", ()):
            if sri == 0:
                merges.append((row0 + r1, 1 + near(bounds[c1_]), row0 + r2, near(bounds[c2_ + 1])))
            covered |= {(rr, cc) for rr in range(r1, r2 + 1) for cc in range(c1_, c2_ + 1)}
        for sci, (c0, c1) in enumerate(spans):
            if c1 > c0 and (sri, sci) not in covered:
                merges.append((ri, c0, ri, c1))
    for (r1, c1, r2, c2) in merges:
        if (r1, c1) != (r2, c2):
            tb.cell(r1, c1).merge(tb.cell(r2, c2))
    return y + sum(sp[2] for sp in spec)

# ── 슬라이드 구조 도구 ─────────────────────────────────────────────────────
def keep_only(slide, ids):
    for sh in list(slide.shapes):
        if sh.shape_id not in ids:
            sh._element.getparent().remove(sh._element)

def clone(src):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst")):
            continue
        new.shapes._spTree.append(copy.deepcopy(el))
    return new

def page(slide, title_id, title, band, foot_id=None, foot=None):
    set_lines(find(slide, title_id), [title])
    band_label(slide, band, 0)
    if foot_id is not None:
        if foot is None:
            sh = find(slide, foot_id); sh._element.getparent().remove(sh._element)
        else:
            set_lines(find(slide, foot_id), foot if isinstance(foot, list) else [foot])

# v10 원본 슬라이드 참조
S3, S4, S5, S6, S7, S8, S9, S10, S11, S12 = (SL[i - 1] for i in range(3, 13))

# 신규 페이지는 v10 슬4(8_흰배경: 띠 1.18") 골격을 복제
T_PG = clone(S4); keep_only(T_PG, {2, 9, 16})      # 경과 및 향후 사업일정
F_PG = clone(S4); keep_only(F_PG, {2, 9, 16})      # 펀딩 현황
L_PG = clone(S4); keep_only(L_PG, {2, 9, 16})      # 매매계약 (3/3) 승계 임대차

# 데이터 보관 (v10 표에서 원문 그대로 읽어 사용)
CAP = [[c.text for c in r.cells] for r in find(S7, 19).table.rows]
foot3 = copy.deepcopy(find(S3, 109)._element)

# ══════════════════════════════════════════════════════════════════════
# 목차
# ══════════════════════════════════════════════════════════════════════
toc = find(SL[1], 9)
items = [("dash", "투자심의위원회 심의 개요"), ("dash", "투자심의위원회 심의 안건"),
         ("sub", "안건 Ⅰ. 매매계약서 체결의 건"), ("sub", "안건 Ⅱ. 주주간계약서 체결의 건"),
         ("sub", "안건 Ⅲ. 코람코라이프인프라리츠 매입확약서(LOC) 날인의 건"),
         ("dash", "참고. 관계사 인수확약서, 참여의향서 등"), ("dash", "참고. 당사 예상수익"),
         ("none", ""), ("none", "[별첨] 투자 대상 물건 설명자료")]
set_lines(toc, [t for _, t in items])
toc.left = IN(toc.left / 914400 - 0.45); toc.width = IN(toc.width / 914400 + 0.45)
for p, (kind, _) in zip(toc.text_frame.paragraphs, items):
    pPr = p._p.get_or_add_pPr()
    if kind != "dash":
        for tag in ("a:buFontTx", "a:buChar"):
            el = pPr.find(qn(tag))
            if el is not None: pPr.remove(el)
        etree.SubElement(pPr, qn("a:buNone"))
        pPr.set("marL", "342900" if kind == "sub" else "0"); pPr.set("indent", "0")

# ══════════════════════════════════════════════════════════════════════
# [개요 1] 경과 및 향후 사업일정 — 2트랙 일정 + 본 투심 안건 콜아웃
# ══════════════════════════════════════════════════════════════════════
s = T_PG
page(s, 9, "> 심의안건 및 추진일정", "경과 및 향후 사업일정", 16, None)
s.shapes._spTree.append(foot3)                       # v10 각주(1·2) 원문 유지

Y1, Y2 = 3.62, 6.02                                  # 트랙 기준선
past = [(0.98, "8월 5일", ["입찰참여"]),
        (1.92, "8월 18일", ["우선협상자", "통지"]),
        (2.95, "8월 24일", ["예비투자심의위원회", "• 양해각서 체결의 건", "• 실사기관 선정의 건"]),
        (3.98, "8월 27일", ["양해각서 체결"]),
        (4.92, "9월 11일", ["발기설립 완료^1)^"]),
        (5.92, "9월 23일", ["영업인가 접수", "선순위 대출심의 완료"])]
NOW = 6.98
fut = [(8.05, "10월 16일", ["• 주주간계약 체결", "• 매매계약 체결", "• 대출약정 체결", "• KLI 이사회"]),
       (9.18, "10월 20일", ["주금 납입"]),
       (10.30, "10월 22일", ["**거래종결**", "(소유권이전)"])]
# 트랙 라벨
box(s, L0, Y1 - 0.80, 2.25, 0.28, [("코람코라이프로지스리츠 (자리츠)", {"size": 8.5, "font": F_M})], line=GRAY)
box(s, L0, Y2 - 0.62, 2.25, 0.28, [("코람코라이프인프라리츠 (모리츠)", {"size": 8.5, "font": F_M})], line=GRAY)
# 과거 구간
line(s, 0.70, Y1, NOW, Y1, color=INK, width=1.0)
shape(s, MSO_SHAPE.OVAL, 0.66, Y1 - 0.04, 0.08, 0.08, fill=INK)
for x, d, desc in past:
    shape(s, MSO_SHAPE.OVAL, x - 0.08, Y1 - 0.08, 0.16, 0.16, fill=LORANGE if "예비" in desc[0] else "FFFFFF", line=INK, lw=1)
    text(s, x - 0.55, Y1 - 0.42, 1.1, 0.24, [(d, {"size": 9})])
    wd = 1.34 if len(desc) > 2 else (1.22 if x > 5 else 1.0)
    text(s, x - wd / 2, Y1 + 0.16, wd, 0.18 * len(desc) + 0.08, desc, size=8 if len(desc) < 3 else 7.5,
         align=PP_ALIGN.LEFT if desc[0].startswith("예비") else PP_ALIGN.CENTER)
# 향후 구간 (빗금 띠)
band = shape(s, MSO_SHAPE.RECTANGLE, NOW + 0.18, Y1 - 0.12, 10.98 - NOW - 0.18, 0.24)
band.fill.patterned(); band.fill.pattern = MSO_PATTERN.LIGHT_VERTICAL
band.fill.fore_color.rgb = RGBColor.from_string(MGRAY); band.fill.back_color.rgb = RGBColor.from_string("FFFFFF")
line(s, 10.98, Y1, 11.20, Y1, color=INK, width=1.0, head=True)
shape(s, MSO_SHAPE.OVAL, NOW - 0.11, Y1 - 0.11, 0.22, 0.22, fill="C55A11", line=INK, lw=1)
text(s, NOW - 0.6, Y1 - 0.44, 1.2, 0.26, [("10월 14일", {"size": 9.5, "font": F_B})])
text(s, NOW - 0.6, Y1 + 0.16, 1.2, 0.40, [("투자심의위원회", {"size": 8.5, "font": F_B}), ("개최", {"size": 8.5, "font": F_B})])
for x, d, desc in fut:
    shape(s, MSO_SHAPE.OVAL, x - 0.15, Y1 - 0.15, 0.30, 0.30, fill="FFFFFF", line=NAVY, lw=2.25)
    text(s, x - 0.55, Y1 - 0.48, 1.1, 0.26, [(d, {"size": 9.5, "font": F_B, "color": NAVY})])
    text(s, x - 0.60, Y1 + 0.20, 1.2, 0.18 * len(desc) + 0.08, desc, size=8,
         align=PP_ALIGN.LEFT if desc[0].startswith("•") else PP_ALIGN.CENTER)
# 기간 표시 (독점협상 / 실사)
line(s, 3.98, 2.48, 3.98, 2.64, color=GRAY, dash=True)
line(s, 3.98, 2.48, 11.10, 2.48, color=GRAY, dash=True)
line(s, 11.10, 2.48, 11.10, Y1 - 0.10, color=GRAY, dash=True, head=True)
box(s, 7.65, 2.37, 3.3, 0.22, [("〈 양해각서 독점적 협상기간 만료 예정: ’26.10.28 〉", {"size": 8})], fill="FFFFFF")
line(s, 3.98, 2.98, 5.92, 2.98, color=GRAY, dash=True)
line(s, 3.98, 2.98, 3.98, 3.10, color=GRAY, dash=True)
line(s, 5.92, 2.98, 5.92, 3.10, color=GRAY, dash=True)
box(s, 4.35, 2.87, 1.20, 0.22, [("〈 실사기간: 4주 〉", {"size": 8})], fill="FFFFFF")
# 본 투심 콜아웃
box(s, 5.20, 1.52, 3.75, 0.84, [("본 투자심의위원회 개최 (10/14)", {"font": F_B, "size": 9.5}),
                                ("1. 매매계약서 체결의 건", {"size": 8.5}),
                                ("2. 주주간계약서 체결의 건", {"size": 8.5}),
                                ("3. 코람코라이프인프라리츠 매입확약서(LOC) 날인의 건", {"size": 8.5})],
    fill="FFF4DC", line=ORANGE, lw=1, kind=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.LEFT)
line(s, NOW, 2.36, NOW, Y1 - 0.13, color="C55A11", width=1.25, head=True)
# KLI 트랙
line(s, NOW, Y1 + 0.58, NOW, Y2 - 0.48, color=GRAY, dash=True)
band2 = shape(s, MSO_SHAPE.RECTANGLE, 5.65, Y2 - 0.12, 10.98 - 5.65, 0.24)
band2.fill.patterned(); band2.fill.pattern = MSO_PATTERN.LIGHT_VERTICAL
band2.fill.fore_color.rgb = RGBColor.from_string(MGRAY); band2.fill.back_color.rgb = RGBColor.from_string("FFFFFF")
line(s, L0 + 2.25, Y2 - 0.48, 5.55, Y2 - 0.48, color=GRAY, dash=True)
line(s, 5.55, Y2 - 0.48, 5.55, Y2 - 0.14, color=GRAY, dash=True, head=True)
kli = [(5.95, "10월 12일(E)", ["투자자문위원회", "• KLL 지분증권 취득 및", "  매입확약서 제출 자문"]),
       (8.05, "10월 16일(E)", ["이사회", "• 주주간계약 체결", "  보통주 148억·2-1종 160억", "• 매입확약서 제출",
                              "• 사모사채(140억원) 발행 및", "  인수계약 체결"]),
       (9.18, "10월 20일", ["주금 납입", "KLI 출자 308억원"]),
       (10.30, "10월 22일 이후", ["변경인가·보고", "[확인 필요]"])]
for x, d, desc in kli:
    shape(s, MSO_SHAPE.OVAL, x - 0.15, Y2 - 0.15, 0.30, 0.30, fill="FFFFFF", line=NAVY, lw=2.25)
    text(s, x - 0.6, Y2 - 0.48, 1.2, 0.26, [(d, {"size": 9, "font": F_B, "color": NAVY})])
    lines_ = [(desc[0], {"font": F_B, "size": 8})] + desc[1:]
    bx, bw = {5.95: (5.25, 1.45), 8.05: (7.05, 1.62), 9.18: (8.68, 1.0), 10.30: (9.85, 1.25)}[x]
    text(s, bx, Y2 + 0.20, bw, 0.16 * len(desc) + 0.08, lines_, size=7,
         align=PP_ALIGN.LEFT if len(desc) > 2 else PP_ALIGN.CENTER)
for x in (8.05, 9.18):                                # KLI → 자리츠 연결
    up = shape(s, MSO_SHAPE.UP_ARROW, x - 0.13, Y1 + 0.98 if x == 8.05 else Y1 + 0.48, 0.26,
               (Y2 - 0.50) - (Y1 + (0.98 if x == 8.05 else 0.48)), fill=LBLUE)
s.shapes._spTree.remove(foot3); s.shapes._spTree.append(foot3)   # 각주를 최상단 레이어로

# ══════════════════════════════════════════════════════════════════════
# [개요 2] 심의안건 + 거래관련 주요 Key-Point (v10 슬3 재사용)
# ══════════════════════════════════════════════════════════════════════
s = S3
keep_only(s, {2, 3, 4})
set_lines(find(s, 2), ["> 심의안건 및 추진일정"])
band_label(s, "심의안건", 0)
ag_rows = [
    ("안건 Ⅰ", "매매계약서 체결의 건\n(매도인: 코크렙안양㈜)",
     ["• **매매대금 2,700억원** (VAT 별도, 감정가 3,345억 대비 −19.3%) · 계약금 100억원 체결일 지급",
      "• 매매목적물: 관양동 912-2·934·934-1 토지·건물 (연면적 95,434.55㎡)",
      "• 10/16 체결 → **10/22 거래종결** (소유권이전 · 임대차 승계)"]),
    ("안건 Ⅱ", "주주간계약서 체결의 건\n(KLI · 종류주식 투자자 · 코크렙안양)",
     ["• **총 출자 988억원**: 보통주 148 · 제1종 240 · 제2종 500 · 제3종 100 (억원)",
      "• 1종 연 7.0% · 2종 연 7.5% 누적배당 / 제3종은 매도인 재투자 (매매대금채권 상계납입)",
      "• 이사 3인 중 2인 · 대표이사 KLI 지명 / 주금 납입 10/20"]),
    ("안건 Ⅲ", "코람코라이프인프라리츠\n매입확약서(LOC) 날인의 건",
     ["• **제2종 340억원(12개월) · 제1종 240억원(24개월)** 발행가액 매입확약 (KLI 보유분 제외)",
      "• 불이행 시 1종 연 5.0% 지연손해금 / 2종 연 3.0% 지연이자 + 위약벌 20%",
      "• 주주간계약 체결일(10/16) 제출 · 종류주주는 매도확약서 제출 (KLI 요청 시 15일 내 매도)"])]
cw = (1.0, 2.25, W0 - 3.25)
NONE, THIN, DARK = (9525, None), (6350, RULE), (12700, "000000")
ag = s.shapes.add_table(4, 3, IN(L0), IN(2.55), IN(W0), IN(0.27 + 0.66 * 3)).table
ag_tblPr = ag._tbl.tblPr
for a in ("firstRow", "bandRow"):
    ag_tblPr.set(a, "0")
for i, w in enumerate(cw):
    ag.columns[i].width = IN(w)
ag.rows[0].height = IN(0.27)
for ci, hd in enumerate(("구분", "안건명", "주요내용")):
    cell_style(ag.cell(0, ci), [(hd, {"font": F_B, "size": 9.5})], align=PP_ALIGN.CENTER, fill=MGRAY,
               t=DARK, b=THIN, l=NONE if ci == 0 else THIN, r=NONE if ci == 2 else THIN)
for i, (g, nm, bl) in enumerate(ag_rows, start=1):
    ag.rows[i].height = IN(0.66)
    bot = DARK if i == 3 else THIN
    nm1, nm2 = nm.split("\n")
    cell_style(ag.cell(i, 0), [(g, {"font": F_M, "size": 9.5})], align=PP_ALIGN.CENTER, fill=LGRAY, b=bot, l=NONE)
    cell_style(ag.cell(i, 1), [(nm1, {"size": 9.5}), (nm2, {"size": 8, "color": GRAY})], align=PP_ALIGN.CENTER, b=bot)
    cell_style(ag.cell(i, 2), bl, size=9, b=bot, r=NONE, marL=91440)
# 거래관련 주요 Key-Point (거래 구조도)
KY = 4.93
frame = shape(s, MSO_SHAPE.RECTANGLE, L0, KY, W0, 7.33 - KY, line=RULE)
box(s, L0 + 0.12, KY - 0.12, 1.85, 0.24, [("▌거래관련 주요 Key-Point", {"font": F_B, "size": 9.5, "color": NAVY})],
    fill="FFFFFF", align=PP_ALIGN.LEFT)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.75, KY + 0.22, 2.20, 2.02, fill="EAF0F8")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 3.85, KY + 0.22, 7.22, 2.02, fill="FBEEE6")
text(s, 0.80, KY + 0.24, 1.2, 0.2, [("매도인", {"size": 8, "color": GRAY})], align=PP_ALIGN.LEFT)
text(s, 9.40, KY + 0.24, 1.6, 0.2, [("리츠 · 투자자", {"size": 8, "color": GRAY})], align=PP_ALIGN.RIGHT)
MY, MH = KY + 0.70, 0.60                              # 주 행 (매도인 · 자리츠 · 모리츠)
box(s, 0.95, MY, 1.80, MH, [("코크렙안양㈜", {"font": F_B, "size": 9.5}), ("(매도인 · PFV)", {"size": 8})],
    fill="FFFFFF", line=INK, lw=1)
box(s, 4.30, MY, 2.35, MH, [("코람코라이프로지스리츠", {"font": F_B, "size": 9.5}), ("(자리츠 · 신규설립)", {"size": 8})],
    fill="FFFFFF", line=INK, lw=1.5)
box(s, 8.60, MY, 2.30, MH, [("코람코라이프인프라리츠", {"font": F_B, "size": 9.5}), ("(상장리츠 · 모리츠)", {"size": 8})],
    fill="FFFFFF", line=INK, lw=1)
box(s, 4.30, KY + 0.24, 2.35, 0.30, [("대주: 선순위 Tr.A 1,600 · 중순위 Tr.B 380", {"size": 7.5})], fill="FFFFFF", line=GRAY)
line(s, 5.47, KY + 0.54, 5.47, MY - 0.01, color=GRAY, head=True)
# 매도인 ↔ 자리츠
line(s, 2.78, MY + 0.18, 4.27, MY + 0.18, color=GRAY, head=True)
line(s, 4.27, MY + 0.42, 2.78, MY + 0.42, color=GRAY, head=True)
text(s, 2.70, MY - 0.12, 1.65, 0.24, [("소유권이전 · 임대차 승계", {"size": 7.5})])
text(s, 2.70, MY + 0.46, 1.65, 0.24, [("매매대금 2,700억원", {"size": 7.5})])
# 제3종 재투자 (상계납입)
line(s, 1.85, MY + MH, 1.85, MY + 1.12, color=GRAY)
line(s, 1.85, MY + 1.12, 4.70, MY + 1.12, color=GRAY)
line(s, 4.70, MY + 1.12, 4.70, MY + MH + 0.02, color=GRAY, head=True)
text(s, 1.95, MY + 1.14, 2.75, 0.22, [("제3종 재투자 100억원 (매매대금채권 상계납입)", {"size": 7.5})])
# 모리츠 ↔ 자리츠
line(s, 8.57, MY + 0.18, 6.68, MY + 0.18, color=GRAY, head=True)
line(s, 6.68, MY + 0.42, 8.57, MY + 0.42, color=GRAY, head=True)
text(s, 6.62, MY - 0.14, 2.0, 0.26, [("보통주 148 + 2-1종 160 출자", {"size": 7.5})])
text(s, 6.62, MY + 0.46, 2.0, 0.24, [("의결권 71.4% · 계열 편입", {"size": 7.5})])
# 종류주식 투자자
FY = MY + 0.90
box(s, 7.15, FY, 3.75, 0.60, [("종류주식 투자자 (FI)", {"font": F_M, "size": 8.5}),
                              ("1종 240: 이지스리츠3호 · 키움 · 애큐온 · MG", {"size": 7.5}),
                              ("2종 340: 삼성증권 · 이지스K리츠 · 기계설비조합", {"size": 7.5})],
    fill="FFFFFF", line=GRAY)
line(s, 7.12, FY + 0.30, 6.30, FY + 0.30, color=GRAY)
line(s, 6.30, FY + 0.30, 6.30, MY + MH + 0.02, color=GRAY, head=True)
text(s, 5.15, FY + 0.33, 1.6, 0.2, [("종류주식 출자", {"size": 7.5})])
line(s, 9.75, MY + MH, 9.75, FY - 0.02, color=ORANGE, dash=True, head=True)
text(s, 9.80, MY + MH + 0.02, 1.25, 0.36, [("매입확약", {"size": 7.5, "color": "C55A11"})], align=PP_ALIGN.LEFT)
set_lines(find(s, 3), ["본 심의 안건은 매매계약·주주간계약 체결(10/16) 및 KLI 매입확약서 제출 건임",
                       "주금 납입(10/20) 후 10/22 자리츠 명의 소유권이전으로 거래종결 예정"])
unit(s, R0 - 0.12, 4.98, "(단위: 억원)")
fn = s.shapes.add_textbox(IN(0.31), IN(7.45), IN(10.95), IN(0.25))
write(fn.text_frame, [f"※ 출처: {SPA}, {SHA}, KLI 제1·2종 매입확약 공문(안), 본 자료 재원조달·지분구성표."],
      size=8, align=PP_ALIGN.LEFT, color=INK)

# ══════════════════════════════════════════════════════════════════════
# [개요 3] 펀딩 현황 (재원별 투자자 · 금액 · 확약 현황)
# ══════════════════════════════════════════════════════════════════════
s = F_PG
page(s, 9, "> 심의안건 및 추진일정", "펀딩 현황", 16,
     f"※ 출처: 본 자료 Ⅱ. 투자계획 '재원조달'·지분구성표, {SHA} 제2.1·2.2·2.5조, 참고 인수확약서·참여의향서. "
     "2종 투자자별 물량은 10/8 최종합의 회람 시 확정 예정(세부 변경 가능).")
unit(s, R0, 1.50, "(단위: 억원)")
rows = [["구분", "트랜치", "투자자", "금액", "확약 현황", "비고"],
        ["담보대출", "선순위 Tr.A", "우리은행", "1,600", "금융확약서(LOC) 확보 (9/21)", "금리 4.78%, 수수료 1.45% (All-in 5.50%) · 24개월"],
        ["", "중순위 Tr.B", "키움캐피탈 · MG캐피탈 등", "380", "LOC 확보", "금리 6.00%, 수수료 1.00% (All-in 6.50%) · 24개월"],
        ["Loan 합계", "", "", "1,980", "", ""],
        ["종류주식", "제1-1종", "이지스리츠3호펀드 · 키움캐피탈\n애큐온캐피탈 · MG캐피탈", "240", "우리투자증권 인수확약서(LOC)\n미납입분 잔액인수 (SHA 제2.5조)", "연 7.0% 누적 · 의결권 · 24개월 내 KLI 매입확약"],
        ["", "제2-1종", "코람코라이프인프라리츠", "160", "KLI 이사회 승인 예정 (10/16)", "연 7.5% 누적 · 의결권"],
        ["", "제2-1종 · 2-2종", "삼성증권 (2-1종 120)\n이지스K리츠펀드 · 기계설비조합 (2-2종 220)", "340", "삼성증권 인수확약서(LOC)\n미납입분 잔액인수 (SHA 제2.5조)", "연 7.5% 누적 · 12개월 내 KLI 매입확약\n2-2종 무의결권"],
        ["", "제3종", "코크렙안양주식회사 (매도인)", "100", "매매대금채권과 상계납입 (SHA 제2.2조)", "무배당 · 무의결권 · 발행가 23,000원"],
        ["보통주", "보통주", "코람코라이프인프라리츠", "148", "KLI 이사회 승인 예정 (10/16)", "발행가 5,000원 · 의결권"],
        ["Equity 합계", "", "", "988", "", ""],
        ["임대보증금", "", "승계 임대보증금 (임차인 질권설정금액 제외)", "21", "-", ""],
        ["재원조달 합계", "", "", "2,989", "", "총 투자비 2,989억원"]]
hs = [0.34, 0.46, 0.46, 0.38, 0.58, 0.46, 0.58, 0.46, 0.46, 0.38, 0.46, 0.40]
fw = [1.15, 1.20, 2.75, 0.75, 2.45, 2.37]
table(s, rows, L0, 1.72, fw, size=8.5, heights=hs,
      aligns=[PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.RIGHT, PP_ALIGN.CENTER, PP_ALIGN.LEFT],
      total_rows=(3, 9, 11), merges=((1, 0, 2, 0), (3, 0, 3, 2), (4, 0, 7, 0), (9, 0, 9, 2), (11, 0, 11, 2)))
xk = L0 + sum(fw[:4])
shape(s, MSO_SHAPE.RECTANGLE, xk, 1.70, fw[4], sum(hs) + 0.04, line="FF0000", lw=2)

# ══════════════════════════════════════════════════════════════════════
# 안건 Ⅰ 매매계약서 주요 내용 (1/3) — v10 슬4 재사용
# ══════════════════════════════════════════════════════════════════════
s = S4
keep_only(s, {2, 9, 16})
page(s, 9, "> 안건 Ⅰ. 매매계약서 체결의 건", "매매계약서 주요 내용 (1/3)", 16,
     f"※ 출처: {SPA} 제1~5조·제9조·별첨1. [*] 항목은 협의 중이며, 연면적은 별첨1(일부멸실 등기 반영) 기준으로 IM·재무모델(95,474.59㎡)과 차이. {LEGAL}")
steps = [["구분", "① 계약 체결", "② 선행조건 충족", "③ 거래종결", "④ 종결 후 의무"],
         ["일정", "10/16", "~10/21", "10/22", "~12/22"],
         ["주요 내용", "계약금 100억원 지급", "확약 이행 · 진술보장 진실\n쿠팡 승계동의 · 우선수익자 동의",
          "잔금 지급 · 소유권이전\n임대차 · 보증금 승계", "진술보장 2개월 · 기업결합 신고 30일\n하자보수청구권 양도"]]
grid_tbl(s, 1.62, [
    ("계약당사자\n(전문)", [("b", ["• 매도인: 코크렙안양㈜ (PFV, 한국투자부동산신탁 담보신탁)",
                                 "• 매수인: 코람코라이프로지스리츠"])]),
    ("매매목적물\n(제1.1조·별첨1)", [("b", ["• 관양동 912-2·934·934-1 토지·건물 (대지 15,287.5㎡ / 연면적 95,434.55㎡), 가설건축물·부속 동산"])]),
    ("매매대금 및\n매매조건\n(제1.1·1.3·1.4·3·4.3조)", [("b", [
        "• 매매대금: **2,700억원** (VAT 별도) — 감정가 3,345억원 대비 −19.3%",
        "• 계약금: **100억원** 체결일(10/16) 지급, 해약금 아님",
        "• 종결시 지급금액: 매매대금 − 계약금 − 승계 임대보증금 ± 정산금, VAT는 매수인 대리납부",
        "• 거래종결일 현황 그대로(as-is, where-is) 매수 · 담보책임 배제 (매도인 고의·중과실은 예외)",
        "▲ 검토: 계약금은 체결일 현금 지급인 반면 3종 주금은 매매대금채권과 상계납입(SHA 제2.2조)",
        "   → 체결일 자리츠 자금 3억원뿐, 상계 방식으로 문안 정비 필요"])]),
    ("거래일정", [("b", ["• 계약체결 10/16, **거래종결 10/22(목)** — 잔금 지급 · 소유권이전 · 임대차 승계"]),
                ("t", steps, {"w": [1.0, 1.8, 2.4, 2.0, 2.4], "rh": 0.30, "size": 8, "label_col": True,
                              "al": [PP_ALIGN.CENTER] * 5})]),
    ("진술 및 보장\n(제4.1·4.4조)", [("b", ["• 설립 · 권한 · 소유권 · 소송 · 임대차 · 경계침범",
                                         "• 진술 및 보장은 **거래종결일부터 2개월** 유효"])]),
    ("신고 · 준거법\n(제5.1(c)(e)·9.8조)", [("b", ["• 체결 후 30일 내 거래신고 · 기업결합 사후신고 (종결 후 30일, 11/21)",
                                                "• 관할: 서울중앙지법"])]),
])

# ══════════════════════════════════════════════════════════════════════
# 안건 Ⅰ 매매계약서 주요 내용 (2/3) — v10 슬5 재사용
# ══════════════════════════════════════════════════════════════════════
s = S5
keep_only(s, {2, 14, 13})
page(s, 14, "> 안건 Ⅰ. 매매계약서 체결의 건", "매매계약서 주요 내용 (2/3)", 13,
     f"※ 출처: {SPA} 제4~9조 (종전 2026-09-07 clean본 대비 계약금·CP·매도인 확약·고의·중과실 예외 추가). [*]는 협의 중. {LEGAL}")
grid_tbl(s, 1.62, [
    ("거래종결\n선행조건\n(제6조)", [("b", [
        "• (공통) 확약 이행 · 진술보장 진실 · 중대한 법규 위반 · 소송 부존재",
        "• (매수인) **쿠팡 승계동의서**, **우선수익자 전원 · 수탁자 동의**",
        "• 미충족 시 청구권을 유보하고 종결 가능 (제6.1조)"])]),
    ("매도인의\n확약사항\n(제5.2조)", [("b", [
        "• 담보신탁 해지 · 대출 상환, 보증금 질권 교체, 하자보수청구권 양도",
        "• 현황측량 · 경계 정리, 옥외광고 허가, 소음민원 조치",
        "▲ 검토: 보증금 질권 교체 예치금 약 50.7억원 자금계획 필요"])]),
    ("손해배상\n(제7조)", [("b", [
        "• 한도 **매매대금 5% (135억원)**, 청구기한 **종결 후 2개월**",
        "• 고의 · 중과실 위반은 한도 · 기간 미적용",
        "▲ 검토: 종결 직후 2개월 내 하자 · 진술 점검"])]),
    ("해제\n(제8조)", [("b", [
        "• 서면합의 · 불가항력 30일 · 중요 위반 10영업일 미시정 · 도산",
        "• 귀책 해제 시 위약금(손해배상 예정), 계약금 이자 가산 반환",
        "▲ 검토: 위약금 · 기간 [*] 공란 — 체결 전 확정 필요"])]),
    ("일반조항\n(제9조)", [("b", [
        "• 양도 제한, 비밀유지 (종결 · 해제 후 1년), 비용 각자 부담",
        "• 양해각서 이행보증금(10억원) 반환 관련 권리 · 의무 존속",
        "• 제7~9조는 해제 후에도 존속 (제9.9조)"])]),
], min_h=0.88)

# ══════════════════════════════════════════════════════════════════════
# 안건 Ⅰ 매매계약서 주요 내용 (3/3) — 승계 임대차 (층별 · 임차인별 보증금 · 임대료)
# ══════════════════════════════════════════════════════════════════════
s = L_PG
page(s, 9, "> 안건 Ⅰ. 매매계약서 체결의 건", "매매계약서 주요 내용 (3/3)", 16,
     f"※ 출처: {SPA} 제5.1(d)·6.1(d)조, 본 자료 Ⅱ. 투자계획 '임차인별 임대차 조건'(매도인 제공). 보증금·월 임대료 금액은 임대면적×평당 단가로 표시"
     "(합계는 동 자료 총액 7,205백만원·1,203백만원과 일치). 임대차 만료는 임차인별 임대개시 시점 기준.")
unit(s, R0, 1.40, "(단위: 평, 백만원, 원/평, %)")
rr = [[c.text for c in r.cells] for r in find(SL[31], 46).table.rows]   # v10 슬32 임대차 조건표
num = lambda v: float(re.sub(r"[^0-9.]", "", v)) if re.sub(r"[^0-9.]", "", v) else None
sub = [["층", "임차인", "용도", "임대면적", "보증금", "월 임대료", "보증금\n(평당)", "임대료\n(평당)", "관리비\n(평당)", "임대차 만료", "인상률"]]
floor = ""
for r in rr[2:-1]:
    floor = r[0] or floor
    a, dep, rent = num(r[3]), num(r[4]), num(r[5])
    exp = r[7].replace("2)", "")
    sub.append([r[0], r[1], r[2], r[3].replace("py", ""),
                f"{a * dep / 1e6:,.1f}" if dep else "-", f"{a * rent / 1e6:,.1f}" if rent else "-",
                r[4].replace("원/평", ""), r[5].replace("원/평", ""), r[6].replace("원/평", ""),
                exp + ("*" if r[1] == "쿠팡" else ""), r[8]])
tot = rr[-1]
sub.append(["합계", "", "", tot[3].replace("py", ""), "7,205.1", "1,203.4", "", "", "", "", ""])
tot_dep = sum(num(r[3]) * num(r[4]) for r in rr[2:-1] if num(r[4]))
tot_rent = sum(num(r[3]) * num(r[5]) for r in rr[2:-1] if num(r[5]))
assert abs(tot_dep / 1e6 - 7205.1) < 0.06 and abs(tot_rent / 1e6 - 1203.4) < 0.06, (tot_dep, tot_rent)
merg = []
start = None
for i in range(1, len(sub) - 1):
    if sub[i][0]:
        if start and i - 1 > start: merg.append((start, 0, i - 1, 0))
        start = i
if start and len(sub) - 2 > start: merg.append((start, 0, len(sub) - 2, 0))
merg.append((len(sub) - 1, 0, len(sub) - 1, 2))
coupang = sum(num(r[3]) for r in rr[2:-1] if r[1] == "쿠팡")
grid_tbl(s, 1.62, [
    ("임대차 승계\n(제5.1(d)·6.1(d)조)", [("b", [
        "• 거래종결 시 임대차 전부 승계 · **임대보증금 반환의무 인수** (종결시 지급금액에서 승계 보증금 차감)",
        "• **쿠팡 승계동의서**는 매수인 선행조건"])]),
    ("임대차\n주요조건", [("b", [
        "• 보증금: **총 7,205백만원** / 월 임대료: **총 1,203백만원** / 월 관리비: 총 62,086천원",
        "• 임대율 99.3%, 국내 최대 이커머스 기업 쿠팡이 전체 임대면적의 77.2% 임차"]),
        ("t", sub, {"w": [0.45, 0.85, 0.75, 0.8, 0.75, 0.75, 0.75, 0.7, 0.65, 1.05, 0.55], "rh": 0.19, "size": 7.5,
                    "label_col": True, "tot": (len(sub) - 1,), "merges": merg,
                    "al": [PP_ALIGN.CENTER] * 3 + [PP_ALIGN.RIGHT] * 6 + [PP_ALIGN.CENTER] * 2})]),
    ("비고", [("b", ["• * 쿠팡: 현 임대차계약 상 만기는 ’34년 4월이나, ’28년 10월까지 임차인이 중도해지(연장) 여부를 선택할 수 있음",
                    "• 현 ’29~’30년 만기 예정 임차인(농심 및 키즈밀 등)을 활용한 추가 임대수익 기대"])]),
])
assert abs(coupang / num(tot[3]) - 0.772) < 0.0006, coupang / num(tot[3])

# ══════════════════════════════════════════════════════════════════════
# 안건 Ⅱ 주주간계약서 주요 내용 (1/3) — 당사자 · 출자금액 · 지분율 (v10 슬6)
# ══════════════════════════════════════════════════════════════════════
s = S6
keep_only(s, {2, 9, 16})
page(s, 9, "> 안건 Ⅱ. 주주간계약서 체결의 건", "주주간계약서 주요 내용 (1/3)", 16,
     f"※ 출처: {SHA} 제1~2조·별첨1, 투자자 송부 메일(10/2). 주주명은 확정, 2종 종류주식 물량은 10/8 최종합의 회람 시 확정 예정(세부 변경 가능). "
     "3종 23,000원×434,783주=10,000,009,000원(단수 BKL 확인 중).")
unit(s, R0, 1.40, "(단위: 백만원, 원, 주, %)")
cap = [r[:] for r in CAP]
cmerge = [(1, 0, 4, 0), (5, 0, 6, 0), (7, 0, 8, 0), (len(cap) - 1, 0, len(cap) - 1, 1)]
grid_tbl(s, 1.62, [
    ("대상회사", [("b", ["• 코람코라이프로지스리츠 (자리츠) — 2026.9.11 자본금 3억원 발기설립 (KLI 100% 출자)"])]),
    ("주주 당사자", [("b", [
        "• 보통주: 코람코라이프인프라리츠",
        "• 제1-1종: 이지스리츠3호펀드 · 키움캐피탈 · 애큐온캐피탈 · MG캐피탈",
        "• 제2-1종: 코람코라이프인프라리츠 · 삼성증권 / 제2-2종(무의결): 이지스K리츠펀드 · 기계설비조합",
        "• 제3종(무의결): 코크렙안양주식회사 (매도인)"])]),
    ("주주 당사자 별\n출자금액 및 지분율\n(제2.1·2.4조)", [
        ("b", ["• 총 출자금액: **988억원** (무의결권 1,314,783주 = 발행주식 20.7%, 상법 §344조의3 1/4 이내)"]),
        ("t", cap, {"w": [1.25, 1.75, 0.5, 0.85, 0.75, 0.85, 0.65, 0.95, 0.85], "rh": 0.205, "size": 7.5,
                    "label_col": True, "tot": (len(cap) - 1,), "merges": cmerge,
                    "al": [PP_ALIGN.CENTER] * 3 + [PP_ALIGN.RIGHT] * 6})]),
    ("납입예정일\n(제2.1·2.2조)", [("b", ["• 거래종결 2영업일 전 이사회가 정한 **납입예정일(10/20) 15:00**까지 · 3종은 매매대금채권과 상계납입"])]),
    ("발기주식 감자 ·\n잔액인수\n(제2.3·2.5조)", [("b", [
        "• KLI 발기주식 30만주(3억원) 100% 유상감자",
        "• 미납입분은 **우리투자증권(1종) · 삼성증권(2종) 잔액인수**"])]),
])

# ══════════════════════════════════════════════════════════════════════
# 안건 Ⅱ 주주간계약서 주요 내용 (2/3) — 배당 · 확약 · 지배구조 · 양도 (v10 슬7)
# ══════════════════════════════════════════════════════════════════════
s = S7
keep_only(s, {2, 9, 16})
page(s, 9, "> 안건 Ⅱ. 주주간계약서 체결의 건", "주주간계약서 주요 내용 (2/3)", 16,
     f"※ 출처: {SHA} 제1.3·1.4·3~7조, 투자자 송부 메일(10/2). {LEGAL}")
grid_tbl(s, 1.62, [
    ("배당\n(제1.3조)", [("b", [
        "• **1종 연 7.0% · 2종 연 7.5% 누적** 우선배당",
        "• 보통주 초기 1년 연 7.5% 배당 협조 (배당가능재원 · 결의 없으면 의무 없음)"])], 0.80),
    ("KLI 매입확약\n(제3·4조)", [("b", [
        "• 2종 340억원(12개월, KLI 보유분 제외) · 1종 240억원(24개월) **발행가액 매입**",
        "• 매입완료일에 미수령 누적배당 정산 (세부는 안건 Ⅲ)"])], 0.80),
    ("제3종 조건\n(제5조)", [("b", [
        "• 발행 후 **24+2개월 내 유상감자** · 미감자 시 연 6% → 8% 누적우선배당",
        "• 48+2개월 도과 후 **자산매각청구권** (세부는 다음 페이지)"])], 0.80),
    ("이사회 구성\n(제6조)", [("b", ["• 이사 3인 중 **2인 · 대표이사 KLI 지명**, 감사는 주주총회 선임"])], 0.66),
    ("자산관리회사\n(제1.4조)", [("b", ["• 동일 AMC(코람코) 이해상충 고지"])], 0.66),
    ("양도 제한\n(제7조)", [("b", [
        "• 다른 주주 **전원의 서면동의** 원칙",
        "• 셀다운은 매도확약서 제출 + KLI 사전동의 (우리투자증권 · 삼성증권 최초 셀다운 면제)"])], 0.80),
])

# ══════════════════════════════════════════════════════════════════════
# 안건 Ⅱ 주주간계약서 주요 내용 (3/3) — 제3종 조건 · 배당 순위 (v10 슬8)
# ══════════════════════════════════════════════════════════════════════
s = S8
keep_only(s, {2, 9, 16})
page(s, 9, "> 안건 Ⅱ. 주주간계약서 체결의 건", "주주간계약서 주요 내용 (3/3)", 16,
     f"※ 출처: {SHA} 제1.3·5.1~5.5조, KLL리츠 정관 v8 제11조, 투자자 송부 메일(10/2). {LEGAL}")
tl = [["’26.10 발행", "’28.12 감자약정기한", "미감자 1년차", "미감자 2년차", "’30.12 추가기한"],
      ["100억원 · 23,000원\n무배당 · 무의결권 · 상계납입", "발행 후 24+2개월\n발행가액 유상감자", "연 6% 누적우선배당\n(배당가능재원 내)",
       "연 8% 누적우선배당\n미배당분은 별도 협의", "48+2개월 도과 시\n자산매각청구권 (매각 완료 시까지)"]]
tl_fill = {(0, 0): LGRAY, (0, 1): LBLUE, (0, 2): LORANGE, (0, 3): "F8CBAD", (0, 4): ORANGE}
rank = [["구분", "①", "②", "③", "④", "⑤", "⑥"],
        ["평상시", "1종 7.0% 누적", "2종 7.5% 누적", "3종 가산 6→8%*", "보통주 잔여", "", ""],
        ["매각 · 청산\n(개정 예정)", "1종 배당", "2종 배당", "1종 원본", "2종 원본", "3종 배당 · 원본", "보통주 원본 · 잔여"],
        ["매각 · 청산\n(현행 정관 v8)", "배당 (①~⑥)", "1종 원본", "2종 원본", "3종 원본", "보통주 원본", "잔여 20% 2종 · 80% 보통"]]
grid_tbl(s, 1.62, [
    ("제3종\n감자 · 가산배당\n(제5.1~5.4조)", [("t", tl, {"w": [1] * 5, "rh": 0.36, "size": 8, "fills": tl_fill,
                                                    "al": [PP_ALIGN.CENTER] * 5,
                                                    "colors": {(0, 4): "FFFFFF"}})]),
    ("정관 개정\n(제5.5조)", [("b", ["• 제5.3조 권리를 **거래종결일부터 3개월 내** 정관에 반영 (변경인가 지연 시 연장) · 지급 보장 · 보전 의무 아님"])]),
    ("배당 순위\n(정관 제11조)\n(SHA 제1.3·5.3조)", [
        ("t", rank, {"w": [1.3, 1.2, 1.2, 1.2, 1.2, 1.25, 1.45], "rh": 0.34, "size": 8, "label_col": True,
                     "al": [PP_ALIGN.CENTER] * 7}),
        ("b", ["• * 3종 가산배당은 감자약정기한 도과 시에만 발생 · 매각 · 청산 순위는 1 · 2종 권리 역전이 없도록 개정 시 반영 예정"],
         {"size": 8})]),
    ("2종 권리 영향\n(투자자 안내 기준)", [("b", [
        "• Exit 시점: 2종은 거래종결 후 1년 내 KLI가 매입확약 → 감자약정기한(26개월) 이후 페널티 발생 시점엔 이미 Exit",
        "• 배당 순위: KLI 사유로 2종이 남더라도 3종 가산배당은 2종 배당 후순위 → 워터폴상 2종 배당 영향 없음",
        "• 매각 시: 추가기한 도과로 자산매각청구권 행사 시에도 잔여재산 분배에서 2종 지위 불변",
        "• 기타: 미배당분 처리 별도 협의(제5.3조), 매각청구는 매각 완료 시까지 가능(제5.4조), 매도기한 15일 통일"])]),
])

# ══════════════════════════════════════════════════════════════════════
# 안건 Ⅲ 매입확약서 주요 내용 (1/1) — v10 슬9
# ══════════════════════════════════════════════════════════════════════
s = S9
keep_only(s, {2, 9, 16})
page(s, 9, "> 안건 Ⅲ. 코람코라이프인프라리츠 매입확약서(LOC) 날인의 건", "매입확약서 주요 내용 (1/1)", 16,
     f"※ 출처: KLI 제1·2종 매입확약 공문(안), 매도확약 공문(안), {SHA} 제2.5·3·4조, 투자자 송부 메일(10/2, 확약서 발행회사 기명날인·매도기한 15일 통일). "
     "본 검토는 내부 검토이며 외부 법무법인 확인 필요.")
unit(s, R0, 1.40, "(단위: 억원)")
ms = [["구분", "’26.10.22 거래종결", "’27.10 (12개월)", "’28.10 (24개월)", "KLI 보유 합계"],
      ["KLI 취득", "보통주 148 + 2-1종 160\nKLI 출자 308", "2종 340 매입\n(삼성 120 · 이지스K 100 · 기계설비 120)",
       "1종 240 매입\n(이지스3호 · 키움 · 애큐온 · MG)", "888 = 보통주 148\n+ 1종 240 + 2종 500"]]
cmp_ = [["구분", "제1종 매입확약 (KLI → 종류주주)", "제2종 매입확약 (KLI → 종류주주)", "매도확약 (종류주주 → KLI)"],
        ["대상 · 금액", "1-1종 960,000주 · 240억원", "2종 1,360,000주 · 340억원\n(KLI 보유 2-1종 제외)", "보유 종류주식 전부\n(셀다운 양수인 포함)"],
        ["기한", "발행일부터 24개월 응당일", "발행일부터 12개월 응당일", "KLI 서면 요청일부터 15일 내"],
        ["불이행 시", "미매입잔액 연 5.0% 지연손해금\n매입의무 존속", "연 3.0% 지연이자 + 위약벌 20%\n매입의무 존속", "1종 연 5.0% / 2종 연 3.0%+20%\n매도의무 존속"],
        ["선행조건", "영업인가 · 소유권 취득, 완전한 소유권 · 부담 부존재\n매도확약서 제출", "(좌 동)", "KLI의 보유주식 전부 매도 요청"],
        ["배당 정산 · 제출", "매입완료일에 미수령 누적배당 정산\n주주간계약 체결일(10/16) 제출", "(좌 동)", "인수인 체결일 / 셀다운 양수인 완결 전"]]
grid(s, 1.62, [
    ("확약 당사자", [("b", ["• 매입확약: **코람코라이프인프라리츠 → 제1·2종 종류주주** (매입확약서 날인 · 제출)",
                         "• 매도확약: 종류주주 → KLI (셀다운 양수인 포함)"])]),
    ("KLI 지분 확대\n일정", [("t", ms, {"w": [0.9, 1.9, 2.75, 2.2, 1.75], "rh": 0.46, "size": 8, "label_col": True,
                                    "al": [PP_ALIGN.CENTER] * 5,
                                    "fills": {(0, 4): "FCE4D6"}}),
                          ("b", ["• KLI 지분: 거래종결 시 의결권 71.4% → 1년차 2종 매입 → 2년차 1종 매입 · 3종은 KLI 매입 대상 아님 (제5조 유상감자로 회수)"],
                           {"size": 8.5})]),
    ("매입 · 매도확약\n주요 조건\n(제3·4조)", [("t", cmp_, {"w": [1.15, 2.75, 2.65, 2.45], "rh": 0.40, "size": 8, "label_col": True,
                                                 "al": [PP_ALIGN.CENTER] + [PP_ALIGN.LEFT] * 3})]),
])

# ══════════════════════════════════════════════════════════════════════
# 참고. 인수확약서 · 참여의향서 — 이미지 축소 + 요약표 (v10 슬10·11)
# ══════════════════════════════════════════════════════════════════════
def shrink(slide, sids, k, y0=2.2, anchor=None):
    pics = [find(slide, i) for i in sids]
    ax = anchor if anchor is not None else min(p.left for p in pics) / 914400
    for p in pics:
        l, t, w, h = (v / 914400 for v in (p.left, p.top, p.width, p.height))
        p.left, p.top, p.width, p.height = IN(ax + (l - ax) * k), IN(y0 + (t - y0) * k), IN(w * k), IN(h * k)

s = S10
shrink(s, [17], 0.72, anchor=1.06 + 3.44 / 2); shrink(s, [20], 0.72, anchor=6.44 + 3.48 / 2)
unit(s, R0, 5.72, "(단위: 억원)")
table(s, [["구분", "트랜치", "확약기관", "배정 금액", "확약서 주요내용"],
          ["우선주", "제2종", "삼성증권㈜", "340", "• 제2종 종류주식 인수 확약 · 미납입분 **잔액인수** (SHA 제2.5조) — KLI 보유 2-1종 160억원 별도"],
          ["", "제1종", "우리투자증권㈜", "240", "• 제1-1종 종류주식 인수 확약 · 미납입분 **잔액인수** (SHA 제2.5조)"]],
      L0, 5.95, [1.1, 1.0, 1.5, 1.0, W0 - 4.6], heights=[0.30, 0.48, 0.48], size=8.5,
      aligns=[PP_ALIGN.CENTER] * 3 + [PP_ALIGN.RIGHT, PP_ALIGN.LEFT], merges=((1, 0, 2, 0),))
fn = s.shapes.add_textbox(IN(0.31), IN(7.45), IN(10.95), IN(0.25))
write(fn.text_frame, [f"※ 출처: 각 사 인수확약서(LOC, 2026-07-31), {SHA} 제2.1·2.5조. 배정 금액은 주주간계약 지분구성표 기준."],
      size=8, align=PP_ALIGN.LEFT)

s = S11
shrink(s, [15], 0.72, anchor=0.66 + 3.38 / 2); shrink(s, [20, 22], 0.72, anchor=5.85)
for sid in (20, 22):                               # 좌측 기준 정렬 보정
    p_ = find(s, sid); p_.left = IN(p_.left / 914400 + 0.55)
unit(s, R0, 5.72, "(단위: 억원)")
table(s, [["구분", "트랜치", "확약기관", "금액", "확약서 주요내용"],
          ["담보대출", "선순위 Tr.A", "우리은행", "1,600", "• **금융확약서(LOC) 확보 (9/21)** · 금리 4.78%, 수수료 1.45% (All-in 5.50%) · 24개월"],
          ["", "중순위 Tr.B", "키움캐피탈 · MG캐피탈 등", "380", "• **LOC 확보** · 금리 6.00%, 수수료 1.00% (All-in 6.50%) · 24개월"]],
      L0, 5.95, [1.1, 1.0, 1.9, 0.8, W0 - 4.8], heights=[0.30, 0.48, 0.48], size=8.5,
      aligns=[PP_ALIGN.CENTER] * 3 + [PP_ALIGN.RIGHT, PP_ALIGN.LEFT], merges=((1, 0, 2, 0),))
fn = s.shapes.add_textbox(IN(0.31), IN(7.45), IN(10.95), IN(0.25))
write(fn.text_frame, ["※ 출처: 우리은행 금융확약서(2026-09-21), 키움캐피탈·MG캐피탈 확약서, 본 자료 Ⅱ. 투자계획 '재원조달'."],
      size=8, align=PP_ALIGN.LEFT)

# ══════════════════════════════════════════════════════════════════════
# 참고. 당사 예상수익 — 표준양식 표 구성 (v10 슬12, 수치 동일)
# ══════════════════════════════════════════════════════════════════════
s = S12
old = [[c.text for c in r.cells] for r in find(s, 9).table.rows]
keep_only(s, {2, 8, 9005})
rows = [["구 분", "", "금 액", "비고"],
        ["운용수익", "매입보수", "13.5억원", "• 매입금액 2,700억원^1)^의 **0.5%**"],
        ["", "운용보수", "66.7억원", "• 2,989억원 (매입장부금액의 **0.223%/연**) x 운용기간 약 10년"],
        ["", "고정보수(소계)", "80.2억원", ""],
        ["", "매각기본보수", "18.8억원", "• T. Cap.Rate 4.89%, 사업기간 10년 가정\n• 매각금액의 **0.5%**"],
        ["", "매각성과보수", "124.8억원", "• T. Cap.Rate 4.89%, 사업기간 10년 가정\n• 매각기본보수 차감후 이익(매각금액-매입장부금액-매각부대비용)의 **15%**"],
        ["", "합계", "223.8억원", ""]]
for r_old, r_new in zip(old[1:], rows[1:]):        # v10 수치 그대로인지 확인
    assert r_old[1] == r_new[1] and r_old[2] == r_new[2], (r_old, r_new)
t = table(s, rows, L0, 1.75, [1.6, 1.6, 1.45, W0 - 4.65], heights=[0.34, 0.42, 0.42, 0.40, 0.58, 0.58, 0.42], size=9.5,
          aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.LEFT], total_rows=(3, 6),
          merges=((0, 0, 0, 1), (1, 0, 6, 0)))
for ri in (3, 6):
    t.cell(ri, 1).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
box(s, L0, 1.75 + 3.16 + 0.05, W0, 0.25, [("1)  취득부대비용을 제외한 매입 금액", {"size": 8.5})], align=PP_ALIGN.LEFT)
band_label(s, "당사 예상수익", 0)

# ══════════════════════════════════════════════════════════════════════
# 슬라이드 순서 재배치
# ══════════════════════════════════════════════════════════════════════
order = [SL[0], SL[1], T_PG, S3, F_PG, S4, S5, L_PG, S6, S7, S8, S9, S10, S11, S12] + SL[12:]
lst = prs.slides._sldIdLst
by_part = {prs.part.related_part(e.rId): e for e in lst}
for e in list(lst):
    lst.remove(e)
for sl in order:
    lst.append(by_part[sl.part])

# 세로 중간 정렬 (사내 표준) — 개요 섹션 전체
for sl in order[1:15]:
    for el in sl.shapes._spTree.iter(qn("a:bodyPr")):
        el.set("anchor", "ctr")
    for el in sl.shapes._spTree.iter(qn("a:tcPr")):
        el.set("anchor", "ctr")

prs.save(str(OUT))
print("saved", OUT, len(prs.slides), "slides")
