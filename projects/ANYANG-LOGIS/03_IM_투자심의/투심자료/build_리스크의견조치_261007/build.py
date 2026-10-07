# -*- coding: utf-8 -*-
"""예비투심 리스크관리부서 의견 → [의견]-[조치사항] 1페이지 표 (사내 표준 양식).
사용법: python3 -I build.py <출력.pptx>
"""
import re
import sys
import pathlib

ROOT = pathlib.Path(__file__).resolve().parents[5]
sys.path.insert(0, str(ROOT / "platform/scripts"))
from ic_deck import ICDeck, F_LIGHT, F_BOLD, F_MEDIUM, INK, GRAY, RED, NAVY_DARK, TH_BG, WHITE, LINE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402

OUT = sys.argv[1]
SUB_BG = RGBColor(0xF2, 0xF2, 0xF2)

# (번호, 구분, 항목, [의견 불릿], [조치사항 불릿])  — **굵게**
ROWS = [
    ("1", "AMC", "거래구조상\n이해상충",
     ["동일 AMC가 매도인·매수인·KLI(지분 보유)를 모두 운용해 양측 투자자 이해가 상충할 수 있는 구조",
      "사전 식별·독립 관리의 절차적 증빙 필요"],
     ["매도인 **경쟁입찰**로 우선협상대상자 선정 → 낙찰가 2,700억은 정상가격으로 해석될 가능성 높음(BKL 법률실사)",
      "부투법 §22조의2③ AMC 특별관계자 거래 → **KLL 이사회 승인·주총 특별결의**(10/16)로 절차 요건 충족",
      "KLI는 3종 매입확약에서 제외(이해관계자 거래 회피), 주주간계약에 이해상충 고지 조항(§1.4) 반영"]),
    ("2", "AMC", "거래조건의\n공정성 및\n충실의무",
     ["매매가격·투자조건의 양측 공정성을 외부 가치평가·시장비교로 입증 필요",
      "매도인 측 3종 100억(배당률 0%) 재투자는 관계사 부당지원으로 평가될 가능성",
      "시장성·외부 법률검토·이사회 판단 근거 문서화 필요"],
     ["매매가 2,700억 = 감정가 3,345억(미래새한 '26.09) 대비 **−19.3%**",
      "3종 발행가 **23,000원**/주 ≥ 상증세법 평가액 **22,723원**/주(회계법인 성지 '26.10.7) → 시가 이상 발행",
      "태평양('26.9.14): 투자금 조기 회수·2년 후 원금 회수 등 합리적 근거가 있으면 **부당지원 가능성 낮음**, 3종 조건은 입찰제안서에 기재돼 형식적 입찰 리스크도 낮음",
      "매도인 측 법률(세종)·세무(성지)도 리스크 거의 없음, 미감자 시 연 6%→8% 누적배당(주주간계약 §5.3·5.4)"]),
    ("3", "AMC", "의사결정절차\n및 사후책임",
     ["이해관계 이사·위원의 의결 참여 제한, 리츠별 이사회·주총 승인 절차 점검",
      "검토가정·반대의견·승인조건 의사록 기재"],
     ["KLL 이사회·주총 특별결의, KLI 이사회 결의(10/16). 위 법률의견·평가보고서를 부의자료로 첨부하고 승인조건을 의사록에 기재 예정",
      "이해관계 이사 의결 제한 여부 [BKL 확인 필요]"]),
    ("4", "자리츠", "쿠팡\n퇴거 리스크",
     ["의무임대기간 2029.4 만료, 연장 여부 불확실",
      "면적 77.2%·명목임대료 79.7%(연 약 115억)의 핵심 임차인, 미연장 시 리파이낸싱 불확실"],
     ["쿠팡 시설투자 **91억**: 5층 메자닌·컨베이어·분류기·스파이럴, 6층 선반랙·컨베이어 등(매도인 회신)",
      "5~6층 슬래브 개구·보 보강 대수선('25.2)으로 층간 수직동선 구축(CBRE 물리실사 p.8·19·35) → 이전 부담이 커 **퇴거 가능성 높지 않음**",
      "의무임대차 만료 **6개월 전**(2028.10.31)까지 잔여 60개월(~2034.4) 전체 사용을 통지하지 않으면 **Rent-Free 5개월 미제공**(원계약 별첨5 특약 §3·6)"]),
    ("5", "자리츠", "2029년\n유동성 리스크",
     ["쿠팡 연장 시 Rent-Free 5개월로 임대수입 약 55.3억 감소(해당 연도 NOI 약 92억)",
      "Over funding에 의존한 이자·배당 지급, 금융구조 취약"],
     []),
    ("6", "자리츠", "재차입 및\n매각가치 의존",
     ["'28·'31년 2,130억 재차입(4.8%·4.5%) 가정은 금번 5.69% 대비 낙관적",
      "KLI 보통주 수익률 0.78~12.6%로 매각가치에 의존"],
     []),
]


def write(cell, items, *, size=8.5, bullet=True, align=PP_ALIGN.LEFT, color=INK, font=F_LIGHT):
    tf = cell.text_frame
    tf.word_wrap = True
    # 기존 문단 정리
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, it in enumerate(items):
        p = p0 if i == 0 else tf.add_paragraph()
        p.alignment = align
        if bullet:
            p.space_after = Pt(1)
        parts = re.split(r"(\*\*[^*]+\*\*)", ("• " if bullet else "") + it)
        for part in parts:
            if not part:
                continue
            b = part.startswith("**")
            r = p.add_run()
            r.text = part.strip("*") if b else part
            r.font.name = F_BOLD if b else font
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.color.rgb = color


d = ICDeck(asset="안양물류센터")
s = d.content_slide(
    "심의안건", "예비투자심의 리스크관리부서 의견 및 조치사항",
    lead="리스크관리부서 의견 6건 중 4건은 법률의견·외부평가·임대차 조건으로 조치함",
    subs=["3종 발행가(23,000원)는 상증세법 평가액(22,723원) 이상, 쿠팡은 91억 시설투자와 Rent-Free 조건으로 잔류 유인 확보",
          "2029년 유동성·재차입 2건은 조치사항 작성 예정 (KLI 모리츠 리스크 검토는 본 장표에서 제외)"])
# 제목 위치: v12 '투자심의위원회 개요' 레이아웃 장표와 동일 (lIns 2.80", tIns 0.39", 검정)
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith(">"):
        bp = sh.text_frame._txBody.bodyPr
        bp.set("lIns", "2556000"); bp.set("tIns", "359911")
        for r in sh.text_frame.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
        break
d.band(s, "리스크관리부서 의견 및 조치사항")

hdr = ["No.", "구분", "항목", "리스크관리부서 의견", "조치사항"]
widths = [0.36, 0.58, 1.00, 3.10, 5.66]
rows = [hdr] + [["", "", "", "", ""] for _ in ROWS]
tbl = d.table(s, rows, col_widths=widths, row_height=0.30, size=9, zebra=False)
heights = [0.26, 0.68, 0.92, 0.50, 0.90, 0.50, 0.52]
for i, h in enumerate(heights):
    tbl.rows[i].height = Inches(h)

for ci, h in enumerate(hdr):
    write(tbl.cell(0, ci), [h], size=9, bullet=False, align=PP_ALIGN.CENTER, font=F_BOLD)
for ri, (no, gb, item, op, act) in enumerate(ROWS, start=1):
    for ci in range(5):
        c = tbl.cell(ri, ci)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.fill.solid()
        c.fill.fore_color.rgb = SUB_BG if ci < 3 else WHITE
    write(tbl.cell(ri, 0), [no], size=9, bullet=False, align=PP_ALIGN.CENTER, font=F_BOLD, color=NAVY_DARK)
    write(tbl.cell(ri, 1), [gb], size=8.5, bullet=False, align=PP_ALIGN.CENTER)
    write(tbl.cell(ri, 2), item.split("\n"), size=8.5, bullet=False, align=PP_ALIGN.CENTER, font=F_MEDIUM)
    write(tbl.cell(ri, 3), op, size=8)
    write(tbl.cell(ri, 4), act if act else [""], size=8, bullet=bool(act))

# 구분 열 세로 병합 (AMC 1~3, 자리츠 4~6)
for a, b in ((1, 3), (4, 6)):
    tbl.cell(a, 1).merge(tbl.cell(b, 1))
    write(tbl.cell(a, 1), [ROWS[a - 1][1]], size=8.5, bullet=False, align=PP_ALIGN.CENTER)
    tbl.cell(a, 1).vertical_anchor = MSO_ANCHOR.MIDDLE

d.footnote(s, "※ 의견: 리스크관리부서 「안양물류센터 매입관련 리스크검토의견」 Ⅰ·Ⅲ(Ⅱ. KLI 리스크검토는 제외). 조치 출처: BKL 법률실사('26.9.9) p.15~16·20, "
              "태평양 「신설리츠 투자구조 관련 검토」('26.9.14) p.3~4·10~12, 율촌 의견서('26.4) p.5, 매도인 보고자료('26.9.30), 회계법인 성지 상증세법 평가보고서('26.10.7, "
              "평가기준일 10/31 가정 — 기준 일정상 거래종결 10/22), 주주간계약 BKL v30, 임차인 시설투자 회신('26.9.29), CBRE 물리실사('26.8.31), "
              "매도인 임차인별 계약조건표(쿠팡 원계약 별첨5 특약 원문 대조 필요). 법률 관련 사항은 내부 검토이며 외부 법무법인 확인 필요. "
              "상기 내용 중 일부는 협의 과정에서 변동될 수 있음.", top=7.00, size=7)
d.save(OUT)
print("saved", OUT)
