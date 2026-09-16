#!/usr/bin/env python3
"""안양물류센터 본 투자심의 자료 v06 — 전 수치를 기준 재무모델(2026-09-15 인가ver)에 정합.

기준 파일:
    04_Equity투자자/models/20260915_재무모델_Financial_Modeling_인가ver_v02.xlsm
    → 'A&R' 탭(가정·결과)과 'IM' 탭(투자자별 수익률·NOI 추이)의 값만 인용한다.

v05 대비 변경
 1. 운영가정(2/2) 운영비용·리츠비용·Initial NOI 표를 A&R/IM 탭 값으로 교체
 2. NOI 추이(10년+매각1개년) 표를 IM 탭 B4:M20으로 교체
 3. 투자자별 수익률 4표(1·2·3종·보통주)를 IM 탭 B25:N51 / B77:N103으로 교체,
    좌우 밴드 라벨 뒤바뀜(1·2종 ↔ 3종·보통주) 수정
 4. Sell-down 2p를 Sell-down 탭(신규 파일) 값으로 갱신
 5. DSCR 표를 A&R V67:AI81으로 교체
 6. 민감도 3표(임대료·공실률·금리)를 A&R AV60:BC89로 교체
 7. 보수·수익률 요약, LTV, 감정가, 쿠팡 비중 등 본문 수치 정합
 8. 재무모델과 IM의 수익률 불일치 주석 삭제(신규 모델에서 9.24%/16.44%로 일치)

실행: python3 build_deck_v6.py
"""
import pathlib
import sys

import openpyxl
from pptx import Presentation

HERE = pathlib.Path(__file__).parent
sys.path.insert(0, str(HERE))
from build_deck_v2 import (iter_shapes, by_text, by_y, set_tf, set_cell, replace_runs, fn_add)  # noqa: E402
from build_deck_v5 import normalize_tables, autofit_tables, lead2  # noqa: E402

PROJ = HERE.parent.parent
MODEL = PROJ / "04_Equity투자자/models/20260915_재무모델_Financial_Modeling_인가ver_v02.xlsm"
SRC = HERE / "20260916_투심자료_ANYANG-LOGIS_v05.pptx"
OUT = HERE / "20260916_투심자료_ANYANG-LOGIS_v06.pptx"
MODEL_TAG = "재무모델 2026-09-15 인가ver(v02)"


# ══════════════════════════════════════════════════════════════════════════
# 0. 서식 헬퍼
# ══════════════════════════════════════════════════════════════════════════
def n0(v, dash=True):
    """백만원 정수. 음수는 회계식 괄호, 0/None은 '-'."""
    if v is None:
        return "-" if dash else ""
    if abs(v) < 0.5:
        return "-" if dash else "0"
    return f"({abs(v):,.0f})" if v < 0 else f"{v:,.0f}"


def pc(v, nd=2, dash=False):
    if v is None:
        return "-" if dash else ""
    return f"{v:.{nd}%}"


def num(v, nd=2):
    return "-" if v is None else f"{v:,.{nd}f}"


# ══════════════════════════════════════════════════════════════════════════
# 1. 기준 재무모델 판독
# ══════════════════════════════════════════════════════════════════════════
class Model:
    def __init__(self, path):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        self.ar = {}
        for r in wb["A&R"].iter_rows(min_row=1, max_row=139):
            for c in r:
                if c.value is not None:
                    self.ar[c.coordinate] = c.value
        self.im = {}
        for r in wb["IM"].iter_rows(min_row=1, max_row=169):
            for c in r:
                if c.value is not None:
                    self.im[c.coordinate] = c.value
        self.sd = {}
        for r in wb["Sell-down"].iter_rows(min_row=60, max_row=93, min_col=12, max_col=30):
            for c in r:
                if c.value is not None:
                    self.sd[c.coordinate] = c.value
        wb.close()

    # 셀 접근 (없으면 0)
    def A(self, ref):
        return self.ar.get(ref, 0)

    def I(self, ref):
        return self.im.get(ref, 0)

    def S(self, ref):
        return self.sd.get(ref, 0)

    def Arow(self, row, cols):
        return [self.ar.get(f"{c}{row}", 0) for c in cols]

    def Irow(self, row, cols):
        return [self.im.get(f"{c}{row}", 0) for c in cols]


AR_YR = ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M"]      # IM 탭 1~10년차 + 매각1개년
AR_DS = ["Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI"]  # A&R DSCR Y1~Y10
SENS7 = ["AW", "AX", "AY", "AZ", "BA", "BB", "BC"]                   # 민감도 7열
SENS6 = ["AX", "AY", "AZ", "BA", "BB", "BC"]                         # 금리 민감도 6열


# ══════════════════════════════════════════════════════════════════════════
# 2. 표 채우기 유틸
# ══════════════════════════════════════════════════════════════════════════
def tables(slide):
    return sorted([sh for sh in slide.shapes if sh.has_table], key=lambda x: (round(x.top / 914400, 1), x.left))


def put(tbl, i, j, text):
    if i < len(tbl.rows) and j < len(tbl.columns):
        set_cell(tbl.cell(i, j), text)


def put_row(tbl, i, vals, j0=0):
    for k, v in enumerate(vals):
        put(tbl, i, j0 + k, v)


def set_band_labels(slide, left_txt, right_txt):
    """슬라이드의 좌·우 밴드 그룹 라벨을 위치 기준으로 교체 (제목 밴드 제외)."""
    gs = [g for g in slide.shapes if g.shape_type == 6
          and 2.4 < g.top / 914400 < 2.8
          and any(x.has_text_frame and x.text_frame.text.strip() for x in g.shapes)]
    gs.sort(key=lambda g: g.left)
    for g, txt in zip(gs, (left_txt, right_txt)):
        tb = max([x for x in g.shapes if x.has_text_frame], key=lambda x: x.width)
        set_tf(tb.text_frame, txt)


def label_shapes(slide, names):
    """밴드/라벨 텍스트 상자를 텍스트로 찾아 반환 (좌→우 순서 보장 못하므로 텍스트 매칭)."""
    out = {}
    for sh in iter_shapes(slide.shapes):
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t in names:
                out.setdefault(t, []).append(sh)
    return out


# ══════════════════════════════════════════════════════════════════════════
# 3. 본문
# ══════════════════════════════════════════════════════════════════════════
def build():
    m = Model(MODEL)
    prs = Presentation(str(SRC))
    S = list(prs.slides)
    assert len(S) == 55, len(S)

    appraisal = m.A("U28") / 1000          # 백만원 334,500
    disc = m.A("Q4")                       # -0.19283
    book = m.A("P28") / 1000               # 288,900 매입장부가
    deposit = m.I("M19")                   # 2,100
    noi_y1 = m.I("C17")
    noi_yield = m.A("O28")                 # 4.57%
    ltv_sr_app, ltv_sr_pr = m.A("T32"), m.A("U32")
    ltv_mz_app, ltv_mz_pr = m.A("T35"), m.A("U35")
    ltv_refi = m.A("T33")
    sale_perf = m.A("AJ64") / 1000         # 매각성과보수 12,480 백만원
    sale_base = m.A("AJ62") / 1000         # 매각기본보수 1,879
    buy_fee = m.A("P9") / 1000             # 1,350
    am_fee_y = m.A("R82") / 1000           # 666.5 (연)
    cus_y = (m.A("R83") + m.A("R84")) / 1000   # 자산보관+사무수탁 (연) 38.0
    cus_rate = cus_y * 1000 / m.A("O83") if m.A("O83") else 0
    aud_y = (m.A("R80") + m.A("R81")) / 1000   # 임원급여+회계감사 96
    fix_fee = buy_fee + am_fee_y * 10
    tot_fee = fix_fee + sale_base + sale_perf

    # ── p11. 코람코 예상수익 ────────────────────────────────────────────
    tb = tables(S[10])[0].table
    put(tb, 5, 2, f"{sale_perf/100:,.1f}억원")
    put(tb, 6, 2, f"{tot_fee/100:,.1f}억원")
    put(tb, 3, 2, f"{fix_fee/100:,.1f}억원")
    put(tb, 4, 2, f"{sale_base/100:,.1f}억원")

    # ── p14~15. Investment Highlights ──────────────────────────────────
    for s in (S[13], S[14]):
        replace_runs(s, "감정가 3,343억 대비 −19.2%", f"감정가 {appraisal/100:,.0f}억 대비 \u2212{abs(disc):.1%}")
        replace_runs(s, "감정가 334,300백만원", f"감정가 {appraisal:,.0f}백만원")
        replace_runs(s, "약 ▼19.2%", f"약 ▼{abs(disc):.1%}")
        replace_runs(s, "감정가 대비 약 -19.2%", f"감정가 대비 약 -{abs(disc):.1%}")
        replace_runs(s, "59.9%", f"{ltv_mz_app:.1%}")
        replace_runs(s, "IRR 16.7%", f"IRR {m.I('J102'):.2%}")
        replace_runs(s, "E. Multiple 4.4 예상", f"E. Multiple {m.I('J103'):.2f} 예상")

    # ── p24. 쿠팡 비중 각주 ────────────────────────────────────────────
    coupang = 1 - sum(m.A(f"AG{r}") for r in (26, 28, 29, 30, 31, 32))
    replace_runs(S[23], "주 임차인(76.5%)", f"주 임차인({coupang:.1%})")

    # ── p30. 리츠개요 · 보수 · 예상수익률 ───────────────────────────────
    s = S[29]
    tb = tables(s)
    fee2 = next(t for t in tb if len(t.table.rows) == 3 and len(t.table.columns) == 5).table
    put(fee2, 1, 3, f"{cus_y:,.0f}백만원")
    put(fee2, 2, 3, f"총 자산가액의 연 {cus_rate:.3%}")
    put(fee2, 1, 2, f"{am_fee_y:,.0f}백만원")
    put(fee2, 1, 4, f"{aud_y:,.0f}백만원")
    yl = next(t for t in tb if len(t.table.rows) == 4 and len(t.table.columns) == 6).table
    put_row(yl, 1, [pc(m.I("C49")), pc(m.I("J49")), pc(m.I("C101")), pc(m.I("J101")), pc(m.I("U101"))], 1)
    put_row(yl, 2, [pc(m.A("G66")), pc(m.A("K66")), pc(m.A("G96")), pc(m.A("K96")), pc(m.I("U102"))], 1)
    put_row(yl, 3, [pc(m.I("C50")), pc(m.I("J50")), pc(m.I("C102")), pc(m.I("J102")), pc(m.I("U103"))], 1)
    lead2(s, f"Equity 전체 연평균 CoC {m.I('U101'):.2%}(매각차익 제외), IRR {m.I('U103'):.2%} 예상",
          f"1종 CoC {m.I('C49'):.2%}·IRR {m.I('C50'):.2%}, 2종 CoC {m.I('J49'):.2%}·IRR {m.I('J50'):.2%}, "
          f"보통주 CoC {m.I('J101'):.2%}·IRR {m.I('J102'):.2%}")
    for sh in by_y(s, 7.45, tol=0.12):
        if sh.has_text_frame and "총 자산가액" in sh.text_frame.text:
            set_tf(sh.text_frame, f"총 자산가액: 부동산 장부가액 {book/100:,.0f}억원 및 여유현금 100억원 예정 / "
                                  f"수익률·보수는 {MODEL_TAG} 기준 (Equity 전체 IRR: IM 탭 전체기준 {m.I('U103'):.2%}, "
                                  f"A&R 민감도 Base {m.A('AZ66'):.2%})")

    # ── p31 · p33. 투자구조도 LTV ──────────────────────────────────────
    for s in (S[30], S[32]):
        replace_runs(s, "LTV 59.9%", f"LTV {ltv_mz_app:.1%}")
        replace_runs(s, "감정가(3,343억원)", f"감정가({appraisal/100:,.0f}억원)")
        replace_runs(s, "LTV 61.8%", f"LTV {ltv_refi:.1%}")

    # ── p38. 운영가정(2/2) ─────────────────────────────────────────────
    s = S[37]
    tb = tables(s)
    opex = next(t for t in tb if len(t.table.rows) == 8 and len(t.table.columns) == 5).table
    for i, row in enumerate([24, 25, 26, 27, 28, 29], start=1):
        ppy, mon, up = m.I(f"Q{row}"), m.I(f"R{row}"), m.I(f"S{row}")
        put(opex, i, 1, "-" if not ppy else f"{ppy:,.0f}원/평")
        put(opex, i, 2, "-" if not mon else f"{mon:,.0f}백만원")
        put(opex, i, 3, "-" if up == 0 else (f"({abs(up):.1%})" if up < 0 else f"{up:.1%}"))
    reit = next(t for t in tb if len(t.table.rows) == 6 and len(t.table.columns) == 4).table
    am_m, cus_m, aud_m, exe_m = m.I("R34"), m.I("R35") + m.I("R36"), m.I("R38") / 12, m.I("R37")
    put_row(reit, 1, [f"AUM X {m.I('Q34'):.3%}", f"{am_m:,.0f}백만원/월", "-"], 1)
    put_row(reit, 2, [f"AUM X {cus_rate:.3%}", f"{cus_m:,.0f}백만원/월", "-"], 1)
    put_row(reit, 3, ["일식", f"{aud_m:,.0f}백만원/월", "-"], 1)
    put_row(reit, 4, ["일식", f"{exe_m:,.0f}백만원/월", "-"], 1)
    put(reit, 5, 2, f"{am_m+cus_m+aud_m+exe_m:,.0f}백만원/월")
    ini = next(t for t in tb if len(t.table.rows) == 18 and len(t.table.columns) == 3).table
    rev = m.I("C8")
    seq = [("C8", 1), ("C4", 2), ("C5", 3), ("C6", 4), ("C7", 5),
           ("C16", 6), ("C9", 7), ("C10", 8), ("C11", 9), ("C12", 10),
           ("C13", 11), ("C14", 12), ("C15", 13), ("C17", 14)]
    for ref, i in seq:
        put(ini, i, 1, n0(m.I(ref)) + " ")
    for i, ref in ((1, "C8"), (6, "C16"), (11, "C13"), (14, "C17")):
        put(ini, i, 2, f"{m.I(ref)/rev:.1%}")
    put(ini, 15, 1, n0(deposit) + " ")
    put(ini, 16, 1, n0(book) + " ")
    put(ini, 17, 1, pc(noi_yield))

    # ── p39. NOI 추이 ─────────────────────────────────────────────────
    s = S[38]
    tb = tables(s)[0].table
    order = [("C8", 1), ("C4", 2), ("C5", 3), ("C6", 4), ("C7", 5),
             ("C16", 6), ("C9", 7), ("C10", 8), ("C11", 9), ("C12", 10),
             ("C13", 11), ("C14", 12), ("C15", 13), ("C17", 14)]
    for ref, i in order:
        put_row(tb, i, [n0(v) for v in m.Irow(int(ref[1:]), AR_YR)], 2)
    put(tb, 15, 12, pc(m.I("M18")))
    put(tb, 16, 12, n0(m.I("M19")))
    put(tb, 17, 12, n0(m.I("M20")))
    lead2(s, "임대료 상승으로 NOI는 1년차 {:,.0f} → 10년차 {:,.0f}백만원 증가 예상".format(m.I("C17"), m.I("L17")),
          "3년차는 쿠팡 연장 Rent-Free 5개월 적용으로 일시 감소 / 매각가치 {:,.0f}백만원".format(m.I("M20")))

    # ── p40 · p41. 투자자별 예상 수익률 ─────────────────────────────────
    def inv_rows(cols, r0, rt):
        """IM 탭 블록 → 26행. cols=(잔액,배당,C.G,합계,Yield), r0=FY1 행, rt=합계 행."""
        b, d, g, t, y = cols
        rows = [["구분", "투자잔액", "운영배당", "C.G", "합계", "Yield(%)"]]
        for k in range(20):
            r = r0 + k
            bal = m.I(f"{b}{r}")
            rows.append([f"FY {k+1}", n0(bal), n0(m.I(f"{d}{r}")), n0(m.I(f"{g}{r}")),
                         n0(m.I(f"{t}{r}")), (pc(m.I(f"{y}{r}"), 1) if bal else "-")])
        rows.append(["합계", "", n0(m.I(f"{d}{rt}")), n0(m.I(f"{g}{rt}")), n0(m.I(f"{t}{rt}")), ""])
        rows.append(["투자금액", n0(m.I(f"{b}{rt+1}") or m.I(f"{y}{rt+1}")), "", "", "", ""])
        rows.append(["CoC_C.G제외", pc(m.I(f"{b}{rt+2}")), "", "", "", ""])
        rows.append(["IRR_C.G포함", pc(m.I(f"{b}{rt+3}")), "", "", "", ""])
        rows.append(["E.Multiple", f"x {m.I(f'{b}{rt+4}'):.2f}", "", "", "", ""])
        return rows

    inv = {
        "1종": inv_rows(("C", "D", "E", "F", "G"), 27, 47),
        "2종": inv_rows(("J", "K", "L", "M", "N"), 27, 47),
        "3종": inv_rows(("C", "D", "E", "F", "G"), 79, 99),
        "보통주": inv_rows(("J", "K", "L", "M", "N"), 79, 99),
    }
    # 투자금액 행(헤더 포함 기준 index 22)
    inv["3종"][22][1] = n0(m.I("C100"))
    inv["보통주"][22][1] = n0(m.I("J100"))
    inv["1종"][22][1] = n0(m.I("C48"))
    inv["2종"][22][1] = n0(m.I("J48"))

    def fill_inv(slide, left_key, right_key):
        tbs = sorted([sh for sh in slide.shapes if sh.has_table], key=lambda x: x.left)
        for tbl_sh, key in zip(tbs, (left_key, right_key)):
            t = tbl_sh.table
            for i, row in enumerate(inv[key]):
                put_row(t, i, row)

    fill_inv(S[39], "3종", "보통주")
    fill_inv(S[40], "1종", "2종")
    # 밴드 라벨 정정 (v05에서 좌·우 표와 라벨이 어긋나 있었음)
    set_band_labels(S[39], "3종 우선주", "보통주")
    set_band_labels(S[40], "1종 우선주", "2종 우선주")
    lead2(S[39], f"3종은 무배당·2년 후 원본 유상감자, 보통주 IRR {m.I('J102'):.2%}·E.M {m.I('J103'):.2f}x 예상",
          f"보통주는 초기 1년 연 7.5% 배당 후 잔여배당, 매각차익의 80%({n0(m.I('L98'))}백만원) 배분")
    lead2(S[40], f"사업기간 10년 가정 투자자별 수익률은 1종 IRR {m.I('C50'):.2%}, 2종 IRR {m.I('J50'):.2%} 예상",
          "제1종·제2종 우선주는 누적적 확정 배당수익률 연 7.0%, 7.5%를 각각 지급 예정")
    fn_add(S[39], f"Source: {MODEL_TAG} IM 탭(B77:N103). 3종 E.Multiple 0.00x는 배당·C.G 기준 산출값으로 "
                  "FY4 원본 유상감자(100억원) 회수분 제외 (FY는 반기)")
    for sh in by_y(S[40], 7.48, tol=0.12):
        if sh.has_text_frame and "Source" in sh.text_frame.text:
            set_tf(sh.text_frame, f"Source: {MODEL_TAG} IM 탭(B25:N51)")

    # ── p42 · p43. Sell-down ──────────────────────────────────────────
    def sd_rows(g):
        """g=0:1종,1:2종,2:보통주,3:통합 — Sell-down 탭 N~AC열."""
        base = ["N", "R", "V", "Z"][g]
        cols = {"N": ("N", "O", "P", "Q"), "R": ("R", "S", "T", "U"),
                "V": ("V", "W", "X", "Y"), "Z": ("Z", "AA", "AB", "AC")}[base]
        b, d, c, y = cols
        rows = [["구분", "투자잔액", "운영배당", "C.G", "합계", "Yield(%)"]]
        for k in range(20):
            r = 68 + k
            bal = m.S(f"{b}{r}") / 1000
            div, cg = m.S(f"{d}{r}") / 1000, m.S(f"{c}{r}") / 1000
            rows.append([f"FY {k+1}", n0(bal), n0(div), n0(cg), n0(div + cg),
                         (pc(m.S(f"{y}{r}"), 1) if bal else "-")])
        td, tc = m.S(f"{d}88") / 1000, m.S(f"{c}88") / 1000
        rows.append(["합계", "", n0(td), n0(tc), n0(td + tc), ""])
        rows.append(["투자금액", n0(m.S(f"{y}89") / 1000), "", "", "", ""])
        rows.append(["CoC_C.G제외", pc(m.S(f"{y}90")), "", "", "", ""])
        rows.append(["IRR_C.G포함", pc(m.S(f"{y}92")), "", "", "", ""])
        rows.append(["E.Multiple", f"x {m.S(f'{y}93'):.2f}", "", "", "", ""])
        return rows

    sd = {g: sd_rows(g) for g in range(4)}
    for slide, pair in ((S[42], (0, 1)), (S[41], (2, 3))):
        tbs = sorted([sh for sh in slide.shapes if sh.has_table], key=lambda x: x.left)
        for tbl_sh, g in zip(tbs, pair):
            for i, row in enumerate(sd[g]):
                put_row(tbl_sh.table, i, row)
    set_band_labels(S[41], "보통주 (KLI 초기 출자)", "통합 (1종+2종+보통주)")
    set_band_labels(S[42], "1종 우선주 (KLI 매입, FY5~)", "2종 우선주 (KLI 매입, FY3~)")
    lead2(S[41], f"KLI 통합 기준 투자잔액 {m.S('AC89')/100000:,.0f}억원, IRR {m.S('AC92'):.2%}, E.Multiple {m.S('AC93'):.2f}x 예상",
          f"보통주(148억원)는 초기 출자, 2종·1종 매입 완료 후 FY5부터 통합 Yield 약 {m.S('AC72'):.1%} 수준")
    lead2(S[42], "KLI리츠가 2종(12개월 후)·1종(24개월 후)을 순차 매입할 경우의 배당스케줄 및 수익률",
          f"매입가에 간주취득세를 포함한 Yield는 1종 {m.S('Q90'):.1%}, 2종 {m.S('U90'):.1%} 수준 (반기 기준 FY)")
    fn_add(S[41], f"Source: {MODEL_TAG} Sell-down 탭. 통합 C.G {m.S('AB88')/1000:,.0f}백만원"
                  f"(1종 {m.S('P88')/1000:,.0f}·2종 {m.S('T88')/1000:,.0f}·보통주 {m.S('X88')/1000:,.0f})")
    fns = sorted([sh for sh in S[42].shapes if sh.has_text_frame and 7.2 < (sh.top or 0) / 914400 < 7.9],
                 key=lambda x: x.left)
    if fns:
        set_tf(fns[0].text_frame, "* FY는 운용 개시(2026-10-31) 기준 반기, 단위 백만원 / 간주취득세 포함 매입가 기준")
    if len(fns) > 1:
        set_tf(fns[1].text_frame, f"Source: {MODEL_TAG} Sell-down 탭 ‘Sell-down 인수 시 수익률’ 블록")

    # ── p44. DSCR ─────────────────────────────────────────────────────
    s = S[43]
    tb = tables(s)[0].table
    spec = [(2, 68, n0), (3, 69, n0), (4, 70, n0), (5, 71, num), (6, 72, num),
            (7, 73, n0), (8, 74, num), (9, 75, num), (10, 76, n0), (11, 77, num),
            (12, 78, num), (13, 79, n0), (14, 80, num), (15, 81, num)]
    for i, row, f in spec:
        put(tb, i, 2, f(m.A(f"X{row}")) if m.A(f"X{row}") else "")
        put(tb, i, 3, f(m.A(f"Y{row}")) if m.A(f"Y{row}") else "")
        put_row(tb, i, [f(m.A(f"{c}{row}")) for c in AR_DS], 4)
    lead2(s, f"선순위 단순 DSCR은 Carry 2개년 최소 {m.A('X71'):,.2f}·평균 {m.A('Y71'):,.2f}, 중순위 최소 {m.A('X74'):,.2f} 수준",
          "제1·2종 우선주 배당 재원 부족분은 3종 우선주 투자금(100억원) Overfunding으로 충당 예정")

    # ── p45. 민감도 ───────────────────────────────────────────────────
    s = S[44]
    tb = tables(s)
    rent = next(t for t in tb if len(t.table.rows) == 7).table
    put_row(rent, 1, [f"{m.A(f'{c}61'):,.0f}천원/평" for c in SENS7], 1)
    for i, row in enumerate((62, 63, 64, 65, 66), start=2):
        put_row(rent, i, [pc(m.A(f"{c}{row}")) for c in SENS7], 1)
    vac = next(t for t in tb if len(t.table.rows) == 6).table
    put(vac, 0, 1, f"Base/({m.A('AW70'):.2%})")
    for i, row in enumerate((71, 72, 73, 74, 75), start=1):
        put_row(vac, i, [pc(m.A(f"{c}{row}")) for c in SENS7], 1)
    ir = next(t for t in tb if len(t.table.rows) == 12).table
    for i, row in enumerate((79, 80), start=1):
        put_row(ir, i, [pc(m.A(f"{c}{row}")) for c in SENS6], 2)
    for i, row in enumerate((81, 82, 83, 84, 85), start=3):
        put_row(ir, i, [pc(m.A(f"{c}{row}")) for c in SENS6], 2)
    for i, row in enumerate((86, 87, 88, 89), start=8):
        put_row(ir, i, [num(m.A(f"{c}{row}")) for c in SENS6], 2)

    # ── p48. 리스크 (재무·금리) ────────────────────────────────────────
    s = S[47]
    tb = tables(s)[0].table
    for i in range(len(tb.rows)):
        k = tb.cell(i, 0).text.replace("\n", "")
        if "재무리스크" in k:
            set_cell(tb.cell(i, 1),
                     f"감정가 기준 LTV {ltv_mz_app:.1%}(선순위 {ltv_sr_app:.1%})이나 매입가 기준 실질 LTV 약 {ltv_mz_pr:.1%}로 레버리지 수준 관리 필요"
                     f"/대주 담보인정비율 이내로 설정되어 있으며, Carry 2개년 선순위 단순 DSCR 최소 {m.A('X71'):,.2f}, "
                     f"운영기간 내 2종 우선주 누적 DSCR {min(m.A(f'{c}81') for c in AR_DS):,.2f} 이상 유지")
        if "금리변동" in k:
            set_cell(tb.cell(i, 1),
                     "제반 경제 상황에 따라 금리 상승 및 조달금리 변동 리스크 존재함"
                     f"/선순위 금리 {m.A('Q32'):.2%}(All-in {m.A('S32'):.2%}), 중순위 {m.A('Q35'):.2%}(All-in {m.A('S35'):.2%}) 고정, "
                     "만기 24개월으로 조달. 만기 6개월 전 차환 협의 착수 및 시장 스프레드 상시 모니터링 예정")
        if "Exit" in k:
            set_cell(tb.cell(i, 1), tb.cell(i, 1).text.replace("4.88%", f"{m.A('O4'):.2%}"))

    # ── p49. Checklist ────────────────────────────────────────────────
    tb = tables(S[48])[0].table
    for i in range(len(tb.rows)):
        if "수치 확정" in tb.cell(i, 0).text:
            set_cell(tb.cell(i, 1),
                     f"본 자료 수치는 {MODEL_TAG} A&R·IM 탭 기준으로 정합 완료 "
                     f"(2종 IRR {m.I('J50'):.2%}·보통주 {m.I('J102'):.2%}) / 잔여 확인: WALE 산정기준(IM 6.8년 vs 모델 {m.A('AI3'):,.2f}년)")
            set_cell(tb.cell(i, 2), "실사보고서 최종본(9/17) 반영 시 WALE 기준 통일 후 투자자 Q&A·IM 갱신")
        if "실사 반영" in tb.cell(i, 0).text:
            set_cell(tb.cell(i, 1),
                     f"대지면적 차이(IM·재무DD 15,287.5㎡ vs 모델 {m.A('J22'):,.1f}㎡, 필지 수) / "
                     "법률실사 Recommendation 41건 중 미해소 항목 / 실사보고서 최종본(9/17) 반영")

    normalize_tables(prs)
    autofit_tables(prs)
    prs.save(str(OUT))
    print("생성 완료:", OUT)
    return OUT


if __name__ == "__main__":
    build()
