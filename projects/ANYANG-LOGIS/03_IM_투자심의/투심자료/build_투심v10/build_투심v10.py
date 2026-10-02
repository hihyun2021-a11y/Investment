# -*- coding: utf-8 -*-
"""ANYANG-LOGIS 본 투심자료 v09 → v10.

사용자 지시(2026-10-02): 주주간계약서 최종본(BKL v30.5 markup, 투자자 송부 메일 포함) 반영,
투자심의위원회 개요(안건개요) 섹션을 표·도표 중심으로 가독성 있게 재구성.
"""
import copy, pathlib, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph

HERE = pathlib.Path(__file__).parent
SRC = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else HERE.parent / "20261002_투심자료_ANYANG-LOGIS_v09.pptx"
OUT = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "out.pptx"
F_B, F_M, F_L = "KoPubWorld돋움체 Bold", "KoPubWorld돋움체 Medium", "KoPubWorld돋움체 Light"
NAVY, BAND, INK, GRAY, RED = "002060", "345B86", "1F1F1F", "595959", "C00000"
LBLUE, LGRAY, MGRAY, ORANGE, LORANGE = "DAE3F3", "F2F2F2", "D9D9D9", "F78E3F", "FDE9D9"
SHA = "주주간계약서 BKL v30.5 markup(2026-10-02, 최종)"

prs = Presentation(str(SRC))
S = lambda n: prs.slides[n - 1]
IN = lambda v: Emu(int(v * 914400))

# ── 공통 도구 ────────────────────────────────────────────────────────────
def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)

def find(slide, sid):
    for sh in walk(slide.shapes):
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)

def remove(slide, sid):
    sh = find(slide, sid); sh._element.getparent().remove(sh._element)

def set_font(run, name, size, color=INK, bold=False):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)

def set_para(p, text):
    runs = p.runs
    if not runs:
        r = p.add_run()
        end = p._p.find(qn("a:endParaRPr"))
        if end is not None:
            rp = copy.deepcopy(end); rp.tag = qn("a:rPr")
            old = r._r.find(qn("a:rPr"))
            if old is not None: r._r.remove(old)
            r._r.insert(0, rp)
        runs = [r]
    runs[0].text = text
    for x in runs[1:]:
        p._p.remove(x._r)

def set_lines(shape, lines):
    tf = shape.text_frame if hasattr(shape, "text_frame") else shape
    body = tf._txBody
    while len(tf.paragraphs) < len(lines):
        body.append(copy.deepcopy(tf.paragraphs[-1]._p))
    while len(tf.paragraphs) > len(lines):
        body.remove(tf.paragraphs[-1]._p)
    for p, t in zip(tf.paragraphs, lines):
        set_para(p, t)

def band_label(slide, text, idx=0):
    groups = [g for g in slide.shapes if g.shape_type == 6 and any(
        c.has_text_frame and c.text_frame.text.strip() for c in g.shapes)]
    g = sorted(groups, key=lambda x: x.top)[idx]
    tb = [c for c in g.shapes if c.has_text_frame and c.text_frame.text.strip()][0]
    set_lines(tb, [text])
    return g

def write(tf, lines, size=9, color=INK, font=F_L, align=PP_ALIGN.CENTER, bold_first=False, first_size=None,
          first_color=None, first_font=None):
    """lines: str 또는 (text, dict) — dict로 size/color/font 개별 지정."""
    tf.clear(); tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, ln in enumerate(lines):
        txt, opt = (ln, {}) if isinstance(ln, str) else ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        r = p.add_run(); r.text = txt
        if i == 0 and (bold_first or first_size or first_color or first_font):
            set_font(r, opt.get("font", first_font or (F_B if bold_first else font)), opt.get("size", first_size or size),
                     opt.get("color", first_color or color), bold=False)
        else:
            set_font(r, opt.get("font", font), opt.get("size", size), opt.get("color", color))

def box(slide, x, y, w, h, lines, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75, **kw):
    sh = slide.shapes.add_shape(shape, IN(x), IN(y), IN(w), IN(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    write(sh.text_frame, lines, **kw)
    return sh

def arrow(slide, x, y, w, h, color=MGRAY):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, IN(x), IN(y), IN(w), IN(h))
    sh.shadow.inherit = False
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(color); sh.line.fill.background()
    return sh

def hline(slide, x1, y, x2, color=BAND, width=1.5, dash=False):
    ln = slide.shapes.add_connector(1, IN(x1), IN(y), IN(x2), IN(y))
    ln.line.color.rgb = RGBColor.from_string(color); ln.line.width = Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln

def ln_el(tag, w=3175, color="BFBFBF"):
    el = etree.Element(qn(f"a:{tag}"), w=str(w), cap="flat", cmpd="sng", algn="ctr")
    if color is None:
        etree.SubElement(el, qn("a:noFill"))
    else:
        sf = etree.SubElement(el, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=color)
    return el

def table(slide, rows, x, y, widths, row_h=0.30, size=9, head=True, label_col=True, aligns=None,
          colors=None, head_fill=MGRAY):
    nr, nc = len(rows), len(widths)
    gf = slide.shapes.add_table(nr, nc, IN(x), IN(y), IN(sum(widths)), IN(row_h * nr))
    t = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    for i, w in enumerate(widths):
        t.columns[i].width = IN(w)
    for ri, row in enumerate(rows):
        t.rows[ri].height = IN(row_h)
        for ci in range(nc):
            val = row[ci] if ci < len(row) else ""
            cell = t.cell(ri, ci)
            is_head = head and ri == 0
            is_label = label_col and ci == 0 and not is_head
            kind_font = F_B if is_head else (F_M if is_label else F_L)
            al = PP_ALIGN.CENTER if (is_head or is_label) else (aligns[ci] if aligns else PP_ALIGN.LEFT)
            col = (colors or {}).get((ri, ci), INK)
            tf = cell.text_frame
            lines = val.split("\n") if isinstance(val, str) else val
            write(tf, lines, size=size, font=kind_font, align=al, color=col)
            tcPr = cell._tc.get_or_add_tcPr()
            for ch in list(tcPr): tcPr.remove(ch)
            tcPr.set("anchor", "ctr"); tcPr.set("marL", "54000"); tcPr.set("marR", "36000")
            tcPr.set("marT", "9525"); tcPr.set("marB", "9525")
            tcPr.append(ln_el("lnL", 9525, None) if ci == 0 else ln_el("lnL"))
            tcPr.append(ln_el("lnR", 9525, None) if ci == nc - 1 else ln_el("lnR"))
            tcPr.append(ln_el("lnT", 6350, "000000") if ri == 0 else ln_el("lnT"))
            tcPr.append(ln_el("lnB", 6350, "000000") if ri == nr - 1 else ln_el("lnB"))
            fill = head_fill if is_head else (LGRAY if is_label else None)
            if fill:
                sf = etree.SubElement(tcPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=fill)
            else:
                etree.SubElement(tcPr, qn("a:noFill"))
    return t

def unit(slide, right, y, text):
    tb = slide.shapes.add_textbox(IN(right - 2.6), IN(y), IN(2.6), IN(0.2))
    write(tb.text_frame, [text], size=8, color=INK, align=PP_ALIGN.RIGHT)
    tb.text_frame.margin_right = 0

def kpi(slide, x, y, w, h, label, value, sub, accent=BAND):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(x), IN(y), IN(0.06), IN(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(accent); bar.line.fill.background()
    bar.shadow.inherit = False
    box(slide, x + 0.06, y, w - 0.06, h, [(label, {"size": 9, "color": GRAY, "font": F_M}),
                                          (value, {"size": 18, "color": NAVY, "font": F_B}),
                                          (sub, {"size": 8, "color": GRAY})], fill=LGRAY, align=PP_ALIGN.LEFT)

L0, R0 = 0.57, 11.24          # 본문 좌·우 기준선
W0 = R0 - L0

# ══════════════════════════════════════════════════════════════════════
# 슬3 심의안건 및 추진일정 — 투자자 송부 일정(10/2) 반영
# ══════════════════════════════════════════════════════════════════════
s = S(3)
set_lines(find(s, 3), ["본 심의 안건은 매매계약·주주간계약 체결(10/16) 및 KLI 매입확약서 제출 건임",
                       "주금 납입(10/20) 후 10/22 자리츠 명의 소유권이전으로 거래종결 예정"])
t = find(s, 11).table
for i, txt in enumerate(["코크렙안양㈜와 매매대금 2,700억원(VAT 별도) 부동산매매계약 체결 (10/16 체결, 10/22 종결)",
                         "KLI·이지스·키움·애큐온·MG·삼성증권·기계설비조합·코크렙안양 간 주주간계약 (총 988억원)",
                         "1종 240억원(24개월)·2종 340억원(12개월, KLI 보유분 제외) 매입확약서 제출 (10/16)"], start=1):
    set_lines(t.cell(i, 2).text_frame, [txt])
for sid, lines in ((56, ["10월 14일"]), (67, ["당 투자심의위원회 [10/14]"]),
                   (59, ["주주간계약 체결", "매매계약 체결", "대출약정 체결", "KLI 이사회"]),
                   (29, ["10월 20일"]), (25, ["주금 납입"]), (47, ["10월 22일"]), (45, ["거래종결", "(소유권이전)"])):
    set_lines(find(s, sid), lines)
set_lines(find(s, 109), ["코람코라이프로지스리츠(자리츠) 2026.9.11 자본금 3억원 발기설립 완료 (KLI 100% 출자)",
                         "투자자 송부 일정(10/2) 기준: 10/7 투자자 의견 회신 → 10/8 최종합의 회람 → 10/14 투심위 → 10/16 체결 → 10/20 납입 → 10/22 종결 (독점협상 만료 10/28)"])

# ══════════════════════════════════════════════════════════════════════
# 슬4 매매계약 (1/2) — 핵심지표 카드 + 주요 조건 표
# ══════════════════════════════════════════════════════════════════════
s = S(4)
remove(s, 8)
band_label(s, "매매계약 핵심 조건")
cw, gap, y0 = (W0 - 0.3) / 4, 0.1, 1.58
for i, (lab, val, sub, ac) in enumerate([
        ("매매대금", "2,700억원", "감정가 3,345억 대비 −19.3% · VAT 별도", BAND),
        ("계약금", "100억원", "체결일(10/16) 지급 · 해약금 아님", ORANGE),
        ("거래종결", "10/22(목)", "잔금 지급 · 소유권이전 · 임대차 승계", BAND),
        ("손해배상 한도", "135억원", "매매대금 5% · 청구기한 종결 후 2개월", ORANGE)]):
    kpi(s, L0 + i * (cw + gap), y0, cw, 0.95, lab, val, sub, ac)
rows = [["구분", "주요 내용", "조항"],
        ["체결당사자", "매도인 코크렙안양㈜(PFV, 한국투자부동산신탁 담보신탁) ↔ 매수인 코람코라이프로지스리츠", "전문"],
        ["매매목적물", "관양동 912-2·934·934-1 토지·건물(대지 15,287.5㎡ / 연면적 95,434.55㎡), 가설건축물·부속 동산", "제1.1조·별첨1"],
        ["매매 방식", "거래종결일 현황 그대로(as-is, where-is) · 담보책임 배제 (매도인 고의·중과실은 예외)", "제1.1·4.3조"],
        ["종결시 지급금액", "매매대금 − 계약금 − 승계 임대보증금 ± 정산금 · VAT는 매수인 대리납부", "제1.3·1.4·3조"],
        ["임대차 승계", "임대차 전부 승계·보증금 반환의무 인수 · 쿠팡 승계동의서는 매수인 선행조건", "제5.1(d)·6.1(d)조"],
        ["진술 및 보장", "설립·권한·소유권·소송·임대차·경계침범 · 거래종결일부터 2개월 유효", "제4.1·4.4조"],
        ["신고·준거법", "체결 후 30일 내 거래신고 · 기업결합 사후신고(종결 후 30일, 11/21) · 서울중앙지법", "제5.1(c)(e)·9.8조"]]
table(s, rows, L0, 2.72, [1.55, 7.47, 1.65], row_h=0.52, size=9,
      aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.CENTER])

# ══════════════════════════════════════════════════════════════════════
# 슬5 매매계약 (2/2) — 거래 절차 도식 + 쟁점·검토의견 표
# ══════════════════════════════════════════════════════════════════════
s = S(5)
remove(s, 15)
band_label(s, "거래 절차 및 주요 쟁점")
steps = [("① 계약 체결", "10/16", "계약금 100억원 지급"),
         ("② 선행조건 충족", "~10/21", "확약 이행 · 진술보장 진실\n쿠팡 승계동의 · 우선수익자 동의"),
         ("③ 거래종결", "10/22", "잔금 지급 · 소유권이전\n임대차·보증금 승계"),
         ("④ 종결 후 의무", "~12/22", "진술보장 2개월 · 기업결합 신고 30일\n하자보수청구권 양도")]
cw = (W0 - 0.15 * 3) / 4
for i, (t1, d, t2) in enumerate(steps):
    x = L0 + i * (cw + 0.15)
    box(s, x, 1.58, cw, 0.36, [(f"{t1}  |  {d}", {"size": 10, "color": "FFFFFF", "font": F_B})],
        fill=NAVY if i != 2 else ORANGE, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
    box(s, x, 1.98, cw - 0.12, 0.62, [(l, {"size": 8.5}) for l in t2.split("\n")], fill=LGRAY)
rows = [["조항", "주요 내용", "검토 의견"],
        ["계약금\n(제1.3조(a))", "체결일 현금 100억원 지급, 해약금 아님\n3종 주금은 매매대금채권과 상계납입(SHA 제2.2조)", "▲ 체결일 자리츠 자금 3억원뿐\n상계 방식으로 문안 정비 필요"],
        ["매도인 확약\n(제5.2조)", "담보신탁 해지·대출 상환, 보증금 질권 교체, 하자보수청구권 양도\n현황측량·경계 정리, 옥외광고 허가, 소음민원 조치", "보증금 질권 교체 예치금\n약 50.7억원 자금계획 필요"],
        ["선행조건\n(제6조)", "확약 이행·진술보장 진실·중대한 법규·소송 부존재\n매수인 CP: 쿠팡 승계동의서, 우선수익자 전원·수탁자 동의", "미충족 시 청구권 유보하고\n종결 가능 (제6.1조)"],
        ["손해배상\n(제7조)", "한도 매매대금 5%(135억원), 청구기한 종결 후 2개월\n고의·중과실 위반은 한도·기간 미적용", "종결 직후 2개월 내\n하자·진술 점검"],
        ["해제\n(제8조)", "서면합의·불가항력 30일·중요 위반 10영업일 미시정·도산\n귀책 해제 시 위약금(손해배상 예정), 계약금 이자 가산 반환", "▲ 위약금·기간 [*] 공란\n체결 전 확정 필요"],
        ["일반조항\n(제9조)", "양도 제한, 비밀유지(종결·해제 후 1년), 비용 각자 부담\n양해각서 이행보증금(10억원) 반환 관련 권리·의무 존속", "제7~9조 해제 후 존속\n(제9.9조)"]]
red = {(1, 2): RED, (5, 2): RED}
table(s, rows, L0, 2.78, [1.45, 6.37, 2.85], row_h=0.66, size=8.5, colors=red,
      aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT])

# ══════════════════════════════════════════════════════════════════════
# 슬6 주주간계약 (1/3) — 출자 구조 도식 + 주요 조건 표
# ══════════════════════════════════════════════════════════════════════
s = S(6)
remove(s, 19)
band_label(s, "출자 구조 및 주요 조건")
unit(s, R0, 1.50, "(단위: 억원)")
segs = [("보통주", 148, "KLI리츠", "@5,000원 · 의결권", NAVY),
        ("제1-1종", 240, "이지스3호 100 · 키움 60\n애큐온 50 · MG 30", "7.0% 누적 · 의결권", "2F5597"),
        ("제2-1종", 280, "KLI리츠 160 · 삼성증권 120", "7.5% 누적 · 의결권", BAND),
        ("제2-2종", 220, "이지스K 100 · 기계설비조합 120", "7.5% 누적 · 무의결권", "8EA9DB"),
        ("제3종", 100, "코크렙안양(매도인)", "@23,000원 · 무의결권", ORANGE)]
tot = sum(v for _, v, *_ in segs)
x = L0
for name, v, who, cond, col in segs:
    w = W0 * v / tot
    box(s, x, 1.72, w - 0.03, 0.46, [(f"{name}  {v:,}", {"size": 10, "color": "FFFFFF", "font": F_B})], fill=col)
    box(s, x, 2.20, w - 0.03, 0.56, [(l, {"size": 8}) for l in who.split("\n")] + [(cond, {"size": 7.5, "color": GRAY})])
    x += w
box(s, L0, 2.78, W0, 0.24, [("총 988억원 (무의결권 1,314,783주 = 발행주식 20.7%, 상법 §344조의3 1/4 이내) · 2종 투자자별 물량은 10/8 최종합의 회람 시 확정 예정",
                              {"size": 8, "color": GRAY})], align=PP_ALIGN.LEFT)
rows = [["구분", "주요 내용", "조항"],
        ["주금 납입", "거래종결 2영업일 전 이사회가 정한 납입예정일(10/20) 15:00까지 · 3종은 매매대금채권과 상계납입", "제2.1·2.2조"],
        ["발기주식 감자·인수", "KLI 발기주식 30만주(3억원) 100% 유상감자 · 미납입분은 우리투자증권(1종)·삼성증권(2종) 잔액인수", "제2.3·2.5조"],
        ["배당", "1종 연 7.0%·2종 연 7.5% 누적 · 보통주 초기 1년 연 7.5% 배당 협조(재원·결의 없으면 의무 없음)", "제1.3조"],
        ["KLI 매입확약", "2종 340억원(12개월, KLI 보유분 제외)·1종 240억원(24개월) 발행가액 매입 · 미수령 누적배당 정산", "제3·4조"],
        ["제3종 조건", "24+2개월 내 유상감자 · 미감자 시 연 6%→8% 누적우선배당 · 48+2개월 후 자산매각청구권", "제5조"],
        ["지배구조", "이사 3인 중 2인·대표이사 KLI 지명, 감사 주총 선임 · 동일 AMC(코람코) 이해상충 고지", "제1.4·6조"],
        ["양도 제한", "전원 서면동의 원칙 · 셀다운은 매도확약서 제출+KLI 사전동의(우리투자·삼성 최초 셀다운 면제)", "제7조"]]
table(s, rows, L0, 3.12, [1.65, 7.37, 1.65], row_h=0.48, size=9,
      aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.CENTER])
set_lines(find(s, 16), [f"※ 출처: {SHA} 제1~7조·별첨1, 투자자 송부 메일(10/2). 본 검토는 내부 검토이며 최종 법률자문은 외부 법무법인 확인이 필요함."])

# ══════════════════════════════════════════════════════════════════════
# 슬7 지분구성표 — 투자자 송부 표 기준 명칭
# ══════════════════════════════════════════════════════════════════════
s = S(7)
t = find(s, 19).table
for i, nm in ((1, "이지스리츠3호펀드"), (4, "MG캐피탈"), (7, "이지스K리츠펀드"), (8, "기계설비조합"), (9, "코크렙안양주식회사")):
    set_lines(t.cell(i, 1).text_frame, [nm])
set_lines(find(s, 16), [f"※ 출처: {SHA} 제2.1·2.4조, 투자자 송부 메일(10/2). 주주명은 확정, 2종 종류주식 물량은 10/8 최종합의 회람 시 확정 예정(세부 변경 가능). "
                        "3종 23,000원×434,783주=10,000,009,000원(단수 BKL 확인 중)."])

# ══════════════════════════════════════════════════════════════════════
# 슬8 주주간계약 (3/3) — 제3종 조건 타임라인 + 배당 순위 도식 + 2종 권리 보호
# ══════════════════════════════════════════════════════════════════════
s = S(8)
remove(s, 10)
band_label(s, "제3종 조건 및 배당 순위")
hline(s, L0, 2.12, R0 - 0.1, color=MGRAY, width=6)
tl = [("’26.10 발행", "100억원 · 23,000원\n무배당·무의결권 · 상계납입", LGRAY, INK),
      ("’28.12 감자약정기한", "발행 후 24+2개월\n발행가액 유상감자", LBLUE, INK),
      ("미감자 1년차", "연 6% 누적우선배당\n(배당가능재원 내)", LORANGE, INK),
      ("미감자 2년차", "연 8% 누적우선배당\n미배당분은 별도 협의", "F8CBAD", INK),
      ("’30.12 추가기한", "48+2개월 도과 시\n자산매각청구권(매각 완료 시까지)", ORANGE, "FFFFFF")]
cw = (W0 - 0.12 * 4) / 5
for i, (h1, h2, fill, col) in enumerate(tl):
    x = L0 + i * (cw + 0.12)
    box(s, x + cw / 2 - 0.09, 2.03, 0.18, 0.18, [""], fill=NAVY, shape=MSO_SHAPE.OVAL)
    box(s, x, 1.56, cw, 0.36, [(h1, {"size": 9.5, "font": F_B, "color": NAVY})])
    box(s, x, 2.32, cw, 0.62, [(l, {"size": 8.5, "color": col}) for l in h2.split("\n")], fill=fill)
box(s, L0, 2.98, W0, 0.22, [("정관 개정: 제5.3조 권리를 거래종결일부터 3개월 내 정관에 반영 (변경인가 지연 시 연장, 제5.5조) · 지급 보장·보전 의무 아님",
                              {"size": 8, "color": RED})], align=PP_ALIGN.LEFT)
# 배당 순위 도식
box(s, L0, 3.32, 5.3, 0.28, [("배당 순위 (정관 제11조 · SHA 제1.3·5.3조)", {"size": 9.5, "font": F_B, "color": NAVY})], align=PP_ALIGN.LEFT)
def ladder(y, title, items, fills, h=0.62):
    box(s, L0, y, 0.95, h, [(l, {"size": 8.5, "font": F_M}) for l in title.split("\n")], fill=LGRAY)
    n = len(items); aw = 0.10
    w = (5.3 - 1.0 - aw * (n - 1)) / n
    for i, (it, f) in enumerate(zip(items, fills)):
        x = L0 + 1.0 + i * (w + aw)
        box(s, x, y, w, h, [(l, {"size": 8, "color": "FFFFFF" if f in (NAVY, BAND, ORANGE) else INK})
                            for l in it.split("\n")], fill=f)
        if i < n - 1:
            tri = box(s, x + w + 0.01, y + h / 2 - 0.07, aw - 0.02, 0.14, [""], fill=MGRAY, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
            tri.rotation = 90
ladder(3.68, "평상시", ["1종\n7.0% 누적", "2종\n7.5% 누적", "3종 가산\n6→8%*", "보통주\n잔여"],
       [NAVY, BAND, ORANGE, "8EA9DB"])
ladder(4.42, "매각·청산\n(개정 예정)", ["1종\n배당", "2종\n배당", "1종\n원본", "2종\n원본", "3종\n배당·원본", "보통주\n원본·잔여"],
       [NAVY, BAND, NAVY, BAND, ORANGE, "8EA9DB"])
ladder(5.16, "매각·청산\n(현행 정관 v8)", ["배당\n(①~⑥)", "1종\n원본", "2종\n원본", "3종\n원본", "보통주\n원본", "잔여 20%\n2종·80% 보통"],
       [LGRAY, LGRAY, LGRAY, LGRAY, LGRAY, LGRAY])
box(s, L0, 5.86, 5.3, 0.45, [("* 3종 가산배당은 감자약정기한 도과 시에만 발생, 정관 개정은 거래종결 후 3개월 내(제5.5조)",
                              {"size": 7.5, "color": GRAY, "align": PP_ALIGN.LEFT}),
                             ("  매각·청산 순위는 1·2종 권리 역전이 없도록 개정 시 반영 예정", {"size": 7.5, "color": GRAY, "align": PP_ALIGN.LEFT})],
    align=PP_ALIGN.LEFT)
# 2종 권리 보호
rows = [["구분", "2종 종류주식 권리 영향 검토 (투자자 안내 기준)"],
        ["Exit 시점", "2종은 거래종결 후 1년 내 KLI가 매입확약 → 감자약정기한(26개월) 이후 발생 페널티 시점엔 이미 Exit"],
        ["배당 순위", "KLI 사유로 2종이 남더라도 3종 가산배당은 2종 배당 후순위 → 워터폴상 2종 배당 영향 없음"],
        ["매각 시", "추가기한 도과로 자산매각청구권 행사 시에도 잔여재산 분배에서 2종 지위 불변"],
        ["기타 변경", "미배당분 처리 별도 협의(제5.3조), 매각청구는 매각 완료 시까지 가능(제5.4조), 매도기한 15일 통일"]]
table(s, rows, 6.05, 3.32, [1.05, 4.14], row_h=0.58, size=8.5,
      aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT])
set_lines(find(s, 16), [f"※ 출처: {SHA} 제1.3·5.1~5.5조, KLL리츠 정관 v8 제11조, 투자자 송부 메일(10/2). 본 검토는 내부 검토이며 최종 법률자문은 외부 법무법인 확인이 필요함."])

# ══════════════════════════════════════════════════════════════════════
# 슬9 KLI 매입확약 — 매입 일정 도식 + 확약 비교표
# ══════════════════════════════════════════════════════════════════════
s = S(9)
remove(s, 17); remove(s, 18)
band_label(s, "KLI 지분 확대 일정 (매입확약)", 0)
g2 = band_label(s, "매입·매도확약 주요 조건", 1)
g2.top = IN(4.00)
unit(s, R0, 1.50, "(단위: 억원)")
hline(s, L0, 2.30, R0 - 0.1, color=MGRAY, width=6)
ms = [("’26.10.22 거래종결", "보통주 148 + 2-1종 160\nKLI 출자 308", NAVY),
      ("’27.10 (12개월)", "2종 340 매입\n(삼성 120 · 이지스K 100 · 기계설비 120)", BAND),
      ("’28.10 (24개월)", "1종 240 매입\n(이지스3호·키움·애큐온·MG)", BAND),
      ("KLI 보유 합계", "888 = 보통주 148\n+ 1종 240 + 2종 500", ORANGE)]
cw = (W0 - 0.2 * 3) / 4
for i, (h1, h2, col) in enumerate(ms):
    x = L0 + i * (cw + 0.2)
    box(s, x, 1.62, cw, 0.36, [(h1, {"size": 10, "font": F_B, "color": NAVY})])
    box(s, x + cw / 2 - 0.1, 2.20, 0.2, 0.2, [""], fill=col, shape=MSO_SHAPE.OVAL)
    box(s, x, 2.50, cw, 0.80, [(l, {"size": 9 if j == 0 else 8, "font": F_B if j == 0 else F_L,
                                    "color": "FFFFFF"}) for j, l in enumerate(h2.split("\n"))], fill=col)
box(s, L0, 3.40, W0, 0.40, [("KLI 지분: 거래종결 시 의결권 71.4% → 1년차 2종 매입 → 2년차 1종 매입 · 3종은 KLI 매입 대상 아님(제5조 유상감자로 회수)",
                              {"size": 8.5, "color": GRAY})], align=PP_ALIGN.LEFT)
rows = [["구분", "제1종 매입확약 (KLI → 종류주주)", "제2종 매입확약 (KLI → 종류주주)", "매도확약 (종류주주 → KLI)"],
        ["대상·금액", "1-1종 960,000주 · 240억원", "2종 1,360,000주 · 340억원\n(KLI 보유 2-1종 제외)", "보유 종류주식 전부\n(셀다운 양수인 포함)"],
        ["기한", "발행일부터 24개월 응당일", "발행일부터 12개월 응당일", "KLI 서면 요청일부터 15일 내"],
        ["불이행 시", "미매입잔액 연 5.0% 지연손해금\n매입의무 존속", "연 3.0% 지연이자 + 위약벌 20%\n매입의무 존속", "1종 연 5.0% / 2종 연 3.0%+20%\n매도의무 존속"],
        ["선행조건", "영업인가·소유권 취득, 완전한 소유권·부담 부존재, 매도확약서 제출", "(좌 동)", "KLI의 보유주식 전부 매도 요청"],
        ["배당 정산·제출", "매입완료일에 미수령 누적배당 정산 · 주주간계약 체결일(10/16) 제출", "(좌 동)", "인수인 체결일 / 셀다운 양수인 완결 전"]]
table(s, rows, L0, 4.40, [1.35, 3.15, 3.10, 3.07], row_h=0.48, size=8.5,
      aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
set_lines(find(s, 16), [f"※ 출처: KLI 제1·2종 매입확약 공문(안), 매도확약 공문(안), {SHA} 제2.5·3·4조, 투자자 송부 메일(10/2, 확약서 발행회사 기명날인·매도기한 15일 통일). "
                        "본 검토는 내부 검토이며 외부 법무법인 확인 필요."])

# ══════════════════════════════════════════════════════════════════════
# 일정 변경 반영 (기타 페이지)
# ══════════════════════════════════════════════════════════════════════
def replace_all(slide, old, new):
    for sh in walk(slide.shapes):
        frames = [sh.text_frame] if sh.has_text_frame else []
        if getattr(sh, "has_table", False) and sh.has_table:
            frames += [c.text_frame for r in sh.table.rows for c in r.cells]
        for tf in frames:
            for p in tf.paragraphs:
                if old in p.text:
                    set_para(p, p.text.replace(old, new))
replace_all(S(31), "KLI리츠 이사회 승인 예정(10/14)", "KLI리츠 이사회 승인 예정(10/16)")
replace_all(S(47), "10/14(E) → 10/16 체결", "10/16(E)")
replace_all(S(47), "10/14(E)", "10/16(E)")

# 세로 중간 정렬 (사내 표준)
for sl in list(prs.slides)[2:9]:
    for el in sl.shapes._spTree.iter(qn("a:bodyPr")):
        el.set("anchor", "ctr")
    for el in sl.shapes._spTree.iter(qn("a:tcPr")):
        el.set("anchor", "ctr")

prs.save(str(OUT))
print("saved", OUT)
