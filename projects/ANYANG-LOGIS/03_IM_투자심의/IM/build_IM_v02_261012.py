# -*- coding: utf-8 -*-
"""IM v02 생성기 — KLI 투심 최종본 v5.1(2026-10-12)의 [별첨] 투자대상 물건 설명자료를 기준으로
Information Memorandum을 갱신한다.

- 기준: 03_IM_투자심의/투심자료/20261012_투심자료_ANYANG-LOGIS_KLI최종본_v5.1.pptx (원본 불변, 복사본 편집)
- 유지: v5.1 22~50장(표지·목차·Highlights·Ⅰ 사업개요·Ⅱ 투자계획·Ⅲ 운영계획), 64~65장(리스크), 66장(종료)
- 제외: 1~21장(투심 본문: 당사 손익·심의안건·계약 요약·LOC·당사 예상수익),
        51~61장(Ⅳ 모리츠 투자계획 — 사내 의사결정 자료), 62~63장(리스크관리부서 의견)
- 추가: Ⅴ 시설현황(층별 도면·Floor Plan) 3장 — 종전 IM(20260914 v01) p.25~27 도면 이미지 재사용
- 문구: 표지·목차·마스터 꼬리말을 IM용으로 변경, 리스크 장 번호 Ⅴ→Ⅳ
실행: python3 build_IM_v02_261012.py
"""
import copy
import io
import os

import pymupdf
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
PRJ = os.path.dirname(os.path.dirname(HERE))
SRC = os.path.join(PRJ, "03_IM_투자심의/투심자료/20261012_투심자료_ANYANG-LOGIS_KLI최종본_v5.1.pptx")
OLD_IM = os.path.join(HERE, "20260914_IM_코크렙안양물류센터I_코람코_v01.pdf")
OUT = os.path.join(HERE, "20261012_IM_안양물류센터_InformationMemorandum_v02.pptx")

KEEP = list(range(22, 51)) + [64, 65, 66]          # 1-based (v5.1 기준)
FLOOR_PAGES = [24, 25, 26]                          # 0-based → 종전 IM p.25~27
CLOSING = 66


def replace_in_runs(text_frame_like, mapping):
    """문단 단위로 run 텍스트를 이어 붙여 치환한 뒤 첫 run에 몰아 넣는다(서식은 첫 run 유지)."""
    changed = False
    for p in text_frame_like.paragraphs:
        full = "".join(r.text for r in p.runs)
        new = full
        for a, b in mapping:
            new = new.replace(a, b)
        if new != full and p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""
            changed = True
    return changed


prs = Presentation(SRC)
slides = list(prs.slides)
closing_slide = slides[CLOSING - 1]

# ---------------------------------------------------------------- 1) 표지(22)
cover = slides[21]
for sh in cover.shapes:
    if sh.has_text_frame:
        replace_in_runs(sh.text_frame, [("안양물류센터 매입의 건", "안양물류센터"),
                                        ("투자설명자료", "Information Memorandum")])

# ---------------------------------------------------------------- 2) 목차(23)
toc = slides[22]
for sh in toc.shapes:
    if not sh.has_text_frame:
        continue
    tf = sh.text_frame
    t = tf.text
    if "[별첨]" in t:
        replace_in_runs(tf, [("안양물류센터 매입 프로젝트 l [별첨] 투자 대상 물건 설명자료",
                              "안양물류센터 l Information Memorandum")])
    if "Ⅳ. 모리츠" in t:
        # Ⅳ 모리츠 블록 삭제 → Ⅳ 리스크, Ⅴ 시설현황 순으로 재구성 (문단 서식은 기존 문단 복제)
        paras = list(tf.paragraphs)
        i4 = next(i for i, q in enumerate(paras) if q.text.startswith("Ⅳ. 모리츠"))
        i5 = next(i for i, q in enumerate(paras) if q.text.startswith("Ⅴ. 리스크"))
        head, sub, blank = paras[i5]._p, paras[i5 + 1]._p, paras[i4 - 1]._p

        def clone(el, txt=None):
            c = copy.deepcopy(el)
            if txt is not None:
                rs = c.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}r")
                ts = [r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t") for r in rs]
                ts[0].text = txt
                for x in ts[1:]:
                    x.text = ""
            return c
        new_els = [clone(head, "Ⅳ. 리스크 요인분석 및 관리방안"), clone(sub), clone(blank),
                   clone(head, "Ⅴ. 시설현황"), clone(sub, "     - 층별 도면 및 Floor Plan")]
        for q in paras[i4:]:
            q._p.getparent().remove(q._p)
        anchor = paras[i4 - 1]._p
        for el in new_els:
            anchor.addnext(el)
            anchor = el

# ---------------------------------------------------------------- 3) 마스터 꼬리말 · 레이아웃 머리말
for master in prs.slide_masters:
    for sh in master.shapes:
        if sh.has_text_frame and "투자심의위원회 설명자료" in sh.text_frame.text:
            replace_in_runs(sh.text_frame, [("안양물류센터 매입 프로젝트", "안양물류센터"),
                                            ("본 투자심의위원회 설명자료", "Information Memorandum")])
FLOOR_LAYOUT = None
for lay in prs.slide_layouts:
    for sh in lay.shapes:
        if not sh.has_text_frame:
            continue
        if lay.name == "4. 리스크 요인분석 및 관리방안":
            replace_in_runs(sh.text_frame, [("Ⅴ. 리스크", "Ⅳ. 리스크")])
        if lay.name == "1_4. 리스크 요인분석 및 관리방안" and "모리츠" in sh.text_frame.text:
            replace_in_runs(sh.text_frame, [(sh.text_frame.text.split("\n")[0], "Ⅴ. 시설현황")])
            FLOOR_LAYOUT = lay
assert FLOOR_LAYOUT is not None, "시설현황용 레이아웃을 찾지 못함"

# ---------------------------------------------------------------- 4) 슬라이드 삭제
sldIdLst = prs.slides._sldIdLst
for idx in range(len(slides), 0, -1):
    if idx not in KEEP:
        sld = sldIdLst[idx - 1]
        prs.part.drop_rel(sld.rId)
        sldIdLst.remove(sld)

# ---------------------------------------------------------------- 5) Ⅴ 시설현황 추가 (종료 장 앞)
doc = pymupdf.open(OLD_IM)
for pg in FLOOR_PAGES:
    page = doc[pg]
    r = page.rect
    clip = pymupdf.Rect(r.width * 0.03, r.height * 0.125, r.width * 0.985, r.height * 0.885)
    png = page.get_pixmap(dpi=220, clip=clip).tobytes("png")
    s = prs.slides.add_slide(FLOOR_LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    w_in = 10.75
    h_in = w_in * clip.height / clip.width
    s.shapes.add_picture(io.BytesIO(png), Inches((11.693 - w_in) / 2), Inches(1.15), Inches(w_in), Inches(h_in))
# 종료 장을 맨 끝으로
ids = list(sldIdLst)
closing_id = next(i for i in ids if prs.part.related_part(i.rId) is closing_slide.part)
sldIdLst.remove(closing_id)
sldIdLst.append(closing_id)

cp = prs.core_properties
cp.title = "안양물류센터 Information Memorandum v02"
prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
