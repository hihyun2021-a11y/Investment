import sys, copy, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
ws = wb['2026(예정)']
SH = '2026(예정)'
notice = [("공단1","인천광역시 남동구 남촌동 624",3302,999),("공단2","인천광역시 남동구 고잔동 673-10",3636,1100),
 ("부평공단","인천광역시 부평구 갈산동 181-8",1827,553),("산곡셀프","인천광역시 부평구 산곡동 411-16",1292,391),
 ("둔전셀프","경기도 용인시 처인구 둔전리 52-1, 11, 15, 18, 19",1304,394),("프리미엄셀프","경기도 화성시 동탄구 석우동 3-6",1676,507),
 ("송탄제일","경기도 평택시 장당동 477-2, 4, 6",1173,355),("새광주","경기도 광주시 쌍령동 113-8",1442,436),
 ("노은","대전광역시 유성구 노은동 534-11, 12",1323,400),("제천대로","충청북도 제천시 청전동 125-12",1383,418),
 ("조양","충청북도 제천시 천남동 30-3, 4, 5, 9, 16, 9-12",4886,1478),("교대셀프","충청북도 청주시 분평동 246-3, 33",3164,957),
 ("남대구IC","대구광역시 달서구 월성동 210, 210-6",2216,670),("경대셀프","대구광역시 북구 침산동 13-49, 50, 109, 117",1460,442),
 ("새화랑","경상북도 경주시 동천동 936",2391,723),("열린","경상북도 포항시 남구 연일읍 생지리 379-1, 5, 6, 8",2141,648),
 ("양산대로","경상남도 양산시 상북면 소토리 563",1850,560),("극동셀프","경상남도 창원시 마산회원구 석전동 270-32",2082,630),
 ("밀양현대","경상남도 밀양시 삼문동 358-110",1497,453)]
rows = {ws.cell(r, 6).value: r for r in range(5, 48) if ws.cell(r, 6).value and r not in (29,)}
def find(n):
    for k, r in rows.items():
        if k == n or k.replace('셀프', '') == n.replace('셀프', ''):
            return k, r
    raise KeyError(n)
match = [(i, n, a, m2, py) + find(n) for i, (n, a, m2, py) in enumerate(notice, 1)]
inc = {r: i for i, n, a, m2, py, k, r in match}

# ── 1) 2026(예정) 시트: AB열 '입찰공고' 상태 ─────────────────────────────
def clone(dst, srcc):
    dst.font = copy.copy(srcc.font); dst.fill = copy.copy(srcc.fill); dst.border = copy.copy(srcc.border)
    dst.alignment = copy.copy(srcc.alignment); dst.number_format = srcc.number_format; dst.protection = copy.copy(srcc.protection)
ws.column_dimensions['AB'].width = 13
for r in (3, 4):
    clone(ws.cell(r, 28), ws.cell(r, 27))
ws.merge_cells('AB3:AB4')
ws['AB3'] = "입찰공고\n('26.10.14)"
ws['AB3'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
GRAY = Font(name='KoPub돋움체 Medium', sz=10, color='FF808080')
for r in list(range(5, 29)) + list(range(30, 48)):
    c = ws.cell(r, 28); clone(c, ws.cell(r, 27))
    c.number_format = 'General'
    c.alignment = Alignment(horizontal='center', vertical='center')
    if r in inc:
        c.value = f"No.{inc[r]}"
    else:
        c.value = "미포함"; c.font = GRAY
for r, f in ((29, '=COUNTIF(AB5:AB28,"No.*")'), (48, '=COUNTIF(AB30:AB47,"No.*")'), (49, '=AB29+AB48')):
    c = ws.cell(r, 28); clone(c, ws.cell(r, 27)); c.value = f
    c.number_format = '0"건"'; c.alignment = Alignment(horizontal='center', vertical='center')

# ── 2) 신규 시트: 입찰공고 요약 ─────────────────────────────────────────
ns = wb.create_sheet('입찰공고(261014)', index=2)
F = lambda b=False, sz=10, color=None: Font(name='KoPub돋움체 Medium', sz=sz, b=b, color=color)
thin = Side(style='thin', color='FFBFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='FFE7E6E6'); SUB = PatternFill('solid', fgColor='FFFFFFCC')
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
NUM = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"_);_(@_)'
ns.sheet_view.showGridLines = False
ns['B1'] = '■ KLI 보유 주유소 매각 입찰공고 반영 (공고일 2026.10.14)'; ns['B1'].font = F(True, 11)
ns['B2'] = '출처: (주)코람코라이프인프라위탁관리부동산투자회사 「입찰 공고문 — 보유 자산 매각」(2026.10.14). 금액·면적은 2026(예정) 시트 값을 수식으로 참조함.'
ns['B2'].font = F(sz=9, color='FF595959')

def put(r, c, v, font=None, fill=None, al=CEN, fmt=None, border=True):
    cell = ns.cell(r, c, v); cell.font = font or F(); cell.alignment = al
    if fill: cell.fill = fill
    if fmt: cell.number_format = fmt
    if border: cell.border = BOX
    return cell

# 2-1 공고 개요
ns['B4'] = '1. 매각 개요'; ns['B4'].font = F(True)
info = [
 ('매도자', '(주)코람코라이프인프라위탁관리부동산투자회사'),
 ('매각 방식', '경쟁입찰(매수의향서 LOI 제출), 개별 매각 19개 자산. 유찰 시 수의계약 전환 가능'),
 ('예상 일정', "입찰공고 2026.10.14(수) → 입찰마감 2026.11.19(목) → 우선협상대상자 선정 2026.11월 중 → 매매계약 체결 2026.12월 말 → 소유권이전 2027.01월 이후 (매도자 사정에 따라 변경 가능)"),
 ('LOI 제출', '방문: 2026.11.19(목) 13:00~16:00, 코람코자산신탁 본사 8층 회의실 / 우편: 2026.11.19(목) 18:00 이전 도착분'),
 ('LOI 기재사항', '입찰참여자 정보, 희망 매수금액(건물분 VAT 별도), 자금조달계획(증빙), 매입 후 사용계획, 기타 매매조건·지급방법·일정'),
 ('입찰보증금', '건당 1억원(100,000,000원), 하나은행 183-910051-83204(예금주 KLI). 이행(입찰)보증보험(보증기간 최소 10개월)으로 대체 가능. 미납 시 LOI 무효, 우협 선정 후 계약 미체결·해제 시 미반환'),
 ('계약금 등', '매매계약 시 매매금액의 10% 계약금. MOU 체결 시 매매금액의 5% 이행보증금(우협 입찰보증금은 이행보증금에 충당)'),
 ('특약', '매매계약 체결과 동시에 임차인 에이치디현대오일뱅크㈜와 「토양오염검사 및 그 후속조치에 관한 합의서」 체결. 토지거래허가구역 소재 자산은 허가 후 계약'),
 ('매각주관사', '세빌스코리아(신규원 차장 010-7231-8291), (주)부동산플래닛(김재희 이사 010-5500-8779), 쿠시먼앤드웨이크필드코리아(정예원 차장 010-8695-6466) — CA·개인정보동의서는 3개사 중 택1 제출'),
]
r = 5
for k, v in info:
    put(r, 2, k, F(True), HEAD); ns.merge_cells(start_row=r, start_column=3, end_row=r, end_column=15)
    put(r, 3, v, al=LEFT)
    for c in range(4, 16): ns.cell(r, c).border = BOX
    ns.row_dimensions[r].height = 30 if len(v) > 70 else 18
    r += 1

# 2-2 대상 자산
r += 1; ns.cell(r, 2, '2. 입찰공고 대상 자산 (19개) ↔ 2026(예정) 시트 대조').font = F(True); r += 1
hdr = ['공고\nNo.', '공고 자산명', '시트 주유소명', '시트\nNO', '지역', '공고 상세주소(지번)', '공고\n대지면적(㎡)', '공고\n대지면적(평)',
       '시트\n대지면적(㎡)', '면적 차이\n(㎡)', '보증금', '장부가(계)', '감정평가(담보)', '상환비율', '회수 원본']
for j, h in enumerate(hdr): put(r, 2 + j, h, F(True), HEAD)
ns.row_dimensions[r].height = 30
h0 = r + 1
for i, n, a, m2, py, k, sr in match:
    r += 1
    put(r, 2, i); put(r, 3, n); put(r, 4, k)
    put(r, 5, f"='{SH}'!B{sr}"); put(r, 6, f"='{SH}'!E{sr}"); put(r, 7, a, al=LEFT)
    put(r, 8, m2, fmt='#,##0'); put(r, 9, py, fmt='#,##0')
    put(r, 10, f"='{SH}'!H{sr}", fmt='#,##0.0'); put(r, 11, f"=J{r}-H{r}", fmt='#,##0.0;[Red]-#,##0.0;"-"')
    put(r, 12, f"='{SH}'!L{sr}", fmt=NUM, al=Alignment(horizontal='right', vertical='center'))
    put(r, 13, f"='{SH}'!U{sr}", fmt=NUM, al=Alignment(horizontal='right', vertical='center'))
    put(r, 14, f"='{SH}'!Y{sr}", fmt=NUM, al=Alignment(horizontal='right', vertical='center'))
    put(r, 15, f"='{SH}'!Z{sr}", fmt='0.0%')
    put(r, 16, f"='{SH}'!AA{sr}", fmt=NUM, al=Alignment(horizontal='right', vertical='center'))
h1 = r; r += 1
put(r, 2, '합계 (19개)', F(True), SUB); ns.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
for c in range(3, 8): ns.cell(r, c).border = BOX; ns.cell(r, c).fill = SUB
for c, fmt in ((8, '#,##0'), (9, '#,##0'), (10, '#,##0.0'), (11, '#,##0.0;[Red]-#,##0.0;"-"'), (12, NUM), (13, NUM), (14, NUM), (16, NUM)):
    put(r, c, f"=SUM({L(c)}{h0}:{L(c)}{h1})", F(True), SUB, Alignment(horizontal='right', vertical='center'), fmt)
put(r, 15, '', F(True), SUB)
tot = r

# 2-3 공고 미포함
r += 2; ns.cell(r, 2, "3. 2026(예정) '신규' 중 이번 공고에 미포함된 자산").font = F(True); r += 1
for c, h in ((2, '시트\nNO'), (3, '지역'), (4, '시트 주유소명'), (5, '상세주소'), (12, '보증금'), (13, '장부가(계)'), (14, '감정평가(담보)'), (16, '회수 원본')):
    put(r, c, h, F(True), HEAD)
ns.merge_cells(start_row=r, start_column=5, end_row=r, end_column=11)
for c in range(6, 12): ns.cell(r, c).border = BOX; ns.cell(r, c).fill = HEAD
ns.row_dimensions[r].height = 30
e0 = r + 1
for sr in range(5, 29):
    if sr in inc: continue
    r += 1
    put(r, 2, f"='{SH}'!B{sr}"); put(r, 3, f"='{SH}'!E{sr}"); put(r, 4, f"='{SH}'!F{sr}"); put(r, 5, f"='{SH}'!G{sr}", al=LEFT)
    ns.merge_cells(start_row=r, start_column=5, end_row=r, end_column=11)
    for c in range(6, 12): ns.cell(r, c).border = BOX
    for c, col in ((12, 'L'), (13, 'U'), (14, 'Y'), (16, 'AA')):
        put(r, c, f"='{SH}'!{col}{sr}", fmt=NUM, al=Alignment(horizontal='right', vertical='center'))
e1 = r; r += 1
put(r, 2, '합계 (5개)', F(True), SUB); ns.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
for c in range(3, 12): ns.cell(r, c).border = BOX; ns.cell(r, c).fill = SUB
for c in (12, 13, 14, 16):
    put(r, c, f"=SUM({L(c)}{e0}:{L(c)}{e1})", F(True), SUB, Alignment(horizontal='right', vertical='center'), NUM)
r += 2
notes = ["※ '미매각' 18개 자산은 이번 공고 대상에 없음.",
         "※ 공고 대지면적은 소수점 이하 반올림 표기. 시트 대지면적과의 차이는 모두 ±0.4㎡ 이내(반올림 차이).",
         "※ 공고 주소는 지번, 시트 주소는 도로명 기준. 자산명 표기 차이: 부평공단↔부평공단셀프, 송탄제일↔송탄제일셀프, 새광주↔새광주셀프, 제천대로↔제천대로셀프, 남대구IC↔남대구IC셀프.",
         "※ 회수 원본 = 장부가 − 보증금 − 감정평가(담보) × 상환비율 (2026(예정) 시트 AA열 산식). 실제 회수액은 낙찰가·정산에 따라 달라짐."]
for t in notes:
    ns.cell(r, 2, t).font = F(sz=9, color='FF595959'); r += 1
widths = {'A': 2, 'B': 7, 'C': 13, 'D': 14, 'E': 6, 'F': 7, 'G': 40, 'H': 11, 'I': 10, 'J': 11, 'K': 10, 'L': 14, 'M': 16, 'N': 16, 'O': 9, 'P': 16}
for k, v in widths.items(): ns.column_dimensions[k].width = v
ns.freeze_panes = 'B5'
wb.save(out)
print('ok', 'total row', tot, 'included rows', sorted(inc))
