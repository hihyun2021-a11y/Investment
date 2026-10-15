# -*- coding: utf-8 -*-
"""회계법인 사업타당성 검토보고서 RFP(이메일 초안) Word 파일 생성: python3 build_RFP_사업타당성검토_261015.py"""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "20261015_RFP_사업타당성검토보고서_회계법인_v01.docx")
FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.0)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.line_spacing = 1.25


def para(text="", bold=False, size=None, color=None, align=None, after=None, before=None, indent=None):
    p = doc.add_paragraph()
    if text:
        add_runs(p, text, bold=bold, size=size, color=color)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def add_runs(p, text, bold=False, size=None, color=None):
    """'**굵게**' 표기를 굵은 글씨로 바꿔 넣는다."""
    parts = text.split("**")
    for i, t in enumerate(parts):
        if not t:
            continue
        r = p.add_run(t)
        r.bold = bold or (i % 2 == 1)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color


def bullet(text, level=0, num=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6 + 0.6 * level)
    pf.first_line_indent = Cm(-0.4)
    pf.space_after = Pt(2)
    mark = f"{num}. " if num is not None else ("- " if level else "• ")
    add_runs(p, mark + text)
    return p


def heading(text):
    p = para(text, bold=True, size=11, color=NAVY, before=10, after=4)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F3A5F")):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    pPr.append(bdr)
    return p


def subheading(text):
    return para(text, bold=True, size=10, before=6, after=2)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def set_widths(table, widths_cm):
    """표 열 너비를 고정한다(Word·LibreOffice 공통)."""
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for r in table.rows:
        for c, w in zip(r.cells, widths_cm):
            c.width = Cm(w)


# ---------------------------------------------------------------- 머리 메모(작성 안내)
note = para("※ 작성 안내: 거래 개요 수치는 기준 재무모델 v05 및 투심자료 KLI 최종본 v5.1 기준이며, [ ] 표시는 발송 전 기입할 부분입니다.",
            size=8.5, color=GREY, after=8)

# ---------------------------------------------------------------- 제목·수신
para("RFP (제안요청서)", bold=True, size=9, color=GREY, after=0)
para("안양물류센터 매입 관련 사업타당성 검토보고서 작성 용역 제안 요청", bold=True, size=15, color=NAVY, after=2)
para("코람코자산신탁 상장리츠팀 | 2026년 10월", size=9, color=GREY, after=10)

meta = doc.add_table(rows=3, cols=2)
meta.style = "Table Grid"
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate([("제목", "[RFP] 안양물류센터 매입 관련 사업타당성 검토보고서 작성 용역 제안 요청 (코람코자산신탁)"),
                            ("수신", "[회계법인명] [담당 파트너/이사]님"),
                            ("발신", "코람코자산신탁 상장리츠팀 [성명/직급]")]):
    c0, c1 = meta.rows[i].cells
    c0.text, c1.text = "", ""
    add_runs(c0.paragraphs[0], k, bold=True)
    add_runs(c1.paragraphs[0], v)
    shade(c0, "E8EEF5")
    for c in (c0, c1):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c0.width, c1.width = Cm(2.6), Cm(14.0)
set_widths(meta, [2.6, 14.0])
para(after=4)

para("[회계법인명] [담당 파트너/이사]님께,", after=6)
para("안녕하십니까. 코람코자산신탁 상장리츠팀 [성명]입니다.")
para("당사가 자산관리회사(AMC)로서 추진 중인 안양물류센터 매입 건과 관련하여, 매입의 적정성을 독립적으로 검토하는 사업타당성 검토보고서 작성 용역을 의뢰드리고자 합니다. "
     "아래 내용을 검토하시어 제안서와 견적을 보내주시면 감사하겠습니다.", after=6)

# ---------------------------------------------------------------- 1
heading("1. 배경 및 목적")
bullet("본 건은 신설 사모 위탁관리리츠인 코람코라이프로지스위탁관리부동산투자회사(이하 \"KLL리츠\")가 안양물류센터를 매입하는 거래입니다.")
bullet("당사가 운용하는 상장리츠인 코람코라이프인프라위탁관리부동산투자회사(이하 \"KLI리츠\")가 KLL리츠에 보통주와 종류주식으로 출자합니다.")
bullet("매도인은 코크렙안양㈜이며, 지분은 ㈜엘에프 95%, 당사 5%입니다. 당사는 매도인 주주이면서 매수 리츠의 AMC라서, 본 건은 **이해관계자 간 거래**에 해당할 수 있습니다.")
bullet("이에 다음 네 가지를 독립적인 제3자 관점에서 검토받아, KLI리츠 주주와 이사회, 필요 시 감독당국에 소명할 근거 자료로 쓰고자 합니다.")
for t in ("매입가격의 적정성", "수익률 산정", "적용가정의 타당성", "Exit 가정의 타당성"):
    bullet(t, level=1)

# ---------------------------------------------------------------- 2
heading("2. 대상 거래 개요")
rows = [
    ("대상 자산", "안양물류센터. 경기도 안양시 동안구 관양동 934 외 2필지, 2023-11 준공 상온 물류센터, 지하 2층~지상 8층, 연면적 95,474.59㎡(28,881평)"),
    ("임대 현황", "임대율 99.3%. 주요 임차인 쿠팡(임대면적 약 77%), 농심 외. WALE 약 6.7년(거래종결일 기준)"),
    ("매입가", "2,700억원(평당 약 935만원). 담보 감정가 3,345억원(미래새한, 2026-09) 대비 −19.3%"),
    ("총 사업비·재원", "총 2,989억원 = 선순위 대출 1,600억 + 중순위 대출 380억 + Equity 988억(1종 240·2종 500·3종 100·보통주 148) + 보증금 21억"),
    ("운용 기간", "약 10년(2026-10 ~ 2036-10). 24개월 후 리파이낸싱 가정"),
    ("주요 가정(당사 모델)", "Going-in·Exit Cap Rate 4.89%, 전체 Equity IRR 10.21% [기준 재무모델 v05 기준, 최종 모델 제공 예정]"),
    ("일정", "매매계약·대출약정 체결 2026-10-16(예정), 거래종결 2026-10-22(예정)"),
]
t = doc.add_table(rows=len(rows) + 1, cols=2)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, h in enumerate(("항목", "내용")):
    c = t.rows[0].cells[j]
    c.text = ""
    add_runs(c.paragraphs[0], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(c, "1F3A5F")
for i, (k, v) in enumerate(rows, start=1):
    c0, c1 = t.rows[i].cells
    c0.text, c1.text = "", ""
    add_runs(c0.paragraphs[0], k, bold=True)
    add_runs(c1.paragraphs[0], v)
    shade(c0, "E8EEF5")
for r in t.rows:
    r.cells[0].width, r.cells[1].width = Cm(3.6), Cm(13.0)
    for c in r.cells:
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(1)
set_widths(t, [3.6, 13.0])
para(after=2)

# ---------------------------------------------------------------- 3
heading("3. 용역 범위 (Scope of Work)")
subheading("(1) 매입가격의 적정성")
bullet("감정평가서(미래새한, 담보 목적)의 평가 방법, 주요 가정, 결론 검토")
bullet("권역 내(수도권 남부) 물류센터 거래 사례 비교: 평당 가격, Cap Rate, 임대료 수준")
bullet("수익환원법(DCF 및 직접환원)에 의한 독립적 가치 범위 산정, 매입가와 비교")
subheading("(2) 수익률 산정의 적정성")
bullet("재무모델의 계산 정확성과 구조 검토: 현금흐름, 배당재원, 트랜치별 IRR·CoC, DSCR")
bullet("모델 산출 수익률과 KLI리츠 투자 수익률의 정합성 확인")
subheading("(3) 적용가정의 타당성")
bullet("임대료: 현행 임대료 대비 시장임대료, 상승률, 만기 갱신 가정(쿠팡 2029년 중도해지 옵션 포함), 공실·Rent-Free 가정")
bullet("운영비용, 자본적 지출(CAPEX), 보험료, 제세공과금")
bullet("금융비용: 선·중순위 금리, 수수료, 24개월 후 리파이낸싱 가정")
bullet("매입 부대비용, 세금(취득세 등)")
subheading("(4) Exit 가정의 타당성")
bullet("Exit Cap Rate(4.89%), 매각 시점, 매각 비용 가정이 시장 수준에 비추어 합리적인지")
bullet("Going-in Cap Rate와 같은 Exit Cap을 적용한 것이 적정한지")
subheading("(5) 민감도·시나리오 분석")
bullet("핵심 변수의 변동이 매입가 적정성과 수익률에 미치는 영향")
bullet("변수: 임대료, 공실, Exit Cap ±25~50bp, 금리, 쿠팡 중도해지 시나리오", level=1)
subheading("(6) 이해관계자 거래 관점의 검토 의견")
bullet("거래가격 결정 과정과 거래조건이 독립 당사자 간 거래(Arm's length)와 비교해 합리적인지에 대한 의견")
bullet("귀 법인이 이 의견을 공정성 의견(Fairness Opinion) 형식으로 낼 수 있는지와 그 범위도 함께 제안해 주시기 바랍니다.")

# ---------------------------------------------------------------- 4
heading("4. 산출물 및 일정")
bullet("**산출물:**")
bullet("사업타당성 검토보고서(국문): 요약 보고서(Executive Summary) + 본문", level=1)
bullet("보고 자료(KLI리츠 이사회·주주 설명용 요약본)", level=1)
bullet("**평가 기준일:** [2026-10-XX(거래종결일 또는 계약 체결일)] — 귀 법인 의견을 받아 협의")
bullet("**희망 일정:** 착수 [10/XX] → 초안 [착수 후 3주 이내] → 최종 [초안 후 1주 이내]. 가능한 일정을 제안 부탁드립니다.")
bullet("**보고서 이용 범위:** KLI리츠 주주·이사회 제공, 감독당국 제출, 필요 시 공시 인용. 이에 따른 Reliance Letter 또는 이용 동의 조건을 함께 제시해 주시기 바랍니다.")

# ---------------------------------------------------------------- 5
heading("5. 제공 예정 자료 (비밀유지약정 체결 후 제공)")
bullet("재무모델(최종본), 투자심의 자료, IM")
bullet("감정평가서(미래새한), 시장보고서, 물리실사(CBRE)·법률실사(태평양)·재무실사 보고서")
bullet("부동산매매계약서, 임대차계약서 일체, 렌트롤, 임대보증금 내역")
bullet("대출 조건(Term Sheet) 및 대출약정서 주요 조건")

# ---------------------------------------------------------------- 6
heading("6. 제안서 포함 요청 사항")
bullet("**독립성 확인:** 아래 각 당사자와의 감사, 자문, 기타 이해관계 유무", num=1)
bullet("당사(코람코자산신탁), KLI리츠, KLL리츠, 매도인(코크렙안양㈜), ㈜엘에프", level=1)
bullet("특히 위 회사들의 외부감사인인 경우 수행이 어려울 수 있으니 반드시 알려주시기 바랍니다.", level=1)
bullet("**수행팀 구성 및 유사 실적:** 리츠, 물류센터, 이해관계자 거래 관련 가치평가·타당성 검토 실적", num=2)
bullet("**수행 방법론 및 세부 일정**", num=3)
bullet("**용역 보수:** 부가세 별도, 실비 포함 여부, 지급 조건", num=4)
bullet("**보고서의 책임 범위 및 이용 제한 조건**", num=5)

# ---------------------------------------------------------------- 7
heading("7. 제출 안내")
bullet("제출 기한: 2026년 [10월 XX일(X)] [17:00]까지")
bullet("제출 방법: 본 메일 회신(PDF)")
bullet("문의: 코람코자산신탁 상장리츠팀 [성명/직급], [전화], [이메일]")

para(before=8)
para("본 메일과 첨부 정보는 대외비입니다. 제안 검토 목적으로만 사용해 주시고 제3자에게 공개하지 말아 주시기 바랍니다.", size=9, color=GREY, after=8)
para("감사합니다.", after=8)
para("[성명] 드림")
para("코람코자산신탁 상장리츠팀", after=14)

# ---------------------------------------------------------------- 내부 참고(발송 전 삭제)
box = doc.add_table(rows=1, cols=1)
box.style = "Table Grid"
box.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = box.rows[0].cells[0]
shade(cell, "FFF4E5")
set_widths(box, [16.6])
p0 = cell.paragraphs[0]
add_runs(p0, "[내부 참고 — 발송 전 삭제] 보내기 전에 확인하실 점", bold=True, color=RGBColor(0x9C, 0x57, 0x00))
for txt in (
    "**우리회계법인:** 이미 재무실사를 맡은 우리회계법인은 독립성 면에서 불리할 수 있습니다. 새 회계법인 2~3곳에 동시에 보내 비교 견적을 받는 것을 권합니다.",
    "**평가 기준일:** 거래종결일 이후 작성을 검토하실 수 있습니다. 다만 당국과 주주에게 소명하려면 계약 체결 전 시점 기준이 더 유리합니다. 일정이 허락하면 체결 전 착수를 권합니다.",
    "**최신 모델과 수치 차이:** 선순위 금리를 반영한 최신 모델(1713 버전)은 아직 기준 모델로 지정되지 않았습니다. 그래서 메일의 IRR 등은 v05 기준입니다. 자료를 줄 때 최종 모델로 통일하시면 됩니다.",
):
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(2)
    add_runs(p, "• " + txt, size=9)

doc.core_properties.title = "안양물류센터 사업타당성 검토보고서 RFP"
doc.core_properties.author = "코람코자산신탁 상장리츠팀"
doc.save(OUT)
print("saved", OUT)
