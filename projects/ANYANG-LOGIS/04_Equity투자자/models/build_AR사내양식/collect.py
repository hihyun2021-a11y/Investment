import openpyxl, json, datetime
wbv=openpyxl.load_workbook('lo_out/calc.xlsx', data_only=True, read_only=True)['A&R(사내양식)']
wbf=openpyxl.load_workbook('stage_a.xlsx')['A&R(사내양식)']
formulas={c.coordinate for row in wbf.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('=')}
vals={}
for row in wbv.iter_rows():
    for c in row:
        if hasattr(c,'coordinate') and c.coordinate in formulas:
            v=c.value
            if isinstance(v,datetime.datetime):
                d=v-datetime.datetime(1899,12,30); v=d.days+d.seconds/86400
            vals[c.coordinate]=v
miss=[f for f in formulas if vals.get(f) is None]
errs={k:v for k,v in vals.items() if isinstance(v,str) and v.startswith('#')}
print('formulas',len(formulas),'missing',sorted(miss)[:20],'errs',errs)
json.dump(vals, open('values.json','w'), ensure_ascii=False)
