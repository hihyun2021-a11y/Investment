# -*- coding: utf-8 -*-
"""생성 결과(Markdown)를 Word(.docx)로 변환한다. 제목·표·목록·굵게(**)를 지원한다."""
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt

FONT = "맑은 고딕"


def _runs(par, text: str) -> None:
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            r = par.add_run(part)
            r.bold = i % 2 == 1


def _style(doc) -> None:
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        h = doc.styles[name]
        h.font.name = FONT
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def md_to_docx(md: str) -> bytes:
    doc = Document()
    _style(doc)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    tbl.append(cells)
                i += 1
            if tbl:
                n = max(len(r) for r in tbl)
                t = doc.add_table(rows=len(tbl), cols=n)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, r in enumerate(tbl):
                    for ci in range(n):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        _runs(cell.paragraphs[0], r[ci] if ci < len(r) else "")
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 3))
        elif re.match(r"^\s*[-*]\s+", line):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _runs(doc.add_paragraph(), line.strip())
        elif line.strip() in ("---", "***"):
            doc.add_paragraph("─" * 30)
        elif line.strip():
            _runs(doc.add_paragraph(), line)
        else:
            doc.add_paragraph()
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
