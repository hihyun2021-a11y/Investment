# -*- coding: utf-8 -*-
"""샘플 서류에서 텍스트를 뽑는다(PDF·DOCX·PPTX·XLSX·CSV·MD·TXT·EML·HTML). 결과는 data/cache에 캐시한다.
스캔 PDF처럼 글자가 거의 없으면 needs_pdf=True로 표시해, Claude에 PDF 원본을 직접 보내도록 한다."""
import email
import hashlib
import html
import re
from email import policy
from pathlib import Path

from .config import DATA, MAX_SAMPLE_CHARS

TEXT_EXT = {".md", ".txt", ".csv", ".json"}


def _pdf(p: Path) -> str:
    import pymupdf
    d = pymupdf.open(p)
    return "\n".join(f"\n[p.{i + 1}]\n{pg.get_text()}" for i, pg in enumerate(d))


def _docx(p: Path) -> str:
    import docx
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        parts.append("\n[표]")
        for r in t.rows:
            parts.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(parts)


def _pptx(p: Path) -> str:
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(p).slides, 1):
        out.append(f"\n[slide {i}]")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    out.append(" | ".join(c.text for c in r.cells))
    return "\n".join(out)


def _xlsx(p: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n[sheet {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _eml(p: Path) -> str:
    msg = email.message_from_bytes(p.read_bytes(), policy=policy.default)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        text = _strip_html(text)
    return f"From: {msg['from']}\nTo: {msg['to']}\nDate: {msg['date']}\nSubject: {msg['subject']}\n\n{text}"


def _strip_html(s: str) -> str:
    s = re.sub(r"(?is)<(script|style).*?</\1>", "", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</h\d>|</li>|</tr>", "\n", s)
    return html.unescape(re.sub(r"<[^>]+>", " ", s))


def extract(path: Path) -> dict:
    """{'text': str, 'needs_pdf': bool, 'truncated': bool}"""
    path = Path(path)
    key = hashlib.sha1(f"{path.resolve()}|{path.stat().st_mtime}|{path.stat().st_size}".encode()).hexdigest()
    cache = DATA / "cache" / f"{key}.txt"
    ext = path.suffix.lower()
    if cache.exists():
        text = cache.read_text(encoding="utf-8")
    else:
        try:
            if ext == ".pdf":
                text = _pdf(path)
            elif ext in (".docx", ".docm"):
                text = _docx(path)
            elif ext == ".pptx":
                text = _pptx(path)
            elif ext in (".xlsx", ".xlsm"):
                text = _xlsx(path)
            elif ext == ".eml":
                text = _eml(path)
            elif ext in (".html", ".htm"):
                text = _strip_html(path.read_text(encoding="utf-8", errors="ignore"))
            elif ext in TEXT_EXT:
                text = path.read_text(encoding="utf-8", errors="ignore")
            else:
                text = f"[이 형식({ext})은 텍스트 추출을 지원하지 않습니다]"
        except Exception as e:  # 손상 파일 등
            text = f"[텍스트 추출 실패: {e}]"
        cache.write_text(text, encoding="utf-8")
    needs_pdf = ext == ".pdf" and len(re.sub(r"\[p\.\d+\]|\s", "", text)) < 200
    truncated = len(text) > MAX_SAMPLE_CHARS
    return {"text": text[:MAX_SAMPLE_CHARS], "needs_pdf": needs_pdf, "truncated": truncated}
