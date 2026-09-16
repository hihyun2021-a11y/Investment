#!/usr/bin/env python3
"""투자심의 자료(.pptx) 생성기 — 사내 표준 양식.

규격: platform/standards/ic_deck_style.md
빈 템플릿: platform/templates/ic_deck/투자심의_빈템플릿.pptx (레이아웃·로고·머리말 유지)

사용 예:
    from ic_deck import ICDeck
    d = ICDeck(asset="안양물류센터", meeting="투자심의위원회 심의자료",
               headline="리츠투자부문-상장리츠본부-상장리츠팀 ㅣ신규 투자의 건",
               date="October. 2026")
    d.title_slide()
    d.contents_slide(["투자심의위원회 심의 개요", "안건 Ⅰ. ...", "안건 Ⅱ. ..."])
    s = d.content_slide("사업개요", "1. 투자자산 개요",
                        lead="연면적 28,881평 규모의 신축 프라임 상온 물류센터임",
                        subs=["9,500kW 전기용량·자동화 설비 — 우량 3PL 임차인 선호 스펙"])
    d.band(s, "대상자산 개요")
    d.table(s, [["구 분", "내 용"], ["물건명", "안양물류센터"]], col_widths=[2.2, 8.4])
    d.footnote(s, "* 출처: 00_원본자료/processed/20260916_핵심지표_IM기준_v01.md")
    d.save("투자심의자료.pptx")

CLI:
    python3 platform/scripts/ic_deck.py --demo out.pptx
"""
from __future__ import annotations

import pathlib
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parents[1].parent
TEMPLATE = ROOT / "platform/templates/ic_deck/투자심의_빈템플릿.pptx"

# ── 색상 (ic_deck_style.md §3) ────────────────────────────────────────────
NAVY = RGBColor(0x01, 0x23, 0x53)      # 브랜드 네이비
NAVY_DARK = RGBColor(0x00, 0x20, 0x60)
BAND = RGBColor(0x34, 0x5B, 0x86)      # 섹션 밴드
INK = RGBColor(0x1F, 0x1F, 0x1F)       # 본문 먹색
GRAY = RGBColor(0x40, 0x40, 0x40)
RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x6D, 0x52)
TH_BG = RGBColor(0xDA, 0xE3, 0xF3)     # 표 머리 배경
ROW_BG = RGBColor(0xF2, 0xF2, 0xF2)    # 교차 음영
LINE = RGBColor(0xD9, 0xD9, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ── 서체 ──────────────────────────────────────────────────────────────────
F_LIGHT = "KoPubWorld돋움체 Light"
F_MEDIUM = "KoPubWorld돋움체 Medium"
F_BOLD = "KoPubWorld돋움체 Bold"

# ── 좌표 (inch, ic_deck_style.md §1) ──────────────────────────────────────
TITLE_TOP, TITLE_LEFT, TITLE_W, TITLE_H = 0.24, 0.0, 6.52, 0.58
TITLE_LINS = 1.81          # 제목 텍스트 좌측 내부 여백 (머리말 섹션명 뒤에서 시작)
LEAD_TOP, LEAD_LEFT, LEAD_W = 1.14, 0.31, 10.93
BAND_TOP, BAND_LEFT, BAND_W, BAND_H = 2.17, 0.50, 10.70, 0.27
BODY_TOP, BODY_LEFT, BODY_W = 2.58, 0.57, 10.70
BODY_BOTTOM = 7.40         # 본문이 넘지 말아야 할 하한
NOTE_TOP = 7.45

# ── 섹션 → 레이아웃 (머리말이 레이아웃에 내장돼 있음) ──────────────────────
SECTION_LAYOUT = {
    "표지": "4_사용자 지정 레이아웃",
    "간지": "4_사용자 지정 레이아웃",
    "심의안건": "9_흰배경",
    "심의안건2": "8_흰배경",
    "추진배경": "7_흰배경",
    "하이라이트": "5_흰배경",
    "사업개요": "1. 투자개요",
    "투자계획": "2. 투자계획",
    "운영계획": "1_3. 운영 및 Exit Plan",
    "리스크": "4. 리스크 요인분석 및 관리방안",
    "향후조치": "1_4. 리스크 요인분석 및 관리방안",
    "부록": "10_흰배경",
    "부록2": "11_흰배경",
}

DISCLAIMER = ("This report is solely for the use of client personnel. No part of it may be circulated, "
              "quoted, or reproduced for distribution outside the client or organization without prior "
              "written approval from KORAMCO REITs Management and Trust Co., Ltd.")
TEAM_EN = ("Listed REITs Team │ Listed REITs Department │ REITs Investment Division│ "
           "KORAMCO REITs Management and Trust Co., Ltd.")


def _run(para, text, *, font=F_LIGHT, size=11, bold=False, color=INK):
    r = para.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return r


class ICDeck:
    """사내 표준 양식의 투자심의 자료를 만든다."""

    def __init__(self, asset: str, meeting: str = "투자심의위원회 심의자료",
                 headline: str = "리츠투자부문-상장리츠본부-상장리츠팀 ㅣ신규 투자의 건",
                 date: str = "", template: str | pathlib.Path = TEMPLATE):
        template = pathlib.Path(template)
        if not template.exists():
            raise FileNotFoundError(f"빈 템플릿을 찾을 수 없습니다: {template}")
        self.prs = Presentation(str(template))
        self.asset, self.meeting, self.headline, self.date = asset, meeting, headline, date
        self._layouts = {l.name: l for m in self.prs.slide_masters for l in m.slide_layouts}

    # ── 내부 도구 ────────────────────────────────────────────────────────
    def _slide(self, section: str):
        name = SECTION_LAYOUT.get(section, section)
        if name not in self._layouts:
            raise KeyError(f"레이아웃 없음: {name} (사용 가능: {sorted(self._layouts)})")
        return self.prs.slides.add_slide(self._layouts[name])

    @staticmethod
    def _box(slide, left, top, width, height, *, word_wrap=True):
        tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = tb.text_frame
        tf.word_wrap = word_wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return tb, tf

    # ── 슬라이드 유형 ────────────────────────────────────────────────────
    def title_slide(self):
        """표지."""
        s = self._slide("표지")
        _, tf = self._box(s, 9.18, 0.70, 2.51, 0.27)
        _run(tf.paragraphs[0], "< Strictly Confidential >", size=9, color=GRAY)

        _, tf = self._box(s, 0.41, 2.20, 9.25, 2.18)
        _run(tf.paragraphs[0], self.headline, font=F_MEDIUM, size=12, color=NAVY)
        p = tf.add_paragraph()
        _run(p, self.asset, font=F_BOLD, size=24, color=NAVY)
        p = tf.add_paragraph()
        _run(p, self.meeting, font=F_MEDIUM, size=16, color=INK)

        _, tf = self._box(s, 0.41, 6.06, 10.62, 0.27)
        _run(tf.paragraphs[0], TEAM_EN, size=8, color=GRAY)
        if self.date:
            _, tf = self._box(s, 0.41, 6.32, 3.0, 0.27)
            _run(tf.paragraphs[0], self.date, size=9, color=GRAY)
        _, tf = self._box(s, 0.41, 7.07, 10.87, 0.34)
        _run(tf.paragraphs[0], DISCLAIMER, size=7, color=GRAY)
        return s

    def contents_slide(self, items: list[str], title: str | None = None):
        """목차. title을 주면 좌상단 머리 문구를 바꾼다 (예: 별첨 목차)."""
        s = self._slide("표지")
        _, tf = self._box(s, 0.47, 2.27, 6.5, 0.29)
        _run(tf.paragraphs[0], title or self.meeting, font=F_MEDIUM, size=11, color=NAVY)
        ln = s.shapes.add_connector(1, Inches(2.50), Inches(2.40), Inches(11.20), Inches(2.40))
        ln.line.color.rgb = LINE
        ln.line.width = Pt(0.75)
        _, tf = self._box(s, 0.89, 3.04, 3.0, 0.5)
        _run(tf.paragraphs[0], "Contents", font=F_BOLD, size=20, color=NAVY)
        _, tf = self._box(s, 7.89, 3.04, 3.55, 0.4 * len(items))
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            _run(p, it, size=11, color=INK)
        return s

    def divider_slide(self, section_title: str, toc: list[str], active: str = ""):
        """간지 — 전체 목차에서 현재 섹션만 강조."""
        s = self._slide("간지")
        _, tf = self._box(s, 0.41, 2.20, 6.0, 0.6)
        _run(tf.paragraphs[0], section_title, font=F_BOLD, size=20, color=NAVY)
        _, tf = self._box(s, 0.41, 3.10, 6.5, 0.32 * len(toc))
        for i, it in enumerate(toc):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            on = (it == active)
            _run(p, it, font=F_BOLD if on else F_LIGHT, size=12 if on else 11,
                 color=NAVY if on else GRAY)
        return s

    def content_slide(self, section: str, title: str, lead: str = "", subs: list[str] | None = None):
        """본문 슬라이드 — 제목 + 리드(16pt 결론 + 12pt 불릿)."""
        s = self._slide(section)
        tb, tf = self._box(s, TITLE_LEFT, TITLE_TOP, TITLE_W, TITLE_H)
        tf.margin_left = Emu(int(TITLE_LINS * 914400))
        _run(tf.paragraphs[0], f"> {title}", size=11, bold=True, color=NAVY_DARK)
        if lead or subs:
            _, tf = self._box(s, LEAD_LEFT, LEAD_TOP, LEAD_W, 0.77)
            first = True
            if lead:
                _run(tf.paragraphs[0], lead, size=16, color=INK)
                first = False
            for sub in (subs or []):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                _run(p, f"• {sub}", size=12, color=GRAY)
        return s

    def band(self, slide, label: str, top: float = BAND_TOP, width: float = BAND_W):
        """섹션 구분 밴드 (파란 띠 + 흰 라벨)."""
        from pptx.enum.shapes import MSO_SHAPE
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(BAND_LEFT), Inches(top),
                                    Inches(width), Inches(BAND_H))
        sh.fill.solid()
        sh.fill.fore_color.rgb = BAND
        sh.line.fill.background()
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.margin_left = Inches(0.12)
        tf.margin_top = tf.margin_bottom = 0
        _run(tf.paragraphs[0], label, font=F_BOLD, size=10, color=WHITE)
        return sh

    def table(self, slide, rows: list[list[str]], *, top: float = BODY_TOP, left: float = BODY_LEFT,
              width: float = BODY_W, col_widths: list[float] | None = None,
              row_height: float = 0.30, size: int = 10, header: bool = True, zebra: bool = True):
        """표 — 머리행 연블루 배경 + Bold, 본문 Light."""
        nr, nc = len(rows), max(len(r) for r in rows)
        height = min(row_height * nr, BODY_BOTTOM - top)
        gf = slide.shapes.add_table(nr, nc, Inches(left), Inches(top), Inches(width), Inches(height))
        tbl = gf.table
        if col_widths:
            for i, w in enumerate(col_widths[:nc]):
                tbl.columns[i].width = Inches(w)
        for ri, row in enumerate(rows):
            tbl.rows[ri].height = Inches(row_height)
            for ci in range(nc):
                cell = tbl.cell(ri, ci)
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = 0
                cell.fill.solid()
                if header and ri == 0:
                    cell.fill.fore_color.rgb = TH_BG
                elif zebra and ri % 2 == 0:
                    cell.fill.fore_color.rgb = ROW_BG
                else:
                    cell.fill.fore_color.rgb = WHITE
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER if (header and ri == 0) or ci == 0 else PP_ALIGN.LEFT
                txt = row[ci] if ci < len(row) else ""
                _run(p, txt, font=F_BOLD if (header and ri == 0) else F_LIGHT,
                     size=size, color=INK)
        return tbl

    def bullets(self, slide, items: list[str], *, top: float = BODY_TOP, left: float = BODY_LEFT,
                width: float = BODY_W, size: int = 11):
        """본문 불릿 목록."""
        _, tf = self._box(slide, left, top, width, 0.3 * len(items))
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            _run(p, f"• {it}", size=size, color=INK)
        return tf

    def footnote(self, slide, text: str, *, top: float = NOTE_TOP, size: int = 8):
        """각주 — 출처·단서 문구."""
        _, tf = self._box(slide, 0.31, top, 10.95, 0.30)
        _run(tf.paragraphs[0], text, size=size, color=GRAY)
        return tf

    def closing_slide(self):
        """종료 페이지 — 영문 배포제한 문구."""
        s = self._slide("표지")
        _, tf = self._box(s, 4.07, 3.74, 6.70, 0.33)
        _run(tf.paragraphs[0], DISCLAIMER, size=7, color=GRAY)
        return s

    def save(self, path: str | pathlib.Path):
        path = pathlib.Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


def _demo(out: str):
    d = ICDeck(asset="샘플 물류센터", meeting="투자심의위원회 심의자료", date="October. 2026")
    toc = ["투자심의위원회 심의 개요", "Ⅰ. 사업개요", "Ⅱ. 투자계획", "Ⅲ. 운영계획", "Ⅳ. 리스크 요인분석"]
    d.title_slide()
    d.contents_slide(toc)
    s = d.content_slide("심의안건", "심의안건 및 추진일정",
                        lead="본건은 매매계약 체결 및 주주간계약 체결을 위한 심의 건임",
                        subs=["독점적 협상기간 만료 전 거래종결을 목표로 추진"])
    d.band(s, "심의안건")
    d.table(s, [["구 분", "안 건", "비 고"],
                ["안건 1", "매매계약 체결의 건", "협의 중"],
                ["안건 2", "주주간계약 체결의 건", "협의 중"]],
            col_widths=[1.6, 6.5, 2.6])
    d.footnote(s, "※ 상기 내용 중 일부는 협의 과정에서 변동될 수 있음.")

    d.divider_slide("Ⅰ. 사업개요", toc, active="Ⅰ. 사업개요")
    s = d.content_slide("사업개요", "1. 투자자산 개요",
                        lead="연면적 28,881평 규모의 신축 프라임 상온 물류센터임",
                        subs=["자동화 설비·고스펙 냉난방으로 우량 3PL 임차인 선호 스펙"])
    d.band(s, "대상자산 개요")
    d.table(s, [["구 분", "내 용"],
                ["물건명", "샘플 물류센터"],
                ["소재지", "경기도 ○○시 ○○동"],
                ["연면적", "95,474.59㎡ (28,881평)"],
                ["준공일", "2023-11-24"]],
            col_widths=[2.6, 8.1])
    d.footnote(s, "* 출처: 00_원본자료/processed/ (파일·버전 기재)")
    d.closing_slide()
    p = d.save(out)
    print(f"생성 완료: {p} ({len(d.prs.slides.__iter__.__self__._sldIdLst)}장)")


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "--demo":
        _demo(sys.argv[2])
    else:
        print(__doc__)
