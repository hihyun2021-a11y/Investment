#!/usr/bin/env python3
"""프로젝트 내 모든 원본 파일의 텍스트를 00_원본자료/text/ 에 추출한다 (PDF/DOCX/PPTX/XLSX/XLSM).
사용법: python3 platform/scripts/extract_text.py <코드명>
- `00_원본자료/raw/`(접수 대기함)와 분류 완료된 업무 폴더(01~08)를 모두 훑는다.
- 원본은 수정하지 않는다. 이미 추출된 파일은 원본이 더 새로울 때만 다시 추출한다.
- Bash가 없는 에이전트도 Read 도구로 원본 내용을 읽을 수 있게 하는 캐시다."""
import sys, pathlib

ROOT = pathlib.Path(__file__).resolve().parents[2]
code = sys.argv[1] if len(sys.argv) > 1 else sys.exit("사용법: extract_text.py <코드명>")
proj = ROOT / "projects" / code
if not proj.is_dir():
    sys.exit(f"프로젝트 없음: {proj}")
out = proj / "00_원본자료" / "text"
out.mkdir(parents=True, exist_ok=True)

def pdf(p):
    import pymupdf
    d = pymupdf.open(p)
    return "\n".join(f"\n===== [p.{i+1}] =====\n{pg.get_text()}" for i, pg in enumerate(d))

def docx_(p):
    import docx
    d = docx.Document(p); parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        parts.append("\n[표]")
        for r in t.rows:
            parts.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(parts)

def pptx_(p):
    from pptx import Presentation
    lines = []
    for i, s in enumerate(Presentation(p).slides, 1):
        lines.append(f"\n===== [slide {i}] =====")
        for sh in s.shapes:
            if sh.has_text_frame: lines.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows: lines.append(" | ".join(c.text for c in r.cells))
    return "\n".join(lines)

def xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True); lines = []
    for ws in wb.worksheets:
        lines.append(f"\n===== [sheet {ws.title}] dims={ws.max_row}x{ws.max_column} =====")
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= 400: lines.append("... (400행 초과 생략)"); break
            if any(v is not None for v in row):
                lines.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)

H = {".pdf": pdf, ".docx": docx_, ".pptx": pptx_, ".xlsx": xlsx, ".xlsm": xlsx}
n = skipped = 0
seen = {}
for f in sorted(proj.rglob("*")):
    if not f.is_file() or f.name.startswith(".") or f.suffix.lower() not in H: continue
    if out in f.parents or (proj / "00_원본자료" / "processed") in f.parents: continue
    rel = f.relative_to(proj)
    if f.name in seen:
        print(f"중복 파일명 건너뜀: {rel} (이미 {seen[f.name]})"); continue
    seen[f.name] = rel
    t = out / (f.name + ".txt")
    if t.exists() and t.stat().st_mtime >= f.stat().st_mtime: skipped += 1; continue
    try:
        t.write_text(f"# 텍스트 추출본 (원본: {rel}) — 인용 시 원본 파일명·페이지를 확인할 것\n" + H[f.suffix.lower()](f), encoding="utf-8")
        size = t.stat().st_size
        tag = " [스캔-OCR필요: 텍스트 없음]" if size < 400 else ""
        print(f"추출: {rel} → text/{t.name} ({size//1024} KB){tag}"); n += 1
    except Exception as e:
        print(f"실패: {rel}: {e}")
print(f"완료: {n}건 추출, {skipped}건 최신(건너뜀)")
